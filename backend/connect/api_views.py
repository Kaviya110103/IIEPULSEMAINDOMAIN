import csv
import io
import re
import calendar
import mimetypes
from pathlib import Path
from html import escape
from django.contrib.auth import authenticate
from django.contrib.auth.hashers import check_password, make_password
from django.contrib.auth.models import User
from django.db.models import Q, Count
from django.db.models.functions import Lower, Trim
from django.db import DatabaseError, IntegrityError, OperationalError
from django.utils import timezone
from django.utils.dateparse import parse_date
from rest_framework import generics, status
from rest_framework.decorators import api_view, permission_classes, parser_classes
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from rest_framework.permissions import IsAuthenticated, AllowAny, IsAdminUser
from rest_framework.exceptions import PermissionDenied
from rest_framework.response import Response
from rest_framework.views import APIView
from rest_framework_simplejwt.tokens import RefreshToken
from connect.permissions import is_admin_user, is_staff_employee
# Add this import at the top with your other imports
from django.contrib.auth.decorators import login_required
from datetime import datetime, timedelta
import PyPDF2
from django.db import transaction  # Add this at the top of your views.py
from django.core.cache import cache
from django.conf import settings
from django.http import FileResponse, HttpResponse
import os
import random
import logging
import threading
import requests

# Logger for exception capture
logger = logging.getLogger(__name__)


def get_portal_url(request=None):
    configured_url = getattr(settings, 'APP_PORTAL_URL', '').rstrip('/')
    if configured_url:
        return configured_url
    if request is not None:
        return request.build_absolute_uri('/').rstrip('/')
    return ''

# -- BREVO EMAIL HELPER -------------------------------------------------------
def send_brevo_email(to_email, subject, html_content, text_content=""):
    api_key = os.getenv('BREVO_API_KEY', getattr(settings, 'BREVO_API_KEY', ''))
    from_email = os.getenv('DEFAULT_FROM_EMAIL', getattr(settings, 'DEFAULT_FROM_EMAIL', ''))
    from_name = os.getenv('DEFAULT_FROM_NAME', getattr(settings, 'DEFAULT_FROM_NAME', 'IIE Pulse'))

    if not api_key:
        raise RuntimeError('BREVO_API_KEY is not configured')
    if not from_email:
        raise RuntimeError('DEFAULT_FROM_EMAIL is not configured')

    recipients = to_email if isinstance(to_email, list) else [to_email]
    payload = {
        'sender': {'email': from_email, 'name': from_name},
        'to': [{'email': email} for email in recipients if email],
        'subject': subject,
        'htmlContent': html_content,
    }
    if text_content:
        payload['textContent'] = text_content

    response = requests.post(
        'https://api.brevo.com/v3/smtp/email',
        headers={
            'accept': 'application/json',
            'api-key': api_key,
            'content-type': 'application/json',
        },
        json=payload,
        timeout=20,
    )

    if response.status_code >= 400:
        logger.error(
            'Brevo email failed status=%s response=%s',
            response.status_code,
            response.text,
        )
        response.raise_for_status()

    print(f"Brevo email sent to {recipients}: {subject}")
    logger.info('Brevo email sent to %s subject=%s', recipients, subject)
    return response.json() if response.content else {}


def send_email_async(subject, message, recipient_list):
    """
    Send email asynchronously through Brevo HTTP API. Errors are logged but do
    not fail the caller.
    """
    def send_in_thread():
        try:
            send_brevo_email(
                to_email=recipient_list,
                subject=subject,
                html_content=f"<pre style=\"font-family: Arial, sans-serif; white-space: pre-wrap;\">{escape(message)}</pre>",
                text_content=message,
            )
        except Exception as e:
            print(f"Email error: {e}")
            logger.exception("Failed to send Brevo email to %s", recipient_list)

    thread = threading.Thread(target=send_in_thread, daemon=True)
    thread.start()


from .models import (
    Courses, Employee, Batches, Students,
    StudentAttendance, StudyMaterial, StudyMaterialAssignment, QuizTest, Question, AssignedTest, UserActivity,
    StudentLoginRatingEvent,
    StaffLeaveRequest, StudentLeaveApplication, SupportRequest, StudentSupportRequest,
    CourseSession, DailySessionCompletion, StudentSessionStatus, Student_Session_Progress, DoubtResponse, SessionNotification,
    Announcement, CounselorAnnouncement, CounselorLeaveRequest, CounselorSupportRequest,
    Quiz, QuizQuestion, QuizAttempt, QuizAnswer,
    CompletedStudent, SessionCompletionRequest, TestResult,FeePaymentRequest,
    GalleryItem, VlogItem, NewsItem, CalendarEvent, Referral,
    PublicUser, PublicUserActivity, PublicPracticeResult,
)
from .serializers import (
    CourseSerializer, EmployeeSerializer,
    BatchSerializer, StudentSerializer,
    AttendanceSerializer, StudyMaterialSerializer, TestSerializer,
    QuestionSerializer, AssignedTestSerializer, StaffLeaveRequestSerializer,
    StudentLeaveApplicationSerializer, SupportRequestSerializer,
    StudentSupportRequestSerializer, CourseSessionSerializer,
    AnnouncementSerializer, CounselorLeaveRequestSerializer,
    CounselorSupportRequestSerializer, QuizSerializer, QuizQuestionSerializer,
    QuizAttemptSerializer, CompletedStudentSerializer,
    SessionCompletionRequestSerializer, LoginSerializer,
    GalleryItemSerializer, VlogItemSerializer, NewsItemSerializer,
    CalendarEventSerializer, ReferralSerializer, CounselorAnnouncementSerializer, strip_unsupported_mysql_chars,
)


def get_tokens_for_user(user):
    refresh = RefreshToken.for_user(user)
    return {'refresh': str(refresh), 'access': str(refresh.access_token)}


MSG91_WIDGET_BASE_URL = 'https://control.msg91.com/api/v5/widget'
OTP_SEND_COOLDOWN_SECONDS = 60
OTP_SEND_LIMIT_PER_HOUR = 5
OTP_VERIFY_LIMIT = 5
OTP_REQUEST_TTL_SECONDS = 10 * 60
OTP_VERIFY_TTL_SECONDS = 15 * 60
OTP_HOUR_SECONDS = 60 * 60


def get_client_ip(request):
    forwarded_for = request.META.get('HTTP_X_FORWARDED_FOR')
    if forwarded_for:
        return forwarded_for.split(',')[0].strip()
    return request.META.get('REMOTE_ADDR', 'unknown')


def normalize_indian_mobile(value):
    digits = re.sub(r'\D+', '', str(value or ''))
    if len(digits) == 10:
        return f'91{digits}'
    if len(digits) == 12 and digits.startswith('91'):
        return digits
    return ''


def display_mobile(normalized_mobile):
    if normalized_mobile and len(normalized_mobile) == 12 and normalized_mobile.startswith('91'):
        return normalized_mobile[2:]
    return normalized_mobile or ''


def cache_counter(key, timeout):
    value = cache.get(key)
    if value is None:
        cache.set(key, 1, timeout=timeout)
        return 1
    try:
        value = int(value) + 1
    except (TypeError, ValueError):
        value = 1
    cache.set(key, value, timeout=timeout)
    return value


def check_msg91_config():
    widget_id = getattr(settings, 'MSG91_WIDGET_ID', '').strip()
    auth_token = getattr(settings, 'MSG91_AUTH_TOKEN', '').strip()
    if not widget_id or not auth_token:
        return None, None, Response(
            {'error': 'OTP login is temporarily unavailable. Please contact support.'},
            status=status.HTTP_503_SERVICE_UNAVAILABLE,
        )
    return widget_id, auth_token, None


def find_active_student_by_mobile(normalized_mobile):
    matched_students = []
    for student in Students.objects.select_related('user').filter(user__isnull=False):
        if normalize_indian_mobile(student.mobile_no) == normalized_mobile:
            matched_students.append(student)

    if not matched_students:
        return None, Response({'error': 'No active student account found for this mobile number.'}, status=404)

    if len(matched_students) > 1:
        return None, Response(
            {'error': 'Multiple student accounts use this mobile number. Please contact support.'},
            status=409,
        )

    student = matched_students[0]
    if not student.user or not student.user.is_active:
        return None, Response({'error': 'This student account is inactive. Please contact support.'}, status=403)

    return student, None


def build_student_login_payload(user, student):
    record_user_login(user, 'student', student=student)
    return {
        **get_tokens_for_user(user),
        'user_type': 'student',
        'student_id': student.student_id,
        'student_pk': student.id,
        'name': f"{student.first_name} {student.last_name or ''}".strip(),
    }


def msg91_post(endpoint, payload, timeout=15):
    response = requests.post(
        f'{MSG91_WIDGET_BASE_URL}{endpoint}',
        json=payload,
        headers={'Content-Type': 'application/json'},
        timeout=timeout,
    )
    response.raise_for_status()
    try:
        return response.json()
    except ValueError:
        return {}


def redact_msg91_debug_value(value, secrets=None):
    text = str(value or '')
    for secret in secrets or []:
        if secret:
            text = text.replace(str(secret), '[REDACTED]')
    text = re.sub(r'\b91(\d{2})\d{6}(\d{2})\b', r'91\1******\2', text)
    text = re.sub(r'\b(\d{2})\d{6}(\d{2})\b', r'\1******\2', text)
    return text


def msg91_verify_access_token(auth_token, access_token, timeout=15):
    response = requests.post(
        f'{MSG91_WIDGET_BASE_URL}/verifyAccessToken',
        data={'authkey': auth_token, 'access-token': access_token},
        timeout=timeout,
    )
    response.raise_for_status()
    try:
        return response.json()
    except ValueError:
        return {}


def extract_nested_value(data, names):
    if isinstance(data, dict):
        for key, value in data.items():
            if str(key).lower().replace('_', '').replace('-', '') in names and value:
                return value
        for value in data.values():
            found = extract_nested_value(value, names)
            if found:
                return found
    elif isinstance(data, list):
        for item in data:
            found = extract_nested_value(item, names)
            if found:
                return found
    return None


def extract_msg91_req_id(data):
    return extract_nested_value(data, {'reqid', 'requestid'})


def extract_msg91_access_token(data):
    token = extract_nested_value(data, {'accesstoken', 'access_token', 'token', 'jwttoken'})
    if token:
        return str(token)
    message = data.get('message') if isinstance(data, dict) else None
    if isinstance(message, str) and len(message) > 20 and '.' in message:
        return message
    return None


def extract_msg91_verified_mobile(data):
    value = extract_nested_value(data, {'identifier', 'mobile', 'mobileno', 'phonenumber', 'phone', 'number'})
    return normalize_indian_mobile(value)


ACTIVITY_TIMEOUT = timedelta(minutes=5)


def close_stale_activity_sessions(qs=None):
    now = timezone.now()
    cutoff = now - ACTIVITY_TIMEOUT
    stale_qs = qs or UserActivity.objects.all()
    stale_qs = stale_qs.filter(
        logout_time__isnull=True,
        last_seen__lt=cutoff,
    )

    for activity_id, last_seen in stale_qs.values_list('id', 'last_seen'):
        logout_at = last_seen + ACTIVITY_TIMEOUT if last_seen else cutoff
        UserActivity.objects.filter(id=activity_id, logout_time__isnull=True).update(
            logout_time=logout_at,
            last_seen=logout_at,
        )


def record_user_login(user, user_type, employee=None, student=None):
    now = timezone.now()
    try:
        close_stale_activity_sessions(UserActivity.objects.filter(user=user))
        UserActivity.objects.filter(
            user=user,
            logout_time__isnull=True,
        ).update(logout_time=now, last_seen=now)
        UserActivity.objects.create(
            user=user,
            user_type=user_type,
            employee=employee,
            student=student,
            login_time=now,
            last_seen=now,
        )
        if user_type == 'student' and student:
            record_student_login_rating_event(user, student, 'login', now)
    except Exception:
        logger.exception("Failed to record user login for user_id=%s", getattr(user, 'id', None))


def record_student_login_rating_event(user, student, event_type, occurred_at=None):
    if not user or not student:
        return False
    try:
        StudentLoginRatingEvent.objects.create(
            user=user,
            student=student,
            event_type=event_type,
            occurred_at=occurred_at or timezone.now(),
        )
        return True
    except (DatabaseError, OperationalError):
        logger.warning("Student login rating event table is unavailable. Run migrations.")
        return False
    except Exception:
        logger.exception("Failed to record student login rating event for student_id=%s", getattr(student, 'id', None))
        return False


def record_user_logout(user):
    try:
        activity = UserActivity.objects.filter(
            user=user,
            logout_time__isnull=True,
        ).order_by('-login_time').first()
        if activity:
            now = timezone.now()
            activity.logout_time = now
            activity.last_seen = now
            activity.save(update_fields=['logout_time', 'last_seen'])
    except Exception:
        logger.exception("Failed to record user logout for user_id=%s", getattr(user, 'id', None))


def get_week_start(day=None):
    current_day = day or timezone.localdate()
    return current_day - timedelta(days=current_day.weekday())


def get_local_datetime_range(start_date, days):
    start_dt = timezone.make_aware(datetime.combine(start_date, datetime.min.time()))
    end_dt = timezone.make_aware(datetime.combine(start_date + timedelta(days=days), datetime.min.time()))
    return start_dt, end_dt


def sync_student_login_rating_events_from_activity(student, start_date, week_end):
    try:
        start_dt, end_dt = get_local_datetime_range(start_date, (week_end - start_date).days + 1)
        activity_qs = UserActivity.objects.filter(
            student=student,
            user_type='student',
            login_time__gte=start_dt,
            login_time__lt=end_dt,
        ).values('user_id', 'login_time')

        created = 0
        for item in activity_qs:
            exists = StudentLoginRatingEvent.objects.filter(
                student=student,
                event_type='login',
                occurred_at=item['login_time'],
            ).exists()
            if exists:
                continue
            StudentLoginRatingEvent.objects.create(
                user_id=item['user_id'],
                student=student,
                event_type='login',
                occurred_at=item['login_time'],
            )
            created += 1
        return created
    except (DatabaseError, OperationalError):
        return 0


def build_student_weekly_login_rating(student, week_start=None):
    start_date = week_start or get_week_start()
    rating_days = [start_date + timedelta(days=offset) for offset in range(5)]
    week_end = start_date + timedelta(days=6)
    activity_counts = {day: 0 for day in rating_days}

    try:
        sync_student_login_rating_events_from_activity(student, start_date, week_end)
        start_dt, end_dt = get_local_datetime_range(start_date, 7)
        event_qs = StudentLoginRatingEvent.objects.filter(
            student=student,
            occurred_at__gte=start_dt,
            occurred_at__lt=end_dt,
        ).values('occurred_at')
        for item in event_qs:
            event_day = timezone.localtime(item['occurred_at']).date()
            if event_day in activity_counts:
                activity_counts[event_day] += 1
    except (DatabaseError, OperationalError):
        start_dt, end_dt = get_local_datetime_range(start_date, 7)
        activity_qs = UserActivity.objects.filter(
            student=student,
            user_type='student',
            login_time__gte=start_dt,
            login_time__lt=end_dt,
        ).values('login_time')
        for item in activity_qs:
            login_day = timezone.localtime(item['login_time']).date()
            if login_day in activity_counts:
                activity_counts[login_day] += 1

    days = []
    stars = 0
    for day in rating_days:
        activity_count = activity_counts.get(day, 0)
        earned = activity_count >= 2
        if earned:
            stars += 1
        days.append({
            'date': day.isoformat(),
            'day': day.strftime('%a'),
            'login_count': activity_count,
            'activity_count': activity_count,
            'earned': earned,
        })

    return {
        'week_start': start_date.isoformat(),
        'week_end': week_end.isoformat(),
        'stars': stars,
        'max_stars': 5,
        'days': days,
    }


def get_current_student_for_request(request):
    try:
        return Students.objects.get(user=request.user)
    except Students.DoesNotExist:
        return None


def prepare_activity_monitoring_queryset(qs):
    close_stale_activity_sessions(qs)
    return qs


def get_active_students_queryset():
    qs = Students.objects.all()
    try:
        full_completed_ids = CompletedStudent.objects.filter(
            completion_type='full'
        ).values_list('original_student_id', flat=True)
    except DatabaseError:
        logger.exception("Database error while querying fully completed students")
        full_completed_ids = []

    student_id_values = []
    student_pk_values = []
    for completed_id in full_completed_ids:
        if not completed_id:
            continue
        completed_id = str(completed_id).strip()
        student_id_values.append(completed_id)
        if completed_id.isdigit():
            student_pk_values.append(int(completed_id))

    if not student_id_values and not student_pk_values:
        return qs

    exclude_filter = Q()
    if student_id_values:
        exclude_filter |= Q(student_id__in=student_id_values)
    if student_pk_values:
        exclude_filter |= Q(id__in=student_pk_values)

    return qs.exclude(exclude_filter)


def apply_activity_date_filters(qs, params):
    date_from = parse_date(params.get('date_from') or '')
    date_to = parse_date(params.get('date_to') or '')
    if date_from:
        qs = qs.filter(login_time__date__gte=date_from)
    if date_to:
        qs = qs.filter(login_time__date__lte=date_to)
    return qs


def format_employee_activity(activity):
    employee = activity.employee
    return {
        'id': activity.id,
        'name': f"{employee.first_name} {employee.last_name or ''}".strip() if employee else activity.user.username,
        'email': employee.email if employee else activity.user.email,
        'staff_id': employee.staff_id if employee else '',
        'designation': employee.designation if employee else '',
        'branch': employee.branch if employee else '',
        'login_time': activity.login_time,
        'logout_time': activity.logout_time,
        'last_seen': activity.last_seen,
    }


def format_student_activity(activity):
    student = activity.student
    staff = student.assigned_staff if student else None
    return {
        'id': activity.id,
        'name': f"{student.first_name} {student.last_name or ''}".strip() if student else activity.user.username,
        'email': student.email if student else activity.user.email,
        'student_id': student.student_id if student else '',
        'branch': student.branch if student else '',
        'staff_id': staff.id if staff else None,
        'staff_name': f"{staff.first_name} {staff.last_name or ''}".strip() if staff else '',
        'login_time': activity.login_time,
        'logout_time': activity.logout_time,
        'last_seen': activity.last_seen,
    }


def build_monitoring_pdf_response(title, subtitle, columns, rows, filename):
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.graphics.shapes import Circle, Drawing, Rect, String, Wedge
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    def pdf_text(value):
        if value is None or value == '':
            return '-'
        if hasattr(value, 'isoformat'):
            try:
                value = timezone.localtime(value)
            except Exception:
                pass
            return value.strftime('%d %b %Y, %I:%M %p')
        return str(value)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'MonitoringTitle',
        parent=styles['Title'],
        fontName='Helvetica-Bold',
        fontSize=18,
        alignment=TA_CENTER,
        textColor=colors.HexColor('#0f1b2d'),
        spaceAfter=6,
    )
    subtitle_style = ParagraphStyle(
        'MonitoringSubtitle',
        parent=styles['Normal'],
        fontSize=9,
        alignment=TA_CENTER,
        textColor=colors.HexColor('#64748b'),
        spaceAfter=12,
    )
    cell_style = ParagraphStyle(
        'MonitoringCell',
        parent=styles['Normal'],
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#0f172a'),
    )
    header_style = ParagraphStyle(
        'MonitoringHeaderCell',
        parent=cell_style,
        fontName='Helvetica-Bold',
        textColor=colors.white,
    )

    data = [[Paragraph(label, header_style) for label, _ in columns]]
    for row in rows:
        data.append([Paragraph(pdf_text(row.get(key)), cell_style) for _, key in columns])

    if len(data) == 1:
        data.append([Paragraph('No records found', cell_style)] + [''] * (len(columns) - 1))

    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1572e8')),
        ('GRID', (0, 0), (-1, -1), 0.35, colors.HexColor('#d9e2ec')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))

    generated_at = timezone.localtime(timezone.now()).strftime('%d %b %Y, %I:%M %p')
    logo_path = Path(settings.BASE_DIR) / 'connect' / 'assets' / 'IIE.png'
    header_left = ''
    if logo_path.exists():
        header_left = Image(str(logo_path), width=28 * mm, height=18 * mm)

    header = Table(
        [[
            header_left,
            [
                Paragraph(title, title_style),
                Paragraph(f"{subtitle} | Generated: {generated_at} | Records: {len(rows)}", subtitle_style),
            ],
            '',
        ]],
        colWidths=[35 * mm, None, 35 * mm],
    )
    header.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 0),
        ('TOPPADDING', (0, 0), (-1, -1), 0),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
    ]))

    story = [
        header,
        Spacer(1, 4),
        table,
    ]
    doc.build(story)
    buffer.seek(0)

    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def generate_staff_id():
    """Generate unique Staff ID based on highest existing EMP number"""

    staff_ids = Employee.objects.values_list('staff_id', flat=True)

    max_num = 0

    for sid in staff_ids:
        match = re.search(r'EMP(\d+)', str(sid))
        if match:
            max_num = max(max_num, int(match.group(1)))

    return f"EMP{max_num + 1:04d}"

import re

def generate_batch_number(branch):
    """Generate unique branch-specific batch number using highest existing number"""
    branch_prefix_map = {
        '100ft': '100FT',
        'hopes': 'HOPES',
        'kuniyamuthur': 'KUNIYA',
        'kunniyamuthur': 'KUNIYA',
    }

    prefix = branch_prefix_map.get(branch.lower(), 'BAT')

    if prefix == 'BAT':
        pattern = 'BAT'
    else:
        pattern = f'{prefix}-BAT'

    existing_numbers = Batches.objects.filter(
        batch_number__startswith=pattern
    ).values_list('batch_number', flat=True)

    max_num = 0

    for batch_number in existing_numbers:
        match = re.search(r'BAT(\d+)$', str(batch_number))
        if match:
            max_num = max(max_num, int(match.group(1)))

    next_num = max_num + 1

    if prefix == 'BAT':
        return f'BAT{next_num:04d}'

    return f'{prefix}-BAT{next_num:03d}'


def parse_course_duration(duration):
    value = str(duration or '').strip().lower()
    match = re.fullmatch(r'([1-9]\d*)\s*(day|days|month|months)', value)
    if not match:
        raise ValueError('Selected course has missing or invalid duration. Please update the course duration first.')
    return int(match.group(1)), match.group(2)


def add_calendar_months(start, months):
    month_index = start.month - 1 + months
    year = start.year + month_index // 12
    month = month_index % 12 + 1
    day = min(start.day, calendar.monthrange(year, month)[1])
    return start.replace(year=year, month=month, day=day)


def calculate_batch_end_date(course, start_date_value):
    try:
        start = datetime.strptime(str(start_date_value), '%Y-%m-%d').date()
    except (TypeError, ValueError):
        raise ValueError('Enter a valid batch start date.')

    amount, unit = parse_course_duration(getattr(course, 'duration', ''))
    if unit.startswith('month'):
        return add_calendar_months(start, amount)
    return start + timedelta(days=amount)

def generate_student_id(branch):
    branch_prefix_map = {
        '100ft': '100FT',
        'hopes': 'HOPES',
        'kuniyamuthur': 'KUNIYA',
        'kunniyamuthur': 'KUNIYA',
    }

    prefix = branch_prefix_map.get(branch.lower(), 'STU')
    pattern = f"{prefix}-STU"

    existing_ids = Students.objects.filter(
        student_id__startswith=pattern
    ).values_list('student_id', flat=True)

    max_number = 0

    for sid in existing_ids:
        match = re.search(r'STU(\d+)$', sid)
        if match:
            max_number = max(max_number, int(match.group(1)))

    next_number = max_number + 1
    return f"{prefix}-STU{next_number:03d}"

# -- AUTH ---------------------------------------------------------------------

class LoginView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        serializer = LoginSerializer(data=request.data)
        if not serializer.is_valid():
            return Response(serializer.errors, status=400)

        username = serializer.validated_data['username'].strip().lower()
        password = serializer.validated_data['password']
        user_type = serializer.validated_data['user_type']

        if user_type == 'student':
            student = Students.objects.filter(
                Q(student_id__iexact=username) | Q(email__iexact=username) | Q(user__username__iexact=username)
            ).select_related('user').first()
            if student:
                username = student.user.username

        user = authenticate(username=username, password=password)

        if user is None:
            return Response({'error': 'Invalid username or password.'}, status=401)

        if user_type == 'admin':
            if not (user.is_superuser or user.is_staff):
                return Response({'error': "No admin privileges."}, status=403)
            return Response({**get_tokens_for_user(user), 'user_type': 'admin', 'username': user.username, 'name': user.username})

        elif user_type == 'employee':
            try:
                employee = Employee.objects.get(user=user)
            except Employee.DoesNotExist:
                return Response({'error': 'Employee account not found.'}, status=403)
            record_user_login(user, 'employee', employee=employee)
            return Response({
                **get_tokens_for_user(user),
                'user_type': 'employee',
                'designation': employee.designation,
                'employee_id': employee.id,
                'name': f"{employee.first_name} {employee.last_name or ''}".strip(),
                'branch': employee.branch,
            })

        elif user_type == 'student':
            try:
                student = Students.objects.get(user=user)
            except Students.DoesNotExist:
                return Response({'error': 'Student account not found.'}, status=403)
            return Response(build_student_login_payload(user, student))

        return Response({'error': 'Invalid user type.'}, status=400)


class StudentOtpSendView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        widget_id, auth_token, config_error = check_msg91_config()
        if config_error:
            return config_error

        normalized_mobile = normalize_indian_mobile(request.data.get('mobile_no'))
        if not normalized_mobile:
            return Response({'error': 'Enter a valid 10-digit mobile number.'}, status=400)

        student, lookup_error = find_active_student_by_mobile(normalized_mobile)
        if lookup_error:
            return lookup_error

        ip = get_client_ip(request)
        cooldown_key = f'otp:send:cooldown:{normalized_mobile}'
        if cache.get(cooldown_key):
            return Response({'error': 'Please wait before requesting another OTP.', 'resend_after': OTP_SEND_COOLDOWN_SECONDS}, status=429)

        mobile_count = cache_counter(f'otp:send:hour:mobile:{normalized_mobile}', OTP_HOUR_SECONDS)
        ip_count = cache_counter(f'otp:send:hour:ip:{ip}', OTP_HOUR_SECONDS)
        if mobile_count > OTP_SEND_LIMIT_PER_HOUR or ip_count > OTP_SEND_LIMIT_PER_HOUR * 5:
            return Response({'error': 'Too many OTP requests. Please try again later.'}, status=429)

        try:
            msg91_url = f'{MSG91_WIDGET_BASE_URL}/sendOtpMobile'
            logger.warning('MSG91 send OTP request URL: %s', msg91_url)
            response = requests.post(
                msg91_url,
                json={
                    'widgetId': widget_id,
                    'tokenAuth': auth_token,
                    'identifier': normalized_mobile,
                },
                headers={'Content-Type': 'application/json'},
                timeout=15,
            )
            logger.warning('MSG91 send OTP HTTP status: %s', response.status_code)
            logger.warning('MSG91 send OTP response body: %s', redact_msg91_debug_value(response.text, [auth_token]))
            response.raise_for_status()
            try:
                msg91_response = response.json()
            except ValueError:
                msg91_response = {}
        except requests.RequestException as exc:
            logger.exception('MSG91 send OTP request failed: %s', redact_msg91_debug_value(str(exc), [auth_token]))
            return Response({'error': 'Unable to send OTP right now. Please try again.'}, status=502)

        msg91_type = str(msg91_response.get('type', '')).lower() if isinstance(msg91_response, dict) else ''
        if response.status_code != 200 or msg91_type != 'success':
            logger.warning('MSG91 send OTP failed response: %s', redact_msg91_debug_value(msg91_response, [auth_token]))
            return Response({'error': 'Unable to send OTP right now. Please try again.'}, status=502)

        req_id = extract_msg91_req_id(msg91_response) or str(msg91_response.get('message') or '').strip()
        if not req_id:
            logger.warning('MSG91 send OTP success response did not include a request id: %s', redact_msg91_debug_value(msg91_response, [auth_token]))
            return Response({'error': 'Unable to send OTP right now. Please try again.'}, status=502)

        cache.set(cooldown_key, True, timeout=OTP_SEND_COOLDOWN_SECONDS)
        cache.set(
            f'otp:req:{req_id}',
            {'mobile': normalized_mobile, 'student_id': student.id, 'verify_attempts': 0},
            timeout=OTP_REQUEST_TTL_SECONDS,
        )

        return Response({
            'success': True,
            'message': 'OTP sent successfully.',
            'req_id': req_id,
            'mobile_no': display_mobile(normalized_mobile),
            'resend_after': OTP_SEND_COOLDOWN_SECONDS,
        })


class StudentOtpResendView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        widget_id, auth_token, config_error = check_msg91_config()
        if config_error:
            return config_error

        normalized_mobile = normalize_indian_mobile(request.data.get('mobile_no'))
        req_id = str(request.data.get('req_id') or '').strip()
        if not normalized_mobile or not req_id:
            return Response({'error': 'Mobile number and OTP request id are required.'}, status=400)

        request_state = cache.get(f'otp:req:{req_id}')
        if not request_state or request_state.get('mobile') != normalized_mobile:
            return Response({'error': 'OTP has expired. Please request a new OTP.'}, status=400)

        student, lookup_error = find_active_student_by_mobile(normalized_mobile)
        if lookup_error:
            return lookup_error

        ip = get_client_ip(request)
        cooldown_key = f'otp:send:cooldown:{normalized_mobile}'
        if cache.get(cooldown_key):
            return Response({'error': 'Please wait before requesting another OTP.', 'resend_after': OTP_SEND_COOLDOWN_SECONDS}, status=429)

        mobile_count = cache_counter(f'otp:send:hour:mobile:{normalized_mobile}', OTP_HOUR_SECONDS)
        ip_count = cache_counter(f'otp:send:hour:ip:{ip}', OTP_HOUR_SECONDS)
        if mobile_count > OTP_SEND_LIMIT_PER_HOUR or ip_count > OTP_SEND_LIMIT_PER_HOUR * 5:
            return Response({'error': 'Too many OTP requests. Please try again later.'}, status=429)

        try:
            msg91_response = msg91_post('/retryOtp', {
                'widgetId': widget_id,
                'tokenAuth': auth_token,
                'reqId': req_id,
                'retryChannel': 'SMS',
            })
        except requests.RequestException:
            logger.exception('MSG91 resend OTP request failed')
            return Response({'error': 'Unable to resend OTP right now. Please try again.'}, status=502)

        next_req_id = extract_msg91_req_id(msg91_response) or req_id
        cache.set(cooldown_key, True, timeout=OTP_SEND_COOLDOWN_SECONDS)
        cache.set(
            f'otp:req:{next_req_id}',
            {'mobile': normalized_mobile, 'student_id': student.id, 'verify_attempts': 0},
            timeout=OTP_REQUEST_TTL_SECONDS,
        )

        return Response({
            'message': 'OTP resent successfully.',
            'req_id': next_req_id,
            'mobile_no': display_mobile(normalized_mobile),
            'resend_after': OTP_SEND_COOLDOWN_SECONDS,
        })


class StudentOtpVerifyView(APIView):
    permission_classes = [AllowAny]

    def post(self, request):
        widget_id, auth_token, config_error = check_msg91_config()
        if config_error:
            return config_error

        normalized_mobile = normalize_indian_mobile(request.data.get('mobile_no'))
        req_id = str(request.data.get('req_id') or '').strip()
        otp = re.sub(r'\D+', '', str(request.data.get('otp') or ''))
        if not normalized_mobile or not req_id or not otp:
            return Response({'error': 'Mobile number, request id and OTP are required.'}, status=400)

        request_state = cache.get(f'otp:req:{req_id}')
        if not request_state or request_state.get('mobile') != normalized_mobile:
            return Response({'error': 'OTP has expired. Please request a new OTP.'}, status=400)

        attempts = int(request_state.get('verify_attempts') or 0) + 1
        if attempts > OTP_VERIFY_LIMIT:
            cache.delete(f'otp:req:{req_id}')
            return Response({'error': 'Too many invalid OTP attempts. Please request a new OTP.'}, status=429)

        request_state['verify_attempts'] = attempts
        cache.set(f'otp:req:{req_id}', request_state, timeout=OTP_VERIFY_TTL_SECONDS)

        ip = get_client_ip(request)
        if cache_counter(f'otp:verify:ip:{ip}', OTP_HOUR_SECONDS) > OTP_VERIFY_LIMIT * 10:
            return Response({'error': 'Too many verification attempts. Please try again later.'}, status=429)

        try:
            verify_response = msg91_post('/verifyOtp', {
                'widgetId': widget_id,
                'tokenAuth': auth_token,
                'reqId': req_id,
                'otp': otp,
            })
        except requests.RequestException:
            logger.exception('MSG91 verify OTP request failed')
            return Response({'error': 'Unable to verify OTP right now. Please try again.'}, status=502)

        access_token = extract_msg91_access_token(verify_response)
        if not access_token:
            return Response({'error': 'Invalid OTP or OTP has expired.'}, status=400)

        try:
            token_response = msg91_verify_access_token(auth_token, access_token)
        except requests.RequestException:
            logger.exception('MSG91 access token verification failed')
            return Response({'error': 'Unable to verify OTP right now. Please try again.'}, status=502)

        verified_mobile = extract_msg91_verified_mobile(token_response)
        if verified_mobile and verified_mobile != normalized_mobile:
            return Response({'error': 'OTP verification did not match the requested mobile number.'}, status=400)
        if not verified_mobile:
            return Response({'error': 'Unable to verify mobile number. Please try again.'}, status=400)

        student, lookup_error = find_active_student_by_mobile(normalized_mobile)
        if lookup_error:
            return lookup_error
        if student.id != request_state.get('student_id'):
            return Response({'error': 'Student account mismatch. Please request a new OTP.'}, status=400)

        cache.delete(f'otp:req:{req_id}')
        return Response(build_student_login_payload(student.user, student))


class LogoutView(APIView):
    permission_classes = [IsAuthenticated]

    def post(self, request):
        record_user_logout(request.user)
        try:
            token = RefreshToken(request.data.get('refresh', ''))
            token.blacklist()
        except Exception:
            pass
        return Response({'message': 'Logged out successfully.'})


# -- DASHBOARD -----------------------------------------------------------------

class AdminDashboardView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        if not is_admin_user(request.user):
            return Response({'error': 'Admin access required.'}, status=403)
        active_students = get_active_students_queryset()
        branch_counts = active_students.exclude(
            branch__isnull=True
        ).exclude(
            branch=''
        ).values('branch').annotate(
            count=Count('id')
        ).order_by('branch')
        batch_branch_counts = Batches.objects.exclude(
            branch__isnull=True
        ).exclude(
            branch=''
        ).values('branch').annotate(
            count=Count('id')
        ).order_by('branch')
        def canonical_branch(value):
            value = re.sub(r'\s+', ' ', value or '').strip().lower()
            if value == 'kunniyamuthur':
                return 'kuniyamuthur'
            return value

        def branch_values(branch):
            if branch == '100ft':
                return ['100ft', '100FT']
            if branch == 'hopes':
                return ['hopes', 'Hopes', 'HOPES']
            if branch == 'kuniyamuthur':
                return ['kuniyamuthur', 'Kuniyamuthur', 'KUNIYAMUTHUR', 'kunniyamuthur', 'Kunniyamuthur', 'KUNNIYAMUTHUR']
            return [branch]

        completed_students_qs = CompletedStudent.objects.filter(completion_type='full')
        completed_counts_by_branch = {}
        for item in completed_students_qs.exclude(branch__isnull=True).exclude(branch='').values('branch').annotate(count=Count('id')):
            key = canonical_branch(item['branch'])
            if key:
                completed_counts_by_branch[key] = completed_counts_by_branch.get(key, 0) + item['count']
        completed_branch_counts = [
            {'branch': branch, 'count': count}
            for branch, count in sorted(
                completed_counts_by_branch.items(),
                key=lambda item: (['100ft', 'hopes', 'kuniyamuthur'].index(item[0]) if item[0] in ['100ft', 'hopes', 'kuniyamuthur'] else 3, item[0])
            )
        ]

        def clamp_percentage(value):
            return max(0, min(100, round(value)))

        branch_order = ['100ft', 'hopes', 'kuniyamuthur']
        branch_set = {
            canonical_branch(branch)
            for branch in list(Employee.objects.exclude(branch__isnull=True).exclude(branch='').values_list('branch', flat=True)) +
            list(active_students.exclude(branch__isnull=True).exclude(branch='').values_list('branch', flat=True))
            if canonical_branch(branch)
        }
        branches = sorted(branch_set, key=lambda branch: (branch_order.index(branch) if branch in branch_order else len(branch_order), branch))
        branch_usage_stats = []
        branch_tracking_cards = []
        for branch in branches:
            values = branch_values(branch)
            active_branch_students = active_students.filter(branch__in=values)
            staff_total = Employee.objects.filter(branch__in=values).count()
            student_total = active_branch_students.count()
            staff_logged = UserActivity.objects.filter(
                user_type='employee',
                employee__branch__in=values,
                employee__isnull=False,
            ).values('employee_id').distinct().count()
            student_logged = UserActivity.objects.filter(
                user_type='student',
                student_id__in=active_branch_students.values_list('id', flat=True),
                student__isnull=False,
            ).values('student_id').distinct().count()
            total_users = staff_total + student_total
            logged_users = staff_logged + student_logged
            usage_percentage = clamp_percentage((logged_users / total_users) * 100) if total_users else 0
            branch_usage_stats.append({
                'branch': branch,
                'usage_percentage': usage_percentage,
                'staff_usage_percentage': clamp_percentage((staff_logged / staff_total) * 100) if staff_total else 0,
                'student_usage_percentage': clamp_percentage((student_logged / student_total) * 100) if student_total else 0,
            })
            staff_qs = Employee.objects.filter(branch__in=values)
            branch_batches = Batches.objects.filter(branch__in=values)
            branch_students = active_branch_students
            branch_attendance = StudentAttendance.objects.filter(staff__branch__in=values).count()
            branch_daily_sessions = DailySessionCompletion.objects.filter(
                faculty__branch__in=values,
                completed=True,
            ).values('session_id').distinct().count()
            branch_direct_sessions = CourseSession.objects.filter(
                batch__branch__in=values,
                staff_completed=True,
            ).count()
            branch_sessions_completed = max(branch_daily_sessions, branch_direct_sessions)
            branch_material_uploads = StudyMaterial.objects.filter(uploaded_by__branch__in=values).count()
            branch_material_assignments = StudyMaterialAssignment.objects.filter(assigned_by__branch__in=values).count()
            branch_quizzes_created = Quiz.objects.filter(created_by__branch__in=values).count()
            branch_login_days = UserActivity.objects.filter(
                user_type='employee',
                employee__branch__in=values,
                login_time__gte=timezone.now() - timedelta(days=7),
            ).count()
            branch_score_parts = [
                min(branch_sessions_completed * 2, 30),
                min((branch_material_uploads + branch_material_assignments) * 3, 25),
                min(branch_quizzes_created * 4, 25),
                min(branch_login_days * 3, 20),
            ]
            branch_tracking_cards.append({
                'branch': branch,
                'staff_count': staff_qs.count(),
                'student_count': branch_students.count(),
                'batch_count': branch_batches.count(),
                'attendance_marked_count': branch_attendance,
                'sessions_completed': branch_sessions_completed,
                'materials_count': branch_material_uploads + branch_material_assignments,
                'tests_count': 0,
                'quizzes_count': branch_quizzes_created,
                'activity_score': min(round(sum(branch_score_parts), 1), 100),
            })
        return Response({
            'student_count': active_students.count(),
            'student_branch_counts': [
                {'branch': item['branch'], 'count': item['count']}
                for item in branch_counts
            ],
            'batch_branch_counts': [
                {'branch': item['branch'], 'count': item['count']}
                for item in batch_branch_counts
            ],
            'completed_branch_counts': [
                {'branch': item['branch'], 'count': item['count']}
                for item in completed_branch_counts
            ],
            'branch_usage_stats': branch_usage_stats,
            'branch_tracking_cards': branch_tracking_cards,
            'employee_count': Employee.objects.count(),
            'mentor_count': Employee.objects.filter(designation__iexact='mentor').count(),
            'course_count': Courses.objects.count(),
            'batch_count': Batches.objects.count(),
            'completed_count': completed_students_qs.count(),
            'counselor_count': Employee.objects.filter(designation__iexact='counselor').count(),
            'trainer_count': Employee.objects.filter(designation__iexact='trainer').count(),
            'next_staff_id': generate_staff_id(),
        })


class EmployeeDashboardView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        try:
            employee = Employee.objects.get(user=request.user)
        except Employee.DoesNotExist:
            return Response({'error': 'Employee not found'}, status=404)

        completed_ids = CompletedStudent.objects.values_list('original_student_id', flat=True)
        my_batches = Batches.objects.filter(faculty=employee)
        today = timezone.now().date()
        announcements = Announcement.objects.filter(
            Q(recipient_type='all') | Q(recipient_type='staff') | Q(recipient_type='mentors'),
            is_published=True
        ).order_by('-created_at')[:5]

        return Response({
            'employee': EmployeeSerializer(employee).data,
            'my_batches_count': my_batches.count(),
            'my_students_count': Students.objects.filter(
                assigned_batch__faculty=employee
            ).exclude(student_id__in=completed_ids).count(),
            'today_classes_count': my_batches.filter(start_date__lte=today, end_date__gte=today).count(),
            'materials_count': StudyMaterial.objects.filter(uploaded_by=employee).count(),
            'completed_students_count': CompletedStudent.objects.filter(
                batch_number__in=my_batches.values_list('batch_number', flat=True)
            ).count(),
            'announcements_count': announcements.count(),
            'announcements': AnnouncementSerializer(announcements, many=True).data,
        })


class StudentDashboardView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        try:
            student = Students.objects.get(user=request.user)
        except Students.DoesNotExist:
            return Response({'error': 'Student not found'}, status=404)

        attendance = StudentAttendance.objects.filter(student=student)
        total_att = attendance.count()
        present_att = attendance.filter(status='Present').count()

        admin_items = []
        admin_qs = Announcement.objects.prefetch_related('specific_students').filter(
            is_published=True,
            created_by__is_staff=True,
            recipient_type__in=['all', 'students', 'specific'],
        ).order_by('-created_at')
        for announcement in admin_qs:
            if _admin_announcement_visible_to_student(announcement, student):
                is_selected = announcement.specific_students.filter(id=student.id).exists()
                admin_items.append(_serialize_mobile_announcement(
                    AnnouncementSerializer(announcement).data,
                    'admin',
                    'Admin',
                    'Selected Student' if is_selected else None,
                ))

        counselor_qs = CounselorAnnouncement.objects.select_related(
            'specific_batch',
            'created_by',
            'created_by__employee',
        ).prefetch_related('specific_students').filter(
            is_published=True,
            created_by__employee__designation__iexact='counselor',
        ).distinct().order_by('-created_at')
        counselor_items = [
            _serialize_mobile_announcement(
                CounselorAnnouncementSerializer(announcement).data,
                'counselor',
                'Counselor',
                announcement.specific_batch.batch_number
                if announcement.recipient_type == 'specific_batch' and announcement.specific_batch
                else None,
            )
            for announcement in counselor_qs
            if _counselor_announcement_visible_to_student(announcement, student)
        ]

        trainer_user_ids = Employee.objects.filter(
            Q(designation__iexact='trainer') | Q(designation__iexact='mentor')
        ).values_list('user_id', flat=True)
        trainer_qs = CounselorAnnouncement.objects.select_related(
            'specific_batch',
            'specific_batch__faculty',
            'created_by',
        ).prefetch_related('specific_students').filter(
            is_published=True,
            created_by_id__in=trainer_user_ids,
            recipient_type__in=['specific_batch', 'specific_student'],
        ).distinct().order_by('-created_at')
        trainer_items = [
            _serialize_mobile_announcement(
                CounselorAnnouncementSerializer(announcement).data,
                'trainer',
                'Trainer',
                announcement.specific_batch.batch_number
                if announcement.recipient_type == 'specific_batch' and announcement.specific_batch
                else 'Selected Student',
            )
            for announcement in trainer_qs
            if _trainer_announcement_visible_to_student(announcement, student)
        ]

        visible_announcements = sorted(
            admin_items + counselor_items + trainer_items,
            key=lambda item: item.get('created_at') or '',
            reverse=True,
        )

          # -- ADD THIS: Calculate completed sessions count -------------
        completed_sessions_count = Student_Session_Progress.objects.filter(
            student=student,
            completed=True
        ).count()
        # -------------------------------------------------------------

        # Match the student web modules: count assigned items, not only submitted results.
        if student.assigned_batch:
            assigned_test_ids = set(AssignedTest.objects.filter(
                batch=student.assigned_batch
            ).values_list('test_id', flat=True))
            if student.assigned_staff:
                mentor_test_ids = Question.objects.filter(
                    test__created_by=student.assigned_staff,
                ).values_list('test_id', flat=True).distinct()
                assigned_test_ids.update(mentor_test_ids)
            tests_count = len(assigned_test_ids)
            quizzes_count = Quiz.objects.filter(
                batch=student.assigned_batch,
                is_published=True,
            ).distinct().count()
        else:
            tests_count = 0
            quizzes_count = 0
        materials_count = StudyMaterial.objects.filter(
            Q(batch=student.assigned_batch) |
            Q(assignments__batch=student.assigned_batch)
        ).distinct().count() if student.assigned_batch else 0
        # -------------------------------------------------------------


        return Response({
            'student': StudentSerializer(student).data,
            'attendance_percentage': round((present_att / total_att * 100) if total_att else 0, 1),
            'total_classes': total_att,
            'present_classes': present_att,
            'announcements_count': len(visible_announcements),
            'announcements': visible_announcements[:5],
             # -- ADD THESE NEW FIELDS ---------------------------------
            'completed_sessions_count': completed_sessions_count,
            'tests_count': tests_count,
            'quizzes_count': quizzes_count,
            'materials_count': materials_count,
        })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def student_weekly_login_rating(request):
    student = get_current_student_for_request(request)
    if not student:
        return Response({'error': 'Student not found'}, status=404)

    record_student_login_rating_event(request.user, student, 'app_use')

    return Response({
        'current_week': build_student_weekly_login_rating(student),
        'rule': {
            'minimum_logins_per_day': 2,
            'rating_days': 5,
            'reset': 'weekly',
        },
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def student_weekly_login_rating_history(request):
    student = get_current_student_for_request(request)
    if not student:
        return Response({'error': 'Student not found'}, status=404)

    try:
        limit = max(1, min(int(request.query_params.get('limit', 8)), 24))
    except (TypeError, ValueError):
        limit = 8

    current_week_start = get_week_start()
    weeks = [
        build_student_weekly_login_rating(student, current_week_start - timedelta(days=7 * offset))
        for offset in range(1, limit + 1)
    ]
    return Response({'weeks': weeks})


class CounselorDashboardView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        try:
            counselor = Employee.objects.get(user=request.user)
        except Employee.DoesNotExist:
            return Response({'error': 'Not found'}, status=404)

        branch = counselor.branch or ''
        announcements = Announcement.objects.filter(
            Q(recipient_type='all') | Q(recipient_type='counselors') | Q(recipient_type='staff'),
            is_published=True
        ).order_by('-created_at')[:5]

        return Response({
            'counselor': EmployeeSerializer(counselor).data,
            'student_count': Students.objects.filter(branch=branch).count(),
            'mentor_count': Employee.objects.filter(branch=branch, designation__iexact='mentor').count(),
            'counselor_count': Employee.objects.filter(branch=branch, designation__iexact='counselor').count(),
            'course_count': Courses.objects.filter(batches__branch=branch).distinct().count(),
            'batch_count': Batches.objects.filter(branch=branch).count(),
            'pending_completion_requests': SessionCompletionRequest.objects.filter(counselor=counselor, status='pending').count(),
            'announcements': AnnouncementSerializer(announcements, many=True).data,
            'next_batch_id': generate_batch_number(branch) if branch else '',
            'next_student_id': generate_student_id(branch) if branch else '',
        })


# -- STAFF ID GENERATION ------------------------------------------------------

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_next_staff_id(request):
    if not is_staff_employee(request.user):
        return Response({'error': 'Access denied.'}, status=403)
    return Response({'staff_id': generate_staff_id()})


# -- BATCH NUMBER GENERATION --------------------------------------------------

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_next_batch_number(request):
    if not is_staff_employee(request.user):
        return Response({'error': 'Access denied.'}, status=403)
    branch = request.query_params.get('branch', '')
    if not branch:
        return Response({'error': 'Branch is required'}, status=400)
    return Response({
        'batch_number': generate_batch_number(branch),
        'branch': branch,
    })


# -- STUDENT ID GENERATION ----------------------------------------------------

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_next_student_id(request):
    if not is_staff_employee(request.user):
        return Response({'error': 'Access denied.'}, status=403)
    branch = request.query_params.get('branch', '')
    if not branch:
        return Response({'error': 'Branch is required'}, status=400)
    return Response({'student_id': generate_student_id(branch)})


# -- COURSES ------------------------------------------------------------------

class CourseListCreateView(generics.ListCreateAPIView):
    queryset = Courses.objects.all().order_by('-created_at')
    serializer_class = CourseSerializer
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser, JSONParser]


class CourseDetailView(generics.RetrieveUpdateDestroyAPIView):
    queryset = Courses.objects.all()
    serializer_class = CourseSerializer
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def perform_update(self, serializer):
        logsheet_changed = 'course_logsheet' in self.request.FILES
        course = serializer.save()
        if logsheet_changed:
            for batch in Batches.objects.filter(course_name=course):
                sync_missing_sessions_from_logsheet(batch, prefer_course_logsheet=True)


class GalleryItemListCreateView(generics.ListCreateAPIView):
    queryset = GalleryItem.objects.all().order_by('-created_at')
    serializer_class = GalleryItemSerializer
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def get_permissions(self):
        if self.request.method == 'GET':
            return [AllowAny()]
        return [IsAdminUser()]

    def perform_create(self, serializer):
        serializer.save(uploaded_by=self.request.user)


class GalleryItemDetailView(generics.RetrieveUpdateDestroyAPIView):
    queryset = GalleryItem.objects.all()
    serializer_class = GalleryItemSerializer
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def get_permissions(self):
        if self.request.method == 'GET':
            return [AllowAny()]
        return [IsAdminUser()]


class VlogItemListCreateView(generics.ListCreateAPIView):
    queryset = VlogItem.objects.all().order_by('-created_at')
    serializer_class = VlogItemSerializer
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def get_permissions(self):
        if self.request.method == 'GET':
            return [AllowAny()]
        return [IsAdminUser()]

    def perform_create(self, serializer):
        serializer.save(uploaded_by=self.request.user)


class VlogItemDetailView(generics.RetrieveUpdateDestroyAPIView):
    queryset = VlogItem.objects.all()
    serializer_class = VlogItemSerializer
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def get_permissions(self):
        if self.request.method == 'GET':
            return [AllowAny()]
        return [IsAdminUser()]


class NewsItemListCreateView(generics.ListCreateAPIView):
    queryset = NewsItem.objects.all().order_by('-created_at')
    serializer_class = NewsItemSerializer
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def get_permissions(self):
        if self.request.method == 'GET':
            return [AllowAny()]
        return [IsAdminUser()]

    def perform_create(self, serializer):
        serializer.save(uploaded_by=self.request.user)


class NewsItemDetailView(generics.RetrieveUpdateDestroyAPIView):
    queryset = NewsItem.objects.all()
    serializer_class = NewsItemSerializer
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def get_permissions(self):
        if self.request.method == 'GET':
            return [AllowAny()]
        return [IsAdminUser()]


class CalendarEventListCreateView(generics.ListCreateAPIView):
    queryset = CalendarEvent.objects.all().order_by('event_date', 'event_time', '-created_at')
    serializer_class = CalendarEventSerializer
    parser_classes = [JSONParser, FormParser, MultiPartParser]

    def get_permissions(self):
        if self.request.method == 'GET':
            return [AllowAny()]
        return [IsAdminUser()]

    def perform_create(self, serializer):
        serializer.save(uploaded_by=self.request.user)

    def create(self, request, *args, **kwargs):
        data = request.data.copy()
        event_name = strip_unsupported_mysql_chars(data.get('event_name') or data.get('title') or '').strip()
        message = strip_unsupported_mysql_chars(data.get('message') or data.get('description') or '').strip()
        event_date = str(data.get('event_date') or data.get('date') or '').strip()
        event_time = str(data.get('event_time') or data.get('time') or '').strip()

        if event_name:
            data['event_name'] = event_name
        if not message and event_name:
            message = event_name
        if message:
            data['message'] = message

        if event_date:
            for fmt in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y'):
                try:
                    data['event_date'] = datetime.strptime(event_date, fmt).date().isoformat()
                    break
                except ValueError:
                    continue
            else:
                data['event_date'] = event_date

        if event_time:
            event_time = event_time.upper().replace('.', '').strip()
            parsed_time = None
            for fmt in ('%H:%M:%S', '%H:%M', '%I:%M %p', '%I:%M:%S %p'):
                try:
                    parsed_time = datetime.strptime(event_time, fmt).time()
                    break
                except ValueError:
                    continue
            data['event_time'] = parsed_time.strftime('%H:%M:%S') if parsed_time else event_time

        try:
            serializer = self.get_serializer(data=data)
            serializer.is_valid(raise_exception=True)
            self.perform_create(serializer)
            headers = self.get_success_headers(serializer.data)
            return Response(serializer.data, status=status.HTTP_201_CREATED, headers=headers)
        except OperationalError as exc:
            return Response({'error': f'Calendar database error: {exc}'}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)
        except IntegrityError as exc:
            return Response({'error': f'Calendar save error: {exc}'}, status=status.HTTP_400_BAD_REQUEST)


class CalendarEventDetailView(generics.RetrieveUpdateDestroyAPIView):
    queryset = CalendarEvent.objects.all()
    serializer_class = CalendarEventSerializer
    parser_classes = [JSONParser, FormParser, MultiPartParser]

    def get_permissions(self):
        if self.request.method == 'GET':
            return [AllowAny()]
        return [IsAdminUser()]


class ReferralListCreateView(generics.ListCreateAPIView):
    queryset = Referral.objects.all().order_by('-created_at')
    serializer_class = ReferralSerializer
    parser_classes = [JSONParser, FormParser]

    def get_permissions(self):
        if self.request.method == 'POST':
            return [AllowAny()]
        return [IsAdminUser()]


class ReferralDetailView(generics.RetrieveDestroyAPIView):
    queryset = Referral.objects.all()
    serializer_class = ReferralSerializer
    permission_classes = [IsAdminUser]


# -- EMPLOYEES ----------------------------------------------------------------

class EmployeeListView(generics.ListAPIView):
    queryset = Employee.objects.all().order_by('-created_at')
    serializer_class = EmployeeSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        qs = super().get_queryset()
        designation = self.request.query_params.get('designation')
        branch = self.request.query_params.get('branch')
        if designation:
            qs = qs.filter(designation__iexact=designation)
        if branch:
            qs = qs.filter(branch=branch)
        return qs



class EmployeeCreateView(APIView):
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def post(self, request):
        data = request.data
        employee_name = data.get('employee_name', '').strip()
        email = data.get('email', '').strip().lower()
        mobile_no = data.get('mobile_no', '').strip()
        role = data.get('role', data.get('designation', '')).strip()
        branch = data.get('branch', '').strip()
        gender = data.get('gender', '').strip()
        date_of_birth = data.get('date_of_birth', '').strip()
        address = data.get('address', '').strip()
        photo = request.FILES.get('photo')
        id_proof = request.FILES.get('id_proof')

        if not all([employee_name, email, mobile_no, role, branch, date_of_birth]):
            return Response({'error': 'Please fill in all required fields.'}, status=400)

        if not re.match(r'^[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}$', email):
            return Response({'error': 'Please enter a valid email address.'}, status=400)

        if User.objects.filter(username=email).exists():
            return Response({'error': 'This email is already registered!'}, status=400)
        if Employee.objects.filter(email=email).exists():
            return Response({'error': 'Email already exists in employees!'}, status=400)
        if Employee.objects.filter(mobile_no=mobile_no).exists():
            return Response({'error': 'Mobile number already exists!'}, status=400)

        name_parts = employee_name.split(' ', 1)
        first_name = name_parts[0].strip()
        last_name = name_parts[1].strip() if len(name_parts) > 1 else ''
        staff_id = generate_staff_id()

        try:
            user = User.objects.create_user(
                username=email,
                password=mobile_no,
                email=email,
                first_name=first_name,
                last_name=last_name,
            )
            employee = Employee.objects.create(
                user=user,
                staff_id=staff_id,
                email=email,
                first_name=first_name,
                last_name=last_name,
                mobile_no=mobile_no,
                date_of_birth=date_of_birth,
                designation=role,
                branch=branch,
                gender=gender,
                address=address,
                photo=photo if photo else None,
                id_proof=id_proof if id_proof else None,
            )

            # -- Send welcome email with login credentials (async) -------------
            try:
                branch_display = {
                    '100ft': '100 Feet Road',
                    'hopes': 'Hopes',
                    'kuniyamuthur': 'Kuniyamuthur',
                }.get(branch.lower(), branch.capitalize())
                portal_url = get_portal_url(request)

                send_email_async(
                    subject='Welcome to IIE Pulse — Your Login Credentials',
                    message=f"""Dear {first_name},

Welcome to IIE Pulse! Your account has been created successfully.

??????????????????????????????
YOUR LOGIN CREDENTIALS
??????????????????????????????

  Portal URL  :  {portal_url}
  Username    :  {email}
  Password    :  {mobile_no}
  Staff ID    :  {staff_id}
  Role        :  {role.capitalize()}
  Branch      :  {branch_display}

??????????????????????????????

Please log in using the above credentials.
For security, we recommend changing your password after your first login.

If you have any trouble accessing your account, please contact your administrator.

Best regards,
IIE Pulse Team
Indra Institute of Education
{portal_url}
""",
                    recipient_list=[email],
                )
                logger.info("Welcome email queued for employee %s", email)
            except Exception as mail_err:
                logger.warning("Welcome email queueing failed for employee %s: %s", email, mail_err)
            # ---------------------------------------------------------

            return Response({
                'message': f"Employee '{first_name}' added successfully!",
                'staff_id': staff_id,
                'employee': EmployeeSerializer(employee).data,
                'next_staff_id': generate_staff_id(),
            }, status=201)

        except Exception as e:
            return Response({'error': str(e)}, status=400)


class EmployeeDetailView(generics.RetrieveUpdateDestroyAPIView):
    queryset = Employee.objects.all()
    serializer_class = EmployeeSerializer
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def partial_update(self, request, *args, **kwargs):
        employee = self.get_object()
        new_email = request.data.get('email', '').strip().lower()

        # If email changed, update the Django User too
        if new_email and new_email != employee.email:
            # Check not already taken by another user
            if User.objects.filter(username=new_email).exclude(pk=employee.user.pk).exists():
                return Response({'error': 'This email is already registered!'}, status=400)
            employee.user.username = new_email
            employee.user.email = new_email
            employee.user.save()

        return super().partial_update(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        employee = self.get_object()
        user = employee.user
        employee.delete()
        user.delete()
        return Response({'message': 'Employee deleted.'}, status=204)

class EmployeeProfileView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        try:
            emp = Employee.objects.get(user=request.user)
            return Response(EmployeeSerializer(emp).data)
        except Employee.DoesNotExist:
            return Response({'error': 'Not found'}, status=404)


# -- BATCHES ------------------------------------------------------------------

class BatchListCreateView(generics.ListAPIView):
    queryset = Batches.objects.all().order_by('-created_at')
    serializer_class = BatchSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        
        # Get faculty filter from query params
        faculty_id = self.request.query_params.get('faculty')
        if faculty_id:
            qs = qs.filter(faculty__id=faculty_id)
        
        # For non-admin users
        if not (user.is_superuser or user.is_staff):
            try:
                emp = Employee.objects.get(user=user)
                if emp.designation.lower() == 'counselor':
                    qs = qs.filter(branch=emp.branch)
                elif emp.designation.lower() in ['trainer', 'mentor']:
                    # If faculty filter not provided, filter by the logged-in staff
                    if not faculty_id:
                        qs = qs.filter(faculty=emp)
            except Employee.DoesNotExist:
                return qs.none()
        
        return qs

    def list(self, request, *args, **kwargs):
        """Override list to add student_count to each batch"""
        queryset = self.get_queryset()
        serializer = self.get_serializer(queryset, many=True)
        data = serializer.data
        
        # Add student_count to each batch
        for batch_data in data:
            try:
                batch = Batches.objects.get(id=batch_data['id'])
                student_count = Students.objects.filter(assigned_batch=batch).count()
                batch_data['student_count'] = student_count
            except Batches.DoesNotExist:
                batch_data['student_count'] = 0
        
        # Get total count for pagination if needed
        total_count = queryset.count()
        
        # Return in the same format as other list views
        return Response({
            'results': data,
            'count': total_count
        })
    
class BatchCreateView(APIView):
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def post(self, request):
        data = request.data
        batch_number = data.get('batch_number', '').strip()
        course_type = data.get('course_type', '').strip()
        course_name_id = data.get('course_name', '').strip()
        faculty_id = data.get('faculty', '').strip()
        start_date = data.get('start_date', '').strip()
        batch_timing = data.get('batch_timing', '').strip()
        branch = data.get('branch', '').strip()
        logsheet_file = request.FILES.get('logsheet_file')

        employee = None
        try:
            employee = Employee.objects.get(user=request.user)
        except Employee.DoesNotExist:
            pass

        if not batch_number and branch:
            batch_number = generate_batch_number(branch)

        missing = []
        if not batch_number: missing.append('Batch Number')
        if not course_type: missing.append('Course Type')
        if not course_name_id: missing.append('Course Name')
        if not faculty_id: missing.append('Faculty')
        if not start_date: missing.append('Start Date')
        if not batch_timing: missing.append('Batch Timing')
        if not branch: missing.append('Branch')
        if missing:
            return Response({'error': f"Missing: {', '.join(missing)}"}, status=400)

        if employee and employee.designation.lower() == 'counselor':
            if branch != employee.branch:
                return Response({'error': f'You can only create batches for your own branch ({employee.branch})!'}, status=403)
            try:
                selected_faculty = Employee.objects.get(id=faculty_id)
                if selected_faculty.branch != employee.branch:
                    return Response({'error': 'Selected faculty is not from your branch!'}, status=403)
            except Employee.DoesNotExist:
                return Response({'error': 'Selected faculty not found!'}, status=404)

        if Batches.objects.filter(batch_number=batch_number).exists():
            return Response({'error': f"Batch number '{batch_number}' already exists!"}, status=400)

        try:
            course = Courses.objects.get(id=course_name_id)
            faculty = Employee.objects.get(id=faculty_id)
        except Courses.DoesNotExist:
            return Response({'error': 'Selected course not found!'}, status=404)
        except Employee.DoesNotExist:
            return Response({'error': 'Selected faculty not found!'}, status=404)

        try:
            calculated_end_date = calculate_batch_end_date(course, start_date)
        except ValueError as exc:
            return Response({'error': str(exc)}, status=400)

        batch = Batches(
            batch_number=batch_number,
            course_type=course_type,
            course_name=course,
            faculty=faculty,
            start_date=start_date,
            end_date=calculated_end_date,
            batch_timing=batch_timing,
            branch=branch,
        )
        if logsheet_file:
            batch.course_logsheet = logsheet_file
        batch.save()
        if faculty.user:
            _queue_user_notification(
                faculty.user,
                request.user,
                'assignment',
                'New Batch Assigned',
                f"New Batch Assigned: {batch.batch_number} has been assigned to you.",
                True,
            )

        return Response({
            'message': f"Batch '{batch_number}' created successfully!",
            'batch': BatchSerializer(batch).data,
            'next_batch_number': generate_batch_number(branch),
        }, status=201)


class BatchDetailView(generics.RetrieveUpdateDestroyAPIView):
    queryset = Batches.objects.all()
    serializer_class = BatchSerializer
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def partial_update(self, request, *args, **kwargs):
        batch = self.get_object()
        data = request.data

        # -- Update fields manually ----------------------------------------
        if data.get('course_type'):
            batch.course_type = data.get('course_type')

        if data.get('course_name'):
            try:
                course = Courses.objects.get(id=data.get('course_name'))
                batch.course_name = course
            except Courses.DoesNotExist:
                return Response({'error': 'Course not found'}, status=404)

        if data.get('faculty'):
            try:
                faculty = Employee.objects.get(id=data.get('faculty'))
                batch.faculty = faculty
            except Employee.DoesNotExist:
                return Response({'error': 'Faculty not found'}, status=404)

        if data.get('start_date'):
            batch.start_date = data.get('start_date')

        if data.get('batch_timing'):
            batch.batch_timing = data.get('batch_timing')

        if data.get('branch'):
            batch.branch = data.get('branch')

        if data.get('batch_number'):
            batch.batch_number = data.get('batch_number')

        # -- Handle logsheet file ------------------------------------------
        if request.FILES.get('course_logsheet'):
            batch.course_logsheet = request.FILES.get('course_logsheet')

        try:
            batch.end_date = calculate_batch_end_date(batch.course_name, batch.start_date)
        except ValueError as exc:
            return Response({'error': str(exc)}, status=400)

        batch.save()

        return Response({
            'message': f"Batch '{batch.batch_number}' updated successfully!",
            'batch': BatchSerializer(batch).data,
        })

    def destroy(self, request, *args, **kwargs):
        batch = self.get_object()
        batch_number = batch.batch_number
        batch.delete()
        return Response({'message': f"Batch '{batch_number}' deleted successfully!"}, status=204)

# -- STUDENTS -----------------------------------------------------------------
class StudentListView(generics.ListAPIView):
    queryset = Students.objects.all().order_by('-created_at')
    serializer_class = StudentSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        qs = super().get_queryset()

        # -- Exclude ONLY FULLY COMPLETED students --------------------
        # Only exclude students who have a CompletedStudent record with completion_type='full'
        try:
            full_completed_ids = CompletedStudent.objects.filter(
                completion_type='full'
            ).values_list('original_student_id', flat=True)
        except DatabaseError as e:
            # Likely the production DB is missing the column; log and continue without excluding
            logger.exception("Database error while querying CompletedStudent.completion_type")
            full_completed_ids = []

        valid_student_ids = []
        valid_student_pks = []
        for cid in full_completed_ids:
            if not cid:
                continue
            cid_str = str(cid).strip()
            if cid_str.isdigit():
                valid_student_pks.append(int(cid_str))
            valid_student_ids.append(cid_str)

        if valid_student_ids or valid_student_pks:
            exclude_filter = Q()
            if valid_student_ids:
                exclude_filter |= Q(student_id__in=valid_student_ids)
            if valid_student_pks:
                exclude_filter |= Q(id__in=valid_student_pks)
            qs = qs.exclude(exclude_filter)
        # -------------------------------------------------------------

        assigned_staff = self.request.query_params.get('assigned_staff')
        if assigned_staff:
            try:
                assigned_staff_id = int(assigned_staff)
                qs = qs.filter(assigned_staff__id=assigned_staff_id)
            except (ValueError, TypeError):
                pass

        branch = self.request.query_params.get('branch')
        if branch:
            qs = qs.filter(branch__iexact=branch.strip())

        user = self.request.user
        if not (user.is_superuser or user.is_staff):
            try:
                emp = Employee.objects.get(user=user)
                if emp.designation.lower() == 'counselor':
                    qs = qs.filter(branch=emp.branch)
                elif emp.designation.lower() in ['trainer', 'mentor']:
                    qs = qs.filter(assigned_staff=emp)
            except Employee.DoesNotExist:
                return qs.none()

        return qs
    
    def list(self, request, *args, **kwargs):
        try:
            return super().list(request, *args, **kwargs)
        except Exception as e:
            logger.exception("Unhandled exception in StudentListView.list")
            return Response({
                'error': 'Student list retrieval failed',
                'details': str(e),
                'view': 'StudentListView.list'
            }, status=500)
    
class StudentCreateView(APIView):
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def post(self, request):
        data = request.data
        student_id_input = data.get('student_id', '').strip()
        first_name = data.get('first_name', '').strip()
        last_name = data.get('last_name', '').strip()
        email = data.get('email', '').strip().lower()
        mobile_no = data.get('mobile_no', '').strip()
        date_of_birth = data.get('date_of_birth', '').strip()
        city = data.get('city', '').strip()
        state = data.get('state', '').strip()
        qualification = data.get('qualification', '').strip()
        course_name = data.get('course', '').strip()
        gender = data.get('gender', '').strip()
        branch = data.get('branch', '').strip()
        photo = request.FILES.get('photo')

        employee = None
        try:
            employee = Employee.objects.get(user=request.user)
        except Employee.DoesNotExist:
            pass

        if employee and employee.designation.lower() == 'counselor':
            if branch.lower() != employee.branch.lower():
                return Response({'error': f'You can only add students to your own branch ({employee.branch})!'}, status=403)

        if not student_id_input and branch:
            student_id_input = generate_student_id(branch)

        if not all([student_id_input, first_name, email, mobile_no, date_of_birth, city, state, qualification, course_name, branch]):
            return Response({'error': 'Please fill in all required fields.'}, status=400)

        existing_user = User.objects.filter(username=email).first()
        if existing_user:
            linked_employee = Employee.objects.filter(user=existing_user).exists()
            linked_student = Students.objects.filter(user=existing_user).exists()
            if existing_user.is_staff or existing_user.is_superuser or linked_employee or linked_student:
                return Response({'error': 'This email is already registered!'}, status=400)
            existing_user.delete()
        if Students.objects.filter(email=email).exists():
            return Response({'error': 'Email already exists!'}, status=400)
        if Students.objects.filter(student_id=student_id_input).exists():
            return Response({'error': f"Student ID '{student_id_input}' already exists!"}, status=400)

        try:
            user = User.objects.create_user(
                username=email,
                password=mobile_no,
                email=email,
                first_name=first_name,
                last_name=last_name,
            )
            student = Students.objects.create(
                user=user,
                student_id=student_id_input,
                email=email,
                first_name=first_name,
                last_name=last_name,
                mobile_no=mobile_no,
                date_of_birth=date_of_birth,
                city=city,
                state=state,
                qualification=qualification,
                course=course_name,
                gender=gender,
                branch=branch,
                photo=photo if photo else None,
            )

            # -- Send welcome email with login credentials (async) -------------
            try:
                branch_display = {
                    '100ft': '100 Feet Road',
                    'hopes': 'Hopes',
                    'kuniyamuthur': 'Kuniyamuthur',
                }.get(branch.lower(), branch.capitalize())
                portal_url = get_portal_url(request)

                send_email_async(
                    subject='Welcome to IIE Pulse — Your Login Credentials',
                    message=f"""Dear {first_name},

Welcome to IIE Pulse! Your student account has been created successfully.

??????????????????????????????
YOUR LOGIN CREDENTIALS
??????????????????????????????

  Portal URL  :  {portal_url}
  Username    :  {email}
  Password    :  {mobile_no}
  Student ID  :  {student_id_input}
  Course      :  {course_name}
  Branch      :  {branch_display}

??????????????????????????????

Please log in using the above credentials.
For security, we recommend changing your password after your first login.

If you have any trouble accessing your account, please contact your counselor or administrator.

Best regards,
IIE Pulse Team
Indra Institute of Education
{portal_url}
""",
                    recipient_list=[email],
                )
                logger.info("Welcome email queued for student %s", email)
            except Exception as mail_err:
                logger.warning("Welcome email queueing failed for student %s: %s", email, mail_err)
            # ---------------------------------------------------------

            return Response({
                'message': f"Student '{first_name}' added successfully!",
                'student_id': student_id_input,
                'student': StudentSerializer(student).data,
                'next_student_id': generate_student_id(branch),
            }, status=201)

        except Exception as e:
            if 'user' in locals():
                user.delete()

            return Response({'error': str(e)}, status=400)


class StudentDetailView(generics.RetrieveUpdateDestroyAPIView):
    queryset = Students.objects.all()
    serializer_class = StudentSerializer
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def partial_update(self, request, *args, **kwargs):
        student = self.get_object()
        new_email = request.data.get('email', '').strip().lower()

        # If email changed, update the Django User too
        if new_email and new_email != student.email:
            if User.objects.filter(username=new_email).exclude(pk=student.user.pk).exists():
                return Response({'error': 'This email is already registered!'}, status=400)
            student.user.username = new_email
            student.user.email = new_email
            student.user.save()

        return super().partial_update(request, *args, **kwargs)

    def destroy(self, request, *args, **kwargs):
        student = self.get_object()
        user = student.user
        student.delete()
        if user:
            user.delete()
        return Response({'message': 'Student deleted.'}, status=204)

class StudentProfileView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        try:
            student = Students.objects.get(user=request.user)
            return Response(StudentSerializer(student).data)
        except Students.DoesNotExist:
            return Response({'error': 'Not found'}, status=404)


from .models import FeePayment, FeeTransaction

# -- When student is assigned to batch, create fee record ---------------------
# -- When student is assigned to batch, create fee record ---------------------
@api_view(['POST'])
@permission_classes([IsAuthenticated])
def assign_staff_to_student(request, student_id):
    try:
        student = Students.objects.get(student_id=student_id)
        staff_id = request.data.get('staff_id')
        batch_id = request.data.get('batch_id')

        if staff_id:
            student.assigned_staff = Employee.objects.get(id=staff_id)

        if batch_id:
            batch = Batches.objects.get(id=batch_id)
            student.assigned_batch = batch

            sessions = CourseSession.objects.filter(batch=batch).order_by("session_number")

            # If this is a new batch and sessions are not created yet, extract from logsheet
            if not sessions.exists():
                print(f"?? No sessions found for batch {batch.id}. Trying to extract from logsheet...")

                extracted_sessions = extract_sessions_from_logsheet(batch, debug=True)

                for item in extracted_sessions:
                    CourseSession.objects.update_or_create(
                        batch=batch,
                        session_number=item["session_number"],
                        defaults={
                            "title": item.get("title", ""),
                            "topics": item.get("topics", ""),
                            "session_enabled": True,
                        }
                    )

                sessions = CourseSession.objects.filter(batch=batch).order_by("session_number")
                print(f"? Sessions available for batch {batch.id}: {sessions.count()}")

            Student_Session_Progress.objects.filter(student=student).delete()
            StudentSessionStatus.objects.filter(student=student).delete()

            for session in sessions:
                existing_staff_completed = bool(session.staff_completed) or Student_Session_Progress.objects.filter(
                    session=session,
                    staff_completed=True,
                ).exists()
                staff_completed_at = (session.completed_date or timezone.now()) if existing_staff_completed else None
                Student_Session_Progress.objects.create(
                    student=student,
                    session=session,
                    completed=False,
                    staff_completed=existing_staff_completed,
                    staff_completed_at=staff_completed_at,
                    student_status="pending" if existing_staff_completed else "not_started"
                )
                StudentSessionStatus.objects.create(
                    student=student,
                    session=session,
                    staff_completed=existing_staff_completed,
                    staff_completed_at=staff_completed_at,
                    student_status="pending" if existing_staff_completed else "not_started"
                )

            # -- Create fee record -----------------------------------------
            course_fee = batch.course_name.fee if batch.course_name and batch.course_name.fee else 0
            if course_fee:
                FeePayment.objects.get_or_create(
                    student=student,
                    batch=batch,
                    defaults={
                        'total_fee': course_fee,
                        'amount_paid': 0,
                        'balance': course_fee,
                        'is_fully_paid': False,
                    }
                )
            # -------------------------------------------------------------

        student.save()
        recipients = []
        if student.assigned_staff and student.assigned_staff.user:
            recipients.append(student.assigned_staff.user)
        if student.assigned_batch and student.assigned_batch.faculty and student.assigned_batch.faculty.user:
            recipients.append(student.assigned_batch.faculty.user)
        for recipient in set(recipients):
            _queue_user_notification(
                recipient,
                request.user,
                'assignment',
                'New Student Assigned',
                f"New Student Assigned: {student.first_name} {student.last_name or ''} has been assigned to batch {student.assigned_batch.batch_number if student.assigned_batch else 'N/A'}.",
                True,
            )
        return Response({'message': 'Assigned successfully.', 'student': StudentSerializer(student).data})

    except (Students.DoesNotExist, Employee.DoesNotExist, Batches.DoesNotExist):
        return Response({'error': 'Not found'}, status=404)

@api_view(['DELETE'])
@permission_classes([IsAuthenticated])
def remove_staff_from_student(request, student_id):
    try:
        student = Students.objects.get(student_id=student_id)
        student.assigned_staff = None
        student.assigned_batch = None
        student.save()
        return Response({'message': 'Assignment removed.'})
    except Students.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)


# -- ATTENDANCE ----------------------------------------------------------------

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def AttendanceListView(request):
    """Get attendance records - supports ?student=me for students"""
    user = request.user
    student_id = request.GET.get('student')
    batch_id = request.GET.get('batch')
    date = request.GET.get('date')
    
    # Start with base queryset
    qs = StudentAttendance.objects.all().order_by('-date')
    
    # Handle 'me' parameter for student
    if student_id == 'me':
        try:
            student = Students.objects.get(user=user)
            qs = qs.filter(student=student)
        except Students.DoesNotExist:
            return Response([])
    elif student_id:
        qs = qs.filter(student__id=student_id)
    
    # Apply other filters
    if batch_id:
        qs = qs.filter(batch__id=batch_id)
    if date:
        qs = qs.filter(date=date)
    
    # If user is a student (not admin/staff), only show their records
    # If user is a student (not admin/staff), only show their records
    if not (user.is_superuser or user.is_staff):
        try:
            student = Students.objects.get(user=user)
            qs = qs.filter(student=student)
        except Students.DoesNotExist:
        # Employee/counselor — allow access, don't restrict
            pass
    
    serializer = AttendanceSerializer(qs, many=True)
    return Response({'results': serializer.data})


LONG_ABSENT_DAYS = 4


def _full_name(first_name='', last_name=''):
    return ' '.join(part for part in [first_name, last_name] if part).strip()


def _employee_name(employee):
    if not employee:
        return 'Unassigned Mentor'
    return _full_name(employee.first_name, employee.last_name) or employee.user.username


def _student_name(student):
    return _full_name(student.first_name, student.last_name) or student.student_id


def _branch_match_values(branch):
    value = (branch or '').strip()
    if not value:
        return []
    compact = re.sub(r'\s+', '', value).lower()
    known = {
        '100ft': ['100ft', '100ft Road', '100ft road', '100 feet', '100 feet road'],
        'hopes': ['hopes', 'Hopes', 'Hopes College', 'hopes college'],
        'kuniyamuthur': ['kuniyamuthur', 'Kuniyamuthur'],
    }
    for key, values in known.items():
        compact_values = {re.sub(r'\s+', '', item).lower() for item in values}
        if compact == re.sub(r'\s+', '', key).lower() or compact in compact_values:
            return values
    return [value]


def _has_four_day_absent_streak(student, attendance_date):
    records = list(
        StudentAttendance.objects.filter(student=student, date__lte=attendance_date)
        .only('date', 'status')
        .order_by('-date')[:LONG_ABSENT_DAYS]
    )
    if len(records) < LONG_ABSENT_DAYS:
        return False
    for index, record in enumerate(records):
        if record.date != attendance_date - timedelta(days=index):
            return False
        if str(record.status).lower() != 'absent':
            return False
    return True


def _queue_user_notification(to_user, from_user, notification_type, title, message, requires_action=False, dedupe_today=True):
    if not to_user:
        return 0
    existing = SessionNotification.objects.filter(
        to_user=to_user,
        notification_type=notification_type,
        message=message,
    )
    if title:
        existing = existing.filter(title=title)
    if dedupe_today:
        existing = existing.filter(created_at__date=timezone.localdate())
    if existing.exists():
        return 0
    SessionNotification.objects.create(
        from_user=from_user,
        to_user=to_user,
        user=to_user,
        notification_type=notification_type,
        title=title,
        message=message,
        requires_action=requires_action,
    )
    return 1


def _queue_batch_notification(to_user, from_user, notification_type, title, message):
    return _queue_user_notification(
        to_user,
        from_user,
        notification_type,
        title,
        message,
        False,
        False,
    )


def _mentor_name(staff):
    if not staff:
        return 'Unknown'
    return _employee_name(staff)


def _branch_counselors_for_batch(batch):
    branch_values = _branch_match_values(getattr(batch, 'branch', ''))
    counselor_qs = Employee.objects.filter(designation__iexact='counselor').select_related('user')
    if branch_values:
        counselor_qs = counselor_qs.filter(branch__in=branch_values)
    return counselor_qs.distinct()


def _batch_duration_exceeded_message(batch, staff):
    batch_name = batch.batch_number or str(batch)
    mentor_name = _mentor_name(staff)
    start_display = batch.start_date.strftime("%d/%m/%Y") if batch.start_date else '-'
    end_display = batch.end_date.strftime("%d/%m/%Y") if batch.end_date else '-'
    return (
        f'Batch "{batch_name}", handled by Mentor "{mentor_name}", was scheduled from '
        f'{start_display} to {end_display}. '
        'The batch duration has ended, but the batch is still in progress.'
    )


def _send_batch_duration_notifications(batch, staff, from_user, attendance_date):
    if not batch or not staff or not attendance_date:
        return 0

    end_date = getattr(batch, 'end_date', None)
    if not end_date:
        try:
            end_date = calculate_batch_end_date(batch.course_name, batch.start_date)
            batch.end_date = end_date
            batch.save(update_fields=['end_date'])
        except Exception:
            return 0

    batch_name = batch.batch_number or str(batch)
    mentor_name = _mentor_name(staff)
    date_key = attendance_date.strftime('%Y-%m-%d')
    sent_count = 0

    if end_date - timedelta(days=2) <= attendance_date <= end_date:
        message = f'Batch "{batch_name}" is nearing its scheduled end date. Please complete the remaining sessions soon.'
        sent_count += _queue_batch_notification(
            staff.user,
            from_user,
            'batch_ending_soon',
            f'Batch Ending Soon - {batch_name} - {date_key}',
            message,
        )

    if attendance_date > end_date:
        duration_message = _batch_duration_exceeded_message(batch, staff)
        sent_count += _queue_batch_notification(
            staff.user,
            from_user,
            'batch_duration_exceeded',
            f'Batch Duration Exceeded - {batch_name} - {date_key}',
            duration_message,
        )

        for admin_user in User.objects.filter(Q(is_staff=True) | Q(is_superuser=True)).distinct():
            sent_count += _queue_batch_notification(
                admin_user,
                from_user,
                'batch_duration_exceeded',
                f'Batch Duration Exceeded - {batch_name} - {date_key}',
                duration_message,
            )

        for counselor in _branch_counselors_for_batch(batch):
            sent_count += _queue_batch_notification(
                counselor.user,
                from_user,
                'batch_duration_exceeded',
                f'Batch Duration Exceeded - {batch_name} - {date_key}',
                duration_message,
            )

    return sent_count


def _queue_leave_alert(to_user, from_user, message):
    return _queue_user_notification(to_user, from_user, 'leave_alert', 'Leave Alert', message, True, False)


def _send_long_absence_notifications(student, batch, staff, from_user, attendance_date):
    if not _has_four_day_absent_streak(student, attendance_date):
        return 0

    student_display = _student_name(student)
    mentor = student.assigned_staff or getattr(student.assigned_batch, 'faculty', None) or getattr(batch, 'faculty', None) or staff
    mentor_display = _employee_name(mentor)
    start_date = attendance_date - timedelta(days=LONG_ABSENT_DAYS - 1)
    date_range = f"from {start_date.strftime('%d %b %Y')} to {attendance_date.strftime('%d %b %Y')}"
    sent_count = 0

    admin_message = (
        f"Leave Alert: Student {student_display}, under Mentor {mentor_display}, has been on continuous leave for "
        f"{LONG_ABSENT_DAYS} consecutive days ({date_range}). Please review the student's attendance."
    )
    counselor_message = (
        f"Leave Alert: Student {student_display}, under Mentor {mentor_display}, has been on continuous leave for "
        f"{LONG_ABSENT_DAYS} consecutive days ({date_range}). Please follow up with the student."
    )
    mentor_message = (
        f"Leave Alert: Your student {student_display} has been on continuous leave for {LONG_ABSENT_DAYS} "
        f"consecutive days ({date_range}). Please contact the student and update the status."
    )
    student_message = (
        f"Leave Alert: You have been on continuous leave for {LONG_ABSENT_DAYS} consecutive days. "
        f"Leave duration: {date_range}. Please contact your Mentor and update your leave status."
    )

    for admin_user in User.objects.filter(Q(is_staff=True) | Q(is_superuser=True)).distinct():
        sent_count += _queue_leave_alert(admin_user, from_user, admin_message)

    branch_values = _branch_match_values(student.branch or getattr(batch, 'branch', ''))
    counselor_qs = Employee.objects.filter(designation__iexact='counselor')
    if branch_values:
        counselor_qs = counselor_qs.filter(branch__in=branch_values)
    for counselor in counselor_qs.select_related('user').distinct():
        sent_count += _queue_leave_alert(counselor.user, from_user, counselor_message)

    if mentor:
        sent_count += _queue_leave_alert(mentor.user, from_user, mentor_message)
    if student.user:
        sent_count += _queue_leave_alert(student.user, from_user, student_message)

    return sent_count


def _sync_long_absence_notifications():
    student_ids = StudentAttendance.objects.filter(status__iexact='Absent').values_list('student_id', flat=True).distinct()
    created = 0
    for student_id in student_ids:
        latest_absent = (
            StudentAttendance.objects.filter(student_id=student_id, status__iexact='Absent')
            .select_related('student', 'batch', 'staff', 'staff__user')
            .order_by('-date')
            .first()
        )
        if not latest_absent:
            continue
        from_user = latest_absent.staff.user if latest_absent.staff_id else User.objects.filter(is_superuser=True).first()
        if not from_user:
            continue
        created += _send_long_absence_notifications(
            latest_absent.student,
            latest_absent.batch,
            latest_absent.staff,
            from_user,
            latest_absent.date,
        )
    return created


def _student_from_leave_alert_message(message):
    match = re.search(r'Leave Alert: Student\s+(.+?),\s+under Mentor', message or '')
    if not match:
        return None
    alert_name = re.sub(r'\s+', ' ', match.group(1)).strip().lower()
    for student in Students.objects.select_related('assigned_staff', 'assigned_batch', 'assigned_batch__faculty'):
        student_name = re.sub(r'\s+', ' ', _student_name(student)).strip().lower()
        if student_name == alert_name:
            return student
    return None


def _batch_from_duration_notification(notification):
    title_match = re.search(
        r'Batch (?:Ending Soon|Duration Exceeded) - (.+?) - \d{4}-\d{2}-\d{2}$',
        notification.title or '',
    )
    batch_number = title_match.group(1).strip() if title_match else ''
    if not batch_number:
        message_match = re.search(r'Batch\s+"([^"]+)"', notification.message or '')
        batch_number = message_match.group(1).strip() if message_match else ''
    if not batch_number:
        return None
    return Batches.objects.select_related('faculty').filter(batch_number=batch_number).first()


def _notification_action_url(notification, request_user):
    if notification.notification_type == 'batch_duration_exceeded':
        return ''

    if notification.notification_type in ('batch_ending_soon', 'batch_duration_exceeded'):
        batch = _batch_from_duration_notification(notification)
        mentor = getattr(batch, 'faculty', None) if batch else None
        batch_student = (
            Students.objects.filter(assigned_batch=batch).order_by('first_name', 'last_name', 'id').first()
            if batch else None
        )
        employee = Employee.objects.filter(user=request_user).first()
        if employee and (employee.designation or '').lower() in ('trainer', 'mentor'):
            return '/employee/batches'
        if employee and (employee.designation or '').lower() == 'counselor':
            if batch:
                params = {
                    'staff_id': mentor.id if mentor else '',
                    'batch_id': batch.id,
                    'student_id': batch_student.id if batch_student else '',
                    'tab': 'attendance',
                }
                query = '&'.join(
                    f"{key}={requests.utils.quote(str(value))}"
                    for key, value in params.items()
                    if value
                )
                return f"/counselor/students?{query}" if query else '/counselor/students'
            return '/counselor/students'
        if is_admin_user(request_user):
            if batch:
                params = {
                    'branch': batch.branch,
                    'staff_id': mentor.id if mentor else '',
                    'batch_id': batch.id,
                    'student_id': batch_student.id if batch_student else '',
                    'tab': 'attendance',
                }
                query = '&'.join(
                    f"{key}={requests.utils.quote(str(value))}"
                    for key, value in params.items()
                    if value
                )
                return f"/admin/attendance?{query}" if query else '/admin/attendance'
            return '/admin/attendance'
        return ''

    if notification.notification_type == 'leave_alert':
        student = _student_from_leave_alert_message(notification.message)
        if not student:
            return ''
        mentor = student.assigned_staff or getattr(student.assigned_batch, 'faculty', None)
        batch = student.assigned_batch
        employee = Employee.objects.filter(user=request_user).first()
        if is_admin_user(request_user):
            params = {
                'branch': student.branch or getattr(batch, 'branch', ''),
                'staff_id': mentor.id if mentor else '',
                'batch_id': batch.id if batch else '',
                'student_id': student.id,
                'tab': 'attendance',
            }
            query = '&'.join(
                f"{key}={requests.utils.quote(str(value))}"
                for key, value in params.items()
                if value
            )
            return f"/admin/attendance?{query}" if query else '/admin/attendance'
        if employee and (employee.designation or '').lower() == 'counselor':
            params = {
                'staff_id': mentor.id if mentor else '',
                'batch_id': batch.id if batch else '',
                'student_id': student.id,
                'tab': 'attendance',
            }
            query = '&'.join(
                f"{key}={requests.utils.quote(str(value))}"
                for key, value in params.items()
                if value
            )
            return f"/counselor/students?{query}" if query else '/counselor/students'
        return '/employee/attendance-history'
    return ''


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def mark_attendance(request):
    batch_id = request.data.get('batch_id')
    date = request.data.get('date')
    attendance_data = request.data.get('attendance', [])
    attendance_date = parse_date(str(date or ''))
    today = timezone.localdate()
    allowed_dates = {today, today - timedelta(days=1)}

    if not attendance_date:
        return Response({'error': 'Valid attendance date is required.'}, status=400)
    if attendance_date not in allowed_dates:
        return Response({'error': 'Attendance can be marked only for today or yesterday.'}, status=400)

    try:
        batch = Batches.objects.get(id=batch_id)
    except Batches.DoesNotExist:
        return Response({'error': 'Batch not found'}, status=404)

    try:
        staff = Employee.objects.get(user=request.user)
    except Employee.DoesNotExist:
        return Response({'error': 'Staff not found'}, status=404)

    created = []
    leave_alerts = 0
    batch_duration_alerts = 0
    for item in attendance_data:
        try:
            student = Students.objects.get(id=item['student_id'])
            status_value = item.get('status', 'Present')
            att, _ = StudentAttendance.objects.update_or_create(
                student=student,
                date=attendance_date,
                defaults={
                    'batch': batch,
                    'staff': staff,
                    'status': status_value,
                    'remarks': item.get('remarks', ''),
                }
            )
            created.append(AttendanceSerializer(att).data)
            if str(status_value).lower() == 'absent':
                leave_alerts += _send_long_absence_notifications(student, batch, staff, request.user, attendance_date)
        except Students.DoesNotExist:
            pass
    if created:
        batch_duration_alerts = _send_batch_duration_notifications(batch, staff, request.user, attendance_date)
    return Response({
        'message': f'{len(created)} records saved.',
        'records': created,
        'leave_alerts': leave_alerts,
        'batch_duration_alerts': batch_duration_alerts,
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def batch_attendance_records(request, batch_id):
    records = StudentAttendance.objects.filter(batch__id=batch_id).order_by('-date')
    return Response(AttendanceSerializer(records, many=True).data)


# -- STUDY MATERIALS -----------------------------------------------------------

def _get_employee_for_request(request):
    try:
        return Employee.objects.get(user=request.user)
    except Employee.DoesNotExist:
        return None


def _assign_material_to_batch(material, batch_id, assigned_by):
    if not batch_id:
        raise ValueError('Target batch is required.')

    try:
        batch = Batches.objects.get(id=batch_id)
    except (Batches.DoesNotExist, ValueError, TypeError):
        raise Batches.DoesNotExist

    assignment, created = StudyMaterialAssignment.objects.get_or_create(
        material=material,
        batch=batch,
        defaults={'assigned_by': assigned_by},
    )
    return assignment, created


class StudyMaterialListView(generics.ListAPIView):
    queryset = StudyMaterial.objects.all().select_related('batch', 'uploaded_by').prefetch_related('assignments__batch').order_by('-uploaded_at')
    serializer_class = StudyMaterialSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        try:
            emp = Employee.objects.get(user=user)
            if not (user.is_superuser or user.is_staff):
                return qs.filter(uploaded_by=emp)
            return qs
        except Employee.DoesNotExist:
            try:
                student = Students.objects.get(user=user)
                if student.assigned_batch:
                    return qs.filter(
                        Q(batch=student.assigned_batch) |
                        Q(assignments__batch=student.assigned_batch)
                    ).distinct()
                return qs.none()
            except Students.DoesNotExist:
                return qs


class StudyMaterialCreateView(generics.CreateAPIView):
    serializer_class = StudyMaterialSerializer
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser, JSONParser]

    def create(self, request, *args, **kwargs):
        if not request.data.get('batch'):
            return Response({'batch': ['Batch is required.']}, status=status.HTTP_400_BAD_REQUEST)

        serializer = self.get_serializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        emp = Employee.objects.get(user=self.request.user)
        material = serializer.save(uploaded_by=emp, is_library=False)
        try:
            _assign_material_to_batch(material, material.batch_id, emp)
        except Batches.DoesNotExist:
            pass
        headers = self.get_success_headers(serializer.data)
        return Response(
            StudyMaterialSerializer(material, context=self.get_serializer_context()).data,
            status=status.HTTP_201_CREATED,
            headers=headers
        )


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def material_library_list(request):
    emp = _get_employee_for_request(request)
    if not emp:
        return Response({'error': 'Only employees can access the material library.'}, status=status.HTTP_403_FORBIDDEN)

    qs = StudyMaterial.objects.filter(
        is_library=True,
        uploaded_by=emp,
    ).select_related('uploaded_by').prefetch_related('assignments__batch').order_by('-uploaded_at')

    search = request.query_params.get('search')
    branch = request.query_params.get('branch')
    batch = request.query_params.get('batch')

    if search:
        qs = qs.filter(title__icontains=search)
    if branch:
        qs = qs.filter(assignments__batch__branch=branch)
    if batch:
        qs = qs.filter(assignments__batch_id=batch)

    return Response(StudyMaterialSerializer(qs.distinct(), many=True, context={'request': request}).data)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser, JSONParser])
def material_library_upload(request):
    emp = _get_employee_for_request(request)
    if not emp:
        return Response({'error': 'Only employees can upload library materials.'}, status=status.HTTP_403_FORBIDDEN)
    if not request.data.get('file'):
        return Response({'file': ['File is required.']}, status=status.HTTP_400_BAD_REQUEST)

    serializer = StudyMaterialSerializer(data=request.data, context={'request': request})
    serializer.is_valid(raise_exception=True)
    material = serializer.save(uploaded_by=emp, is_library=True, batch=None)
    return Response(StudyMaterialSerializer(material, context={'request': request}).data, status=status.HTTP_201_CREATED)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def material_library_assign(request):
    emp = _get_employee_for_request(request)
    if not emp:
        return Response({'error': 'Only employees can assign library materials.'}, status=status.HTTP_403_FORBIDDEN)

    material_id = request.data.get('material') or request.data.get('material_id')
    batch_id = request.data.get('batch') or request.data.get('batch_id')

    if not material_id:
        return Response({'material': ['Library material is required.']}, status=status.HTTP_400_BAD_REQUEST)
    if not batch_id:
        return Response({'batch': ['Target batch is required.']}, status=status.HTTP_400_BAD_REQUEST)

    try:
        material = StudyMaterial.objects.get(id=material_id, is_library=True, uploaded_by=emp)
    except StudyMaterial.DoesNotExist:
        return Response({'error': 'Library material not found.'}, status=status.HTTP_404_NOT_FOUND)

    if not material.file:
        return Response({'error': 'Material file is missing.'}, status=status.HTTP_404_NOT_FOUND)

    try:
        file_path = material.file.path
    except (NotImplementedError, ValueError):
        file_path = None

    if file_path and not os.path.exists(file_path):
        return Response({'error': 'Material file is missing on server.'}, status=status.HTTP_404_NOT_FOUND)

    try:
        assignment, created = _assign_material_to_batch(material, batch_id, emp)
    except Batches.DoesNotExist:
        return Response({'error': 'Selected batch is invalid.'}, status=status.HTTP_400_BAD_REQUEST)
    except ValueError as exc:
        return Response({'error': str(exc)}, status=status.HTTP_400_BAD_REQUEST)

    return Response({
        'message': 'Material assigned.' if created else 'Material was already assigned to this batch.',
        'created': created,
        'assignment_id': assignment.id,
        'material': StudyMaterialSerializer(material, context={'request': request}).data,
    })


@api_view(['DELETE'])
@permission_classes([IsAuthenticated])
def material_library_delete(request, pk):
    emp = _get_employee_for_request(request)
    if not emp:
        return Response({'error': 'Only employees can delete library materials.'}, status=status.HTTP_403_FORBIDDEN)

    try:
        material = StudyMaterial.objects.get(id=pk, is_library=True, uploaded_by=emp)
    except StudyMaterial.DoesNotExist:
        return Response({'error': 'Library material not found.'}, status=status.HTTP_404_NOT_FOUND)

    if material.assignments.exists():
        return Response({'error': 'Cannot delete a library material that is assigned to batches.'}, status=status.HTTP_400_BAD_REQUEST)

    material.delete()
    return Response({'message': 'Deleted.'}, status=204)


@api_view(['PATCH', 'PUT'])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser, JSONParser])
def update_material(request, pk):
    emp = _get_employee_for_request(request)
    if not emp and not (request.user.is_superuser or request.user.is_staff):
        return Response({'error': 'Only employees can edit materials.'}, status=status.HTTP_403_FORBIDDEN)

    try:
        material = StudyMaterial.objects.get(id=pk)
    except StudyMaterial.DoesNotExist:
        return Response({'error': 'Material not found.'}, status=status.HTTP_404_NOT_FOUND)

    if emp and material.uploaded_by_id != emp.id and not (request.user.is_superuser or request.user.is_staff):
        return Response({'error': 'You can edit only your uploaded materials.'}, status=status.HTTP_403_FORBIDDEN)

    serializer = StudyMaterialSerializer(material, data=request.data, partial=True, context={'request': request})
    serializer.is_valid(raise_exception=True)
    material = serializer.save()

    batch_id = request.data.get('batch') or request.data.get('batch_id')
    if batch_id:
        try:
            _assign_material_to_batch(material, batch_id, emp or material.uploaded_by)
            if not material.is_library:
                material.assignments.exclude(batch_id=batch_id).delete()
        except Batches.DoesNotExist:
            return Response({'error': 'Selected batch is invalid.'}, status=status.HTTP_400_BAD_REQUEST)
        except ValueError as exc:
            return Response({'error': str(exc)}, status=status.HTTP_400_BAD_REQUEST)

    return Response(StudyMaterialSerializer(material, context={'request': request}).data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def download_material(request, pk):
    try:
        material = StudyMaterial.objects.get(id=pk)
    except StudyMaterial.DoesNotExist:
        return Response({'error': 'Material not found.'}, status=404)

    # Use persistent storage if uploaded files must survive deploys, restarts,
    # or rebuilds.
    if not material.file:
        return Response(
            {'error': 'File missing on server. Re-upload required or persistent storage required.'},
            status=404,
        )

    try:
        file_path = material.file.path
    except (NotImplementedError, ValueError):
        return Response(
            {'error': 'File missing on server. Re-upload required or persistent storage required.'},
            status=404,
        )

    if not os.path.exists(file_path):
        return Response(
            {'error': 'File missing on server. Re-upload required or persistent storage required.'},
            status=404,
        )

    filename = os.path.basename(material.file.name)
    content_type = mimetypes.guess_type(filename)[0] or 'application/octet-stream'
    file_handle = open(file_path, 'rb')
    response = FileResponse(
        file_handle,
        as_attachment=True,
        filename=filename,
        content_type=content_type,
    )
    response['Content-Length'] = str(os.path.getsize(file_path))
    return response


@api_view(['DELETE'])
@permission_classes([IsAuthenticated])
def delete_material(request, pk):
    try:
        StudyMaterial.objects.get(id=pk).delete()
        return Response({'message': 'Deleted.'}, status=204)
    except StudyMaterial.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)


# -- TESTS ---------------------------------------------------------------------

class TestListCreateView(generics.ListCreateAPIView):
    serializer_class = TestSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        
        # Superuser/Admin can see all tests
        if user.is_superuser or user.is_staff:
            return QuizTest.objects.all().order_by('-created_at')
        
        # Staff/Trainer/Mentor - Only see tests they created
        try:
            emp = Employee.objects.get(user=user)
            # Filter tests where created_by is the logged-in employee
            return QuizTest.objects.filter(created_by=emp).order_by('-created_at')
        except Employee.DoesNotExist:
            return QuizTest.objects.none()
    
    def perform_create(self, serializer):
        try:
            emp = Employee.objects.get(user=self.request.user)
        except Employee.DoesNotExist:
            raise PermissionDenied("Employee profile not found for this account.")
        serializer.save(created_by=emp)

class TestDetailView(generics.RetrieveUpdateDestroyAPIView):
    queryset = QuizTest.objects.all()
    serializer_class = TestSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        
        if user.is_superuser or user.is_staff:
            return QuizTest.objects.all()
        
        try:
            emp = Employee.objects.get(user=user)
            # Only allow access to tests created by this employee
            return QuizTest.objects.filter(created_by=emp)
        except Employee.DoesNotExist:
            return QuizTest.objects.none()

@api_view(['POST'])
@permission_classes([IsAuthenticated])
def add_question(request, test_id):
    try:
        test = QuizTest.objects.get(id=test_id)
    except QuizTest.DoesNotExist:
        return Response({'error': 'Test not found'}, status=404)

    data = request.data.copy()
    data['test'] = test.id

    required_fields = ['question_text', 'option1', 'option2', 'correct_answer']
    missing_fields = [field for field in required_fields if not data.get(field)]
    if missing_fields:
        return Response(
            {'error': 'Missing required fields', 'missing_fields': missing_fields},
            status=400
        )

    serializer = QuestionSerializer(data=data)
    if serializer.is_valid():
        serializer.save(test=test)
        return Response(serializer.data, status=201)
    return Response({'error': 'Validation failed', 'details': serializer.errors}, status=400)


class AssignedTestListView(generics.ListCreateAPIView):
    queryset = AssignedTest.objects.all()
    serializer_class = AssignedTestSerializer
    permission_classes = [IsAuthenticated]


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def student_tests(request):
    try:
        student = Students.objects.get(user=request.user)
        if student.assigned_batch:
            assigned = AssignedTest.objects.filter(batch=student.assigned_batch)
            return Response(AssignedTestSerializer(assigned, many=True).data)
    except Students.DoesNotExist:
        pass
    return Response([])


# -- LEAVE ---------------------------------------------------------------------

class StaffLeaveListView(generics.ListCreateAPIView):
    queryset = StaffLeaveRequest.objects.all().order_by('-applied_at')
    serializer_class = StaffLeaveRequestSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        if not (user.is_superuser or user.is_staff):
            try:
                emp = Employee.objects.get(user=user)
                return qs.filter(staff=emp)
            except Employee.DoesNotExist:
                return qs.none()
        return qs

    def perform_create(self, serializer):
        emp = Employee.objects.get(user=self.request.user)
        serializer.save(staff=emp)


@api_view(['PATCH'])
@permission_classes([IsAuthenticated])
def process_staff_leave(request, pk):
    try:
        leave = StaffLeaveRequest.objects.get(id=pk)
        leave.status = request.data.get('status', leave.status)
        leave.save()
        return Response(StaffLeaveRequestSerializer(leave).data)
    except StaffLeaveRequest.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)



from datetime import datetime
from rest_framework import generics, status
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from .models import StudentLeaveApplication, Students
from .serializers import StudentLeaveApplicationSerializer

class StudentLeaveListView(generics.ListCreateAPIView):
    serializer_class = StudentLeaveApplicationSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        try:
            student = Students.objects.get(user=user)
            return StudentLeaveApplication.objects.filter(student=student).order_by('-applied_at')
        except Students.DoesNotExist:
            return StudentLeaveApplication.objects.none()

    def get(self, request, *args, **kwargs):
        try:
            queryset = self.get_queryset()
            serializer = self.get_serializer(queryset, many=True)
            return Response({'results': serializer.data})
        except Exception as e:
            print(f"Error in get leaves: {e}")
            return Response({'results': []})

    def post(self, request, *args, **kwargs):
        try:
            print("=" * 50)
            print("POST request received to /student-leave/")
            print("Request data:", request.data)
            
            # Get student
            student = Students.objects.get(user=request.user)
            print(f"Student found: {student.first_name}")
            
            # Check assigned staff
            if not student.assigned_staff:
                print("No assigned staff found!")
                return Response({
                    'error': 'No staff assigned to you. Please contact administration.'
                }, status=status.HTTP_400_BAD_REQUEST)
            
            print(f"Assigned staff: {student.assigned_staff.first_name}")
            
            # Get form data
            start_date = request.data.get('start_date')
            end_date = request.data.get('end_date')
            leave_type = request.data.get('leave_type')
            reason = request.data.get('reason')
            contact_info = request.data.get('contact_info', '')
            
            # Validate
            if not start_date:
                return Response({'error': 'Start date is required'}, status=status.HTTP_400_BAD_REQUEST)
            if not end_date:
                return Response({'error': 'End date is required'}, status=status.HTTP_400_BAD_REQUEST)
            if not reason:
                return Response({'error': 'Reason is required'}, status=status.HTTP_400_BAD_REQUEST)
            
            # Calculate days
            start = datetime.strptime(start_date, '%Y-%m-%d').date()
            end = datetime.strptime(end_date, '%Y-%m-%d').date()
            number_of_days = (end - start).days + 1
            
            print(f"Leave dates: {start_date} to {end_date} = {number_of_days} days")
            
            # Create leave application
            leave = StudentLeaveApplication.objects.create(
                student=student,
                assigned_staff=student.assigned_staff,
                start_date=start_date,
                end_date=end_date,
                number_of_days=number_of_days,
                leave_type=leave_type,
                reason=reason,
                contact_info=contact_info,
                status='pending'
            )
            
            print(f"Leave application created successfully! ID: {leave.id}")
            _queue_user_notification(
                student.assigned_staff.user,
                request.user,
                'leave_application',
                'Leave Application',
                f"Leave Application: {student.first_name} {student.last_name or ''} requested {number_of_days} day(s) leave from {start_date} to {end_date}.",
                True,
            )
            
            # Return response
            serializer = self.get_serializer(leave)
            return Response(serializer.data, status=status.HTTP_201_CREATED)
            
        except Students.DoesNotExist:
            print("Student not found for this user")
            return Response({
                'error': 'Student profile not found'
            }, status=status.HTTP_404_NOT_FOUND)
        except Exception as e:
            print(f"Unexpected error: {str(e)}")
            import traceback
            traceback.print_exc()
            return Response({
                'error': str(e)
            }, status=status.HTTP_400_BAD_REQUEST)
        

@api_view(['PATCH'])
@permission_classes([IsAuthenticated])
def process_student_leave(request, pk):
    try:
        leave = StudentLeaveApplication.objects.get(id=pk)
        leave.status = request.data.get('status', leave.status)
        leave.save()
        return Response(StudentLeaveApplicationSerializer(leave).data)
    except StudentLeaveApplication.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)


# -- SUPPORT -------------------------------------------------------------------

class StaffSupportListView(generics.ListCreateAPIView):
    queryset = SupportRequest.objects.all().order_by('-created_at')
    serializer_class = SupportRequestSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        if not (user.is_superuser or user.is_staff):
            return qs.filter(staff=user)
        return qs

    def perform_create(self, serializer):
        serializer.save(staff=self.request.user)


class StudentSupportListView(generics.ListCreateAPIView):
    queryset = StudentSupportRequest.objects.all().order_by('-created_at')
    serializer_class = StudentSupportRequestSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        if not (user.is_superuser or user.is_staff):
            return qs.filter(student=user)
        return qs

    def perform_create(self, serializer):
        support = serializer.save(student=self.request.user)
        student = Students.objects.filter(user=self.request.user).select_related('assigned_staff__user').first()
        if student and student.assigned_staff and student.assigned_staff.user:
            _queue_user_notification(
                student.assigned_staff.user,
                self.request.user,
                'support',
                'Student Support Request',
                f"Support: {student.first_name} {student.last_name or ''} submitted a support request.",
                True,
            )
        return support


class CounselorSupportListView(generics.ListCreateAPIView):
    queryset = CounselorSupportRequest.objects.all().order_by('-created_at')
    serializer_class = CounselorSupportRequestSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        qs = super().get_queryset()
        if not (self.request.user.is_superuser or self.request.user.is_staff):
            return qs.filter(counselor=self.request.user)
        return qs

    def perform_create(self, serializer):
        serializer.save(counselor=self.request.user)


# -- ANNOUNCEMENTS -------------------------------------------------------------


class AnnouncementListView(generics.ListAPIView):
    serializer_class = AnnouncementSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        if user.is_superuser or user.is_staff:
            return Announcement.objects.all().order_by('-created_at')
        try:
            emp = Employee.objects.get(user=user)
            if emp.designation == 'counselor':
                return Announcement.objects.filter(
                    Q(recipient_type='all') | Q(recipient_type='staff') | Q(recipient_type='counselors'),
                    is_published=True,
                    created_by__is_staff=True,   # only admin-created
                ).order_by('-created_at')
            return Announcement.objects.filter(
                Q(recipient_type='all') | Q(recipient_type='staff') | Q(recipient_type='mentors'),
                is_published=True,
                created_by__is_staff=True,   # only admin-created
            ).order_by('-created_at')
        except Employee.DoesNotExist:
            return Announcement.objects.filter(
                Q(recipient_type='all') | Q(recipient_type='students'),
                is_published=True,
                created_by__is_staff=True,   # only admin-created
            ).order_by('-created_at')


class AnnouncementCreateView(generics.CreateAPIView):
    serializer_class = AnnouncementSerializer
    permission_classes = [IsAuthenticated]

    def perform_create(self, serializer):
        announcement = serializer.save(created_by=self.request.user)
        recipient_type = (announcement.recipient_type or '').lower()
        notify_mentors = recipient_type in ['all', 'staff', 'mentors']
        notify_counselors = recipient_type in ['all', 'staff', 'counselors']
        employee_filter = Q()
        if notify_mentors:
            employee_filter |= Q(designation__iexact='mentor') | Q(designation__iexact='trainer')
        if notify_counselors:
            employee_filter |= Q(designation__iexact='counselor')
        if employee_filter:
            for employee in Employee.objects.filter(employee_filter).select_related('user').distinct():
                _queue_user_notification(
                    employee.user,
                    self.request.user,
                    'announcement',
                    'Admin Announcement',
                    f"Admin Announcement: {announcement.title}",
                    False,
                )
        return announcement


def _admin_announcement_visible_to_student(announcement, student):
    if announcement.recipient_type == 'all':
        return True
    if announcement.recipient_type == 'students':
        selected = announcement.specific_students.all()
        return not selected.exists() or selected.filter(id=student.id).exists()
    if announcement.recipient_type == 'specific':
        return announcement.specific_students.filter(id=student.id).exists()
    return False


def _serialize_mobile_announcement(item, source, source_label, audience_label=None):
    data = dict(item)
    data['id'] = f"{source}-{data.get('id')}"
    data['source'] = source
    data['source_label'] = source_label
    if audience_label:
        data['audience_label'] = audience_label
    return data


def _counselor_announcement_visible_to_student(announcement, student):
    if announcement.recipient_type == 'specific_batch':
        return bool(student.assigned_batch_id and announcement.specific_batch_id == student.assigned_batch_id)

    if announcement.recipient_type == 'specific_student':
        return announcement.specific_students.filter(id=student.id).exists()

    if announcement.recipient_type not in ['all', 'students']:
        return False

    student_branch = (student.branch or '').strip().lower()
    announcement_branch = (announcement.branch or '').strip().lower()
    if announcement_branch and announcement_branch == student_branch:
        return True

    creator_emp = Employee.objects.filter(user=announcement.created_by).first()
    creator_branch = (getattr(creator_emp, 'branch', '') or '').strip().lower()
    return bool(creator_branch and creator_branch == student_branch)


def _trainer_announcement_visible_to_student(announcement, student):
    creator_emp = Employee.objects.filter(user=announcement.created_by).first()
    if not creator_emp:
        return False
    if (creator_emp.designation or '').strip().lower() not in ['trainer', 'mentor']:
        return False

    if announcement.recipient_type == 'specific_batch':
        return bool(
            student.assigned_batch_id
            and announcement.specific_batch_id == student.assigned_batch_id
            and (
                student.assigned_staff_id == creator_emp.id
                or student.assigned_batch.faculty_id == creator_emp.id
            )
        )

    if announcement.recipient_type == 'specific_student':
        return announcement.specific_students.filter(id=student.id).exists() and (
            student.assigned_staff_id == creator_emp.id
            or (
                student.assigned_batch_id
                and student.assigned_batch.faculty_id == creator_emp.id
            )
        )

    return False


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def student_announcements(request):
    try:
        student = Students.objects.select_related(
            'assigned_batch',
            'assigned_staff',
            'assigned_staff__user',
        ).get(user=request.user)
    except Students.DoesNotExist:
        return Response({'error': 'Student profile not found'}, status=404)

    requested_sources = {
        source.strip().lower()
        for source in (request.query_params.get('sources') or '').split(',')
        if source.strip()
    }
    include_all_sources = not requested_sources

    admin_items = []
    if include_all_sources or 'admin' in requested_sources:
        admin_qs = Announcement.objects.prefetch_related('specific_students').filter(
            is_published=True,
            created_by__is_staff=True,
            recipient_type__in=['all', 'students', 'specific'],
        ).order_by('-created_at')
        for announcement in admin_qs:
            if _admin_announcement_visible_to_student(announcement, student):
                is_selected = announcement.specific_students.filter(id=student.id).exists()
                admin_items.append(_serialize_mobile_announcement(
                    AnnouncementSerializer(announcement).data,
                    'admin',
                    'Admin',
                    'Selected Student' if is_selected else None,
                ))

    counselor_items = []
    if include_all_sources or 'counselor' in requested_sources or 'counsellor' in requested_sources:
        counselor_qs = CounselorAnnouncement.objects.select_related(
            'specific_batch',
            'created_by',
            'created_by__employee',
        ).prefetch_related('specific_students').filter(
            is_published=True,
            created_by__employee__designation__iexact='counselor',
        ).distinct().order_by('-created_at')
        counselor_items = [
            _serialize_mobile_announcement(
                CounselorAnnouncementSerializer(announcement).data,
                'counselor',
                'Counselor',
                announcement.specific_batch.batch_number
                if announcement.recipient_type == 'specific_batch' and announcement.specific_batch
                else None,
            )
            for announcement in counselor_qs
            if _counselor_announcement_visible_to_student(announcement, student)
        ]

    trainer_items = []
    if include_all_sources or 'trainer' in requested_sources or 'mentor' in requested_sources:
        trainer_user_ids = Employee.objects.filter(
            Q(designation__iexact='trainer') | Q(designation__iexact='mentor')
        ).values_list('user_id', flat=True)
        trainer_qs = CounselorAnnouncement.objects.select_related(
            'specific_batch',
            'specific_batch__faculty',
            'created_by',
        ).prefetch_related('specific_students').filter(
            is_published=True,
            created_by_id__in=trainer_user_ids,
            recipient_type__in=['specific_batch', 'specific_student'],
        ).distinct().order_by('-created_at')
        trainer_items = [
            _serialize_mobile_announcement(
                CounselorAnnouncementSerializer(announcement).data,
                'trainer',
                'Trainer',
                announcement.specific_batch.batch_number
                if announcement.recipient_type == 'specific_batch' and announcement.specific_batch
                else 'Selected Student',
            )
            for announcement in trainer_qs
            if _trainer_announcement_visible_to_student(announcement, student)
        ]

    results = sorted(
        admin_items + counselor_items + trainer_items,
        key=lambda item: item.get('created_at') or '',
        reverse=True,
    )
    return Response({
        'results': results,
        'admin': admin_items,
        'counselor': counselor_items,
        'trainer': trainer_items,
        'count': len(results),
    })


@api_view(['PATCH'])
@permission_classes([IsAuthenticated])
def toggle_announcement(request, pk):
    try:
        ann = Announcement.objects.get(id=pk)
        ann.is_published = not ann.is_published
        ann.save()
        return Response({'is_published': ann.is_published})
    except Announcement.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)


@api_view(['DELETE'])
@permission_classes([IsAuthenticated])
def delete_announcement(request, pk):
    try:
        Announcement.objects.get(id=pk).delete()
        return Response(status=204)
    except Announcement.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)


# -- COUNSELOR LEAVE -----------------------------------------------------------

class CounselorLeaveListView(generics.ListCreateAPIView):
    queryset = CounselorLeaveRequest.objects.all().order_by('-applied_at')
    serializer_class = CounselorLeaveRequestSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        qs = super().get_queryset()
        user = self.request.user
        if not (user.is_superuser or user.is_staff):
            try:
                emp = Employee.objects.get(user=user)
                return qs.filter(counselor=emp)
            except Employee.DoesNotExist:
                return qs.none()
        return qs

    def perform_create(self, serializer):
        emp = Employee.objects.get(user=self.request.user)
        serializer.save(counselor=emp)


@api_view(['PATCH'])
@permission_classes([IsAuthenticated])
def process_counselor_leave(request, pk):
    try:
        leave = CounselorLeaveRequest.objects.get(id=pk)
        leave.status = request.data.get('status', leave.status)
        leave.save()
        return Response(CounselorLeaveRequestSerializer(leave).data)
    except CounselorLeaveRequest.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)


# -- QUIZ ----------------------------------------------------------------------

class QuizListView(generics.ListAPIView):
    serializer_class = QuizSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        
        # Admin can see all quizzes
        if user.is_superuser or user.is_staff:
            return Quiz.objects.all().order_by('-created_at')
        
        try:
            emp = Employee.objects.get(user=user)
            # Staff/Trainer/Mentor - Only see quizzes they uploaded
            return Quiz.objects.filter(created_by=emp).order_by('-created_at')
            
        except Employee.DoesNotExist:
            try:
                student = Students.objects.get(user=user)
                if student.assigned_batch:
                    # Students see published quizzes from their batch
                    return Quiz.objects.filter(
                        batch=student.assigned_batch, 
                        is_published=True
                    ).order_by('-created_at')
                return Quiz.objects.none()
            except Students.DoesNotExist:
                return Quiz.objects.none()


@api_view(['POST'])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser, JSONParser])
def upload_quiz(request):
    try:
        emp = Employee.objects.get(user=request.user)
    except Employee.DoesNotExist:
        if request.user.is_superuser or request.user.is_staff:
            emp = None
        else:
            return Response({'error': 'Employee not found'}, status=404)

    raw_batch_ids = []
    if hasattr(request.data, 'getlist'):
        raw_batch_ids.extend(request.data.getlist('batch_ids'))
        raw_batch_ids.extend(request.data.getlist('batches'))
    for key in ('batch_ids', 'batches', 'batch', 'batch_id'):
        value = request.data.get(key)
        if isinstance(value, str):
            raw_batch_ids.extend([part.strip() for part in value.split(',') if part.strip()])
        elif value:
            raw_batch_ids.append(str(value).strip())
    raw_batch_ids = [value for idx, value in enumerate(raw_batch_ids) if value and value not in raw_batch_ids[:idx]]

    target_batches = []
    is_practice_quiz = any(value in ('practice', 'all', 'public', '0') for value in raw_batch_ids)
    if is_practice_quiz:
        target_batches = [None]
    elif not raw_batch_ids:
        target_batches = [None]
    else:
        try:
            target_batches = list(Batches.objects.filter(id__in=raw_batch_ids))
        except ValueError:
            return Response({'error': 'Invalid batch selected'}, status=400)
        if len(target_batches) != len(raw_batch_ids):
            return Response({'error': 'One or more selected batches were not found'}, status=404)

    file = request.FILES.get('source_file')
    if not file:
        return Response({'error': 'Quiz file is required'}, status=400)

    def find_correct_option(correct_text, opt_a, opt_b, opt_c, opt_d):
        correct_text_lower = correct_text.lower().strip()
        if opt_a and opt_a.lower().strip() == correct_text_lower:
            return 'A'
        if opt_b and opt_b.lower().strip() == correct_text_lower:
            return 'B'
        if opt_c and opt_c.lower().strip() == correct_text_lower:
            return 'C'
        if opt_d and opt_d.lower().strip() == correct_text_lower:
            return 'D'
        if correct_text_lower in ('a', '1', 'option 1', 'option_1'):
            return 'A'
        if correct_text_lower in ('b', '2', 'option 2', 'option_2'):
            return 'B'
        if correct_text_lower in ('c', '3', 'option 3', 'option_3'):
            return 'C'
        if correct_text_lower in ('d', '4', 'option 4', 'option_4'):
            return 'D'
        return 'A'

    try:
        file.seek(0)
        if file.name.lower().endswith(('.xlsx', '.xls')):
            import openpyxl
            workbook = openpyxl.load_workbook(file, data_only=True)
            sheet = workbook.active
            headers = [str(cell.value or '').strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            reader = [
                {headers[idx]: cell.value for idx, cell in enumerate(row) if idx < len(headers)}
                for row in sheet.iter_rows(min_row=2)
            ]
        else:
            content = file.read().decode('utf-8-sig')
            reader = csv.DictReader(io.StringIO(content))

        questions = []
        for i, row in enumerate(reader, 1):
            question_text = str(row.get('Question') or row.get('question') or '').strip()
            option_a = str(row.get('Option_1') or row.get('Option 1') or row.get('option_a') or row.get('A') or '').strip()
            option_b = str(row.get('Option_2') or row.get('Option 2') or row.get('option_b') or row.get('B') or '').strip()
            option_c = str(row.get('Option_3') or row.get('Option 3') or row.get('option_c') or row.get('C') or '').strip()
            option_d = str(row.get('Option_4') or row.get('Option 4') or row.get('option_d') or row.get('D') or '').strip()
            correct_answer_text = str(row.get('Correct Answer') or row.get('correct_answer') or '').strip()
            if not question_text or not option_a or not option_b:
                continue
            questions.append({
                'question_number': len(questions) + 1,
                'question_text': question_text,
                'option_a': option_a,
                'option_b': option_b,
                'option_c': option_c,
                'option_d': option_d,
                'correct_answer': find_correct_option(correct_answer_text, option_a, option_b, option_c, option_d),
                'marks': int(row.get('marks', 1) or 1),
            })
    except Exception as e:
        logger.exception("Error parsing quiz upload")
        return Response({'error': f'Could not parse quiz file: {e}'}, status=400)

    if not questions:
        return Response({'error': 'No valid questions found in the uploaded file.'}, status=400)

    created_quizzes = []
    with transaction.atomic():
        for batch in target_batches:
            file.seek(0)
            quiz = Quiz.objects.create(
                title=request.data.get('title', 'Quiz'),
                description=request.data.get('description', ''),
                batch=batch,
                created_by=emp,
                duration_minutes=int(request.data.get('duration_minutes', 30)),
                passing_marks=int(request.data.get('passing_marks', 35)),
                difficulty=request.data.get('difficulty', 'medium') or 'medium',
                source_file=file,
                total_questions=len(questions),
                total_marks=sum(item['marks'] for item in questions),
                is_published=bool(raw_batch_ids),
                publish_date=timezone.now() if raw_batch_ids else None,
            )
            QuizQuestion.objects.bulk_create([
                QuizQuestion(quiz=quiz, **question)
                for question in questions
            ])
            created_quizzes.append(quiz)

    return Response({
        'message': f'Quiz created for {len(created_quizzes)} batch(es)' if raw_batch_ids else 'Quiz uploaded. Assign it to batches from Manage Quizzes.',
        'quiz_ids': [quiz.id for quiz in created_quizzes],
        'quiz_id': created_quizzes[0].id,
        'questions': len(questions),
    }, status=201)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def assign_quiz_to_batches(request, quiz_id):
    try:
        emp = Employee.objects.get(user=request.user)
    except Employee.DoesNotExist:
        if request.user.is_superuser or request.user.is_staff:
            emp = None
        else:
            return Response({'error': 'Employee not found'}, status=404)

    try:
        quiz = Quiz.objects.prefetch_related('questions').get(id=quiz_id)
    except Quiz.DoesNotExist:
        return Response({'error': 'Quiz not found'}, status=404)

    if emp and quiz.created_by_id != emp.id:
        return Response({'error': 'You can assign only your uploaded quizzes.'}, status=403)

    raw_batch_ids = []
    if hasattr(request.data, 'getlist'):
        raw_batch_ids.extend(request.data.getlist('batch_ids'))
        raw_batch_ids.extend(request.data.getlist('batches'))
    for key in ('batch_ids', 'batches', 'batch', 'batch_id'):
        value = request.data.get(key)
        if isinstance(value, list):
            raw_batch_ids.extend([str(item).strip() for item in value if str(item).strip()])
        elif isinstance(value, str):
            raw_batch_ids.extend([part.strip() for part in value.split(',') if part.strip()])
        elif value:
            raw_batch_ids.append(str(value).strip())
    raw_batch_ids = [value for idx, value in enumerate(raw_batch_ids) if value and value not in raw_batch_ids[:idx]]

    if not raw_batch_ids:
        return Response({'error': 'Select at least one batch'}, status=400)

    try:
        batches = list(Batches.objects.filter(id__in=raw_batch_ids))
    except ValueError:
        return Response({'error': 'Invalid batch selected'}, status=400)
    if len(batches) != len(raw_batch_ids):
        return Response({'error': 'One or more selected batches were not found'}, status=404)

    questions = list(quiz.questions.all().order_by('question_number'))
    if not questions:
        return Response({'error': 'Cannot assign a quiz without questions.'}, status=400)

    assigned_quizzes = []
    with transaction.atomic():
        reusable_quiz = quiz if quiz.batch_id is None else None
        for batch in batches:
            existing = Quiz.objects.filter(
                title=quiz.title,
                created_by=quiz.created_by,
                batch=batch,
            ).first()
            if existing:
                existing.is_published = True
                existing.publish_date = existing.publish_date or timezone.now()
                existing.save(update_fields=['is_published', 'publish_date', 'updated_at'])
                assigned_quizzes.append(existing)
                continue

            if reusable_quiz is not None:
                target_quiz = reusable_quiz
                target_quiz.batch = batch
                target_quiz.is_published = True
                target_quiz.publish_date = timezone.now()
                target_quiz.save(update_fields=['batch', 'is_published', 'publish_date', 'updated_at'])
                reusable_quiz = None
            else:
                target_quiz = Quiz.objects.create(
                    title=quiz.title,
                    description=quiz.description,
                    batch=batch,
                    created_by=quiz.created_by,
                    source_file=quiz.source_file,
                    total_questions=quiz.total_questions,
                    total_marks=quiz.total_marks,
                    passing_marks=quiz.passing_marks,
                    duration_minutes=quiz.duration_minutes,
                    difficulty=quiz.difficulty,
                    category=quiz.category,
                    start_date=quiz.start_date,
                    end_date=quiz.end_date,
                    shuffle_questions=quiz.shuffle_questions,
                    shuffle_options=quiz.shuffle_options,
                    number_of_questions=quiz.number_of_questions,
                    is_published=True,
                    publish_date=timezone.now(),
                    deadline=quiz.deadline,
                    allow_retake=quiz.allow_retake,
                    max_attempts=quiz.max_attempts,
                )
                QuizQuestion.objects.bulk_create([
                    QuizQuestion(
                        quiz=target_quiz,
                        question_number=question.question_number,
                        question_text=question.question_text,
                        option_a=question.option_a,
                        option_b=question.option_b,
                        option_c=question.option_c,
                        option_d=question.option_d,
                        correct_answer=question.correct_answer,
                        explanation=question.explanation,
                        marks=question.marks,
                    )
                    for question in questions
                ])
            assigned_quizzes.append(target_quiz)

    return Response({
        'message': f'Quiz assigned to {len(assigned_quizzes)} batch(es).',
        'quiz_ids': [item.id for item in assigned_quizzes],
    })


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def take_quiz(request, quiz_id):
    try:
        student = Students.objects.get(user=request.user)
        quiz = Quiz.objects.get(id=quiz_id)
    except (Students.DoesNotExist, Quiz.DoesNotExist):
        return Response({'error': 'Not found'}, status=404)
    existing_count = QuizAttempt.objects.filter(student=student, quiz=quiz).count()
    if not quiz.allow_retake and existing_count >= quiz.max_attempts:
        return Response({'error': 'Max attempts reached'}, status=400)
    attempt = QuizAttempt.objects.create(
        quiz=quiz, student=student,
        attempt_number=existing_count + 1,
        submitted_at=timezone.now(), is_completed=True
    )
    answers = request.data.get('answers', {})
    score = 0
    for q in quiz.questions.all():
        selected = answers.get(str(q.id), '')
        is_correct = bool(selected) and selected.upper() == q.correct_answer.upper()
        marks = q.marks if is_correct else 0
        score += marks
        QuizAnswer.objects.create(attempt=attempt, question=q, selected_answer=selected, is_correct=is_correct, marks_obtained=marks)
    attempt.score = score
    attempt.percentage = round(score / max(quiz.total_marks, 1) * 100, 1)
    attempt.is_passed = score >= quiz.passing_marks
    attempt.save()
    return Response({'attempt_id': attempt.id, 'score': score, 'total': quiz.total_marks, 'percentage': attempt.percentage, 'is_passed': attempt.is_passed})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def quiz_result(request, attempt_id):
    """Simple quiz result (for backward compatibility)"""
    try:
        attempt = QuizAttempt.objects.get(id=attempt_id)
        return Response(QuizAttemptSerializer(attempt).data)
    except QuizAttempt.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)


@api_view(['PATCH'])
@permission_classes([IsAuthenticated])
def toggle_quiz_publish(request, quiz_id):
    try:
        quiz = Quiz.objects.get(id=quiz_id)
        if not quiz.is_published and quiz.batch_id is None:
            return Response({'error': 'Assign this quiz to a batch before publishing.'}, status=400)
        quiz.is_published = not quiz.is_published
        if quiz.is_published:
            quiz.publish_date = timezone.now()
        quiz.save()
        return Response({'is_published': quiz.is_published})
    except Quiz.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)



@api_view(['DELETE'])
@permission_classes([IsAuthenticated])
def delete_quiz(request, quiz_id):
    try:
        Quiz.objects.get(id=quiz_id).delete()
        return Response(status=204)
    except Quiz.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)


def _serialize_quiz_with_questions(quiz, include_answers=False):
    questions = quiz.questions.all().order_by('question_number')
    questions_data = []
    total_marks = 0

    for q in questions:
        item = {
            'id': q.id,
            'question_text': q.question_text,
            'option_a': q.option_a or '',
            'option_b': q.option_b or '',
            'option_c': q.option_c or '',
            'option_d': q.option_d or '',
            'marks': q.marks,
        }
        if include_answers:
            item['correct_answer'] = q.correct_answer
        questions_data.append(item)
        total_marks += q.marks

    if total_marks == 0 and questions.count() > 0:
        total_marks = questions.count()

    category = getattr(quiz, 'category', '') or ''

    return {
        'id': quiz.id,
        'title': quiz.title,
        'description': quiz.description,
        'category': category,
        'total_questions': questions.count(),
        'total_marks': total_marks,
        'duration_minutes': quiz.duration_minutes,
        'passing_marks': quiz.passing_marks,
        'max_attempts': quiz.max_attempts,
        'questions': questions_data,
    }


@api_view(['GET'])
@permission_classes([AllowAny])
def practice_quizzes(request):
    """Public practice quizzes (batch-less, published)."""
    quizzes = Quiz.objects.filter(batch__isnull=True, is_published=True).order_by('-created_at')
    return Response({
        'results': [_serialize_quiz_with_questions(quiz, include_answers=True) for quiz in quizzes]
    })


def close_public_user_activity(public_user):
    activity = PublicUserActivity.objects.filter(
        public_user=public_user,
        logout_time__isnull=True,
    ).order_by('-login_time').first()
    if activity:
        now = timezone.now()
        activity.logout_time = now
        activity.last_seen = now
        activity.save(update_fields=['logout_time', 'last_seen'])


def serialize_public_user(user):
    return {
        'id': user.id,
        'name': user.name,
        'email': user.email,
        'mobile': user.mobile,
        'qualification': user.qualification,
        'location': user.location,
        'city': user.city,
        'state': user.state,
        'username': user.username,
        'created_at': user.created_at,
    }


@api_view(['POST'])
@permission_classes([AllowAny])
def public_user_register(request):
    username = str(request.data.get('username', '')).strip().lower()
    password = str(request.data.get('password', '')).strip()
    email = str(request.data.get('email', '')).strip().lower()

    if not username or not password:
        return Response({'error': 'Username and password are required.'}, status=400)

    existing_user = PublicUser.objects.filter(username__iexact=username)
    if email:
        existing_user = PublicUser.objects.filter(Q(username__iexact=username) | Q(email__iexact=email))

    if existing_user.exists():
        return Response({'error': 'Username or email already registered.'}, status=400)

    public_user = PublicUser.objects.create(
        name=str(request.data.get('name', '')).strip() or username,
        email=email,
        mobile=str(request.data.get('mobile', '')).strip(),
        qualification=str(request.data.get('qualification', '')).strip(),
        location=str(request.data.get('location', '')).strip(),
        city=str(request.data.get('city', '')).strip(),
        state=str(request.data.get('state', '')).strip(),
        username=username,
        password=make_password(password),
    )
    return Response({'success': True, 'data': serialize_public_user(public_user)}, status=201)


@api_view(['POST'])
@permission_classes([AllowAny])
def public_user_login(request):
    login_id = str(request.data.get('username', '')).strip().lower()
    password = str(request.data.get('password', '')).strip()

    public_user = PublicUser.objects.filter(
        Q(username__iexact=login_id) | Q(email__iexact=login_id)
    ).first()

    if not public_user or not check_password(password, public_user.password):
        return Response({'error': 'Invalid username/email or password.'}, status=401)

    close_public_user_activity(public_user)
    now = timezone.now()
    PublicUserActivity.objects.create(public_user=public_user, login_time=now, last_seen=now)
    return Response({'success': True, 'data': serialize_public_user(public_user)})


@api_view(['POST'])
@permission_classes([AllowAny])
def public_user_logout(request):
    username = str(request.data.get('username', '')).strip().lower()
    public_user = PublicUser.objects.filter(Q(username__iexact=username) | Q(email__iexact=username)).first()
    if public_user:
        close_public_user_activity(public_user)
    return Response({'success': True})


@api_view(['POST'])
@permission_classes([AllowAny])
def submit_practice_quiz(request, quiz_id):
    """Score a public practice quiz without requiring student authentication."""
    try:
        quiz = Quiz.objects.get(id=quiz_id, batch__isnull=True, is_published=True)
    except Quiz.DoesNotExist:
        return Response({'error': 'Practice quiz not found'}, status=404)

    answers_data = request.data.get('answers', {})
    questions = quiz.questions.all()
    total_marks = sum(q.marks for q in questions) or questions.count()
    score = 0
    correct_count = 0
    attempted_count = 0

    for question in questions:
        selected = str(answers_data.get(str(question.id), '')).strip().upper()
        if selected:
            attempted_count += 1
        if selected and selected == str(question.correct_answer).strip().upper():
            score += question.marks or 1
            correct_count += 1

    percentage = round((score / total_marks) * 100, 1) if total_marks else 0
    wrong_count = attempted_count - correct_count
    request_user = request.user if getattr(request, 'user', None) and request.user.is_authenticated else None
    username = str(request.data.get('username', '')).strip().lower()
    name = str(request.data.get('name', '')).strip()
    email = str(request.data.get('email', '')).strip().lower()
    mobile = str(request.data.get('mobile', '')).strip()

    if not username and request_user:
        username = request_user.username

    public_lookup = Q()
    for value in [username, email, mobile]:
        if value:
            public_lookup |= Q(username__iexact=value) | Q(email__iexact=value) | Q(mobile__iexact=value)
    public_user = PublicUser.objects.filter(public_lookup).first() if public_lookup else None

    if not username and public_user:
        username = public_user.username
    if not username and email:
        username = email
    if not username and mobile:
        username = mobile
    if not username and name:
        username = name.lower().replace(' ', '-')
    if not username:
        username = f"guest-{timezone.now().strftime('%Y%m%d%H%M%S%f')}"

    PublicPracticeResult.objects.create(
        public_user=public_user,
        username=username[:120],
        quiz_id=quiz.id,
        quiz_title=quiz.title,
        score=score,
        total_marks=total_marks,
        percentage=percentage,
        is_passed=percentage >= quiz.passing_marks,
        correct_count=correct_count,
        wrong_count=wrong_count,
        attempted_count=attempted_count,
        total_questions=questions.count(),
        completed_at=timezone.now(),
    )

    return Response({
        'score': score,
        'total_marks': total_marks,
        'percentage': percentage,
        'is_passed': percentage >= quiz.passing_marks,
        'correct_count': correct_count,
        'wrong_count': wrong_count,
        'attempted_count': attempted_count,
        'total_questions': questions.count(),
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_public_users(request):
    if not is_admin_user(request.user):
        return Response({'error': 'Admin access required.'}, status=status.HTTP_403_FORBIDDEN)

    tab = request.query_params.get('tab', 'users')

    if tab == 'results':
        results = PublicPracticeResult.objects.select_related('public_user').order_by('-completed_at')
        return Response({'results': [{
            'id': result.id,
            'name': result.public_user.name if result.public_user else '',
            'email': result.public_user.email if result.public_user else '',
            'mobile': result.public_user.mobile if result.public_user else '',
            'username': result.username,
            'quiz_title': result.quiz_title,
            'score': result.score,
            'total_marks': result.total_marks,
            'percentage': result.percentage,
            'is_passed': result.is_passed,
            'correct_count': result.correct_count,
            'attempted_count': result.attempted_count,
            'total_questions': result.total_questions,
            'completed_at': result.completed_at,
        } for result in results]})

    if tab == 'logins':
        activities = PublicUserActivity.objects.select_related('public_user').order_by('-login_time')
        return Response({'results': [{
            'id': activity.id,
            'name': activity.public_user.name,
            'username': activity.public_user.username,
            'email': activity.public_user.email,
            'login_time': activity.login_time,
            'logout_time': activity.logout_time,
            'last_seen': activity.last_seen,
        } for activity in activities]})

    users = PublicUser.objects.all().order_by('-created_at')
    return Response({'results': [serialize_public_user(user) for user in users]})


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def staff_quiz_results(request):
    attempts = QuizAttempt.objects.filter(is_completed=True, quiz__batch__isnull=False)
    include_public_results = is_admin_user(request.user) or request.user.is_superuser or request.user.is_staff

    if include_public_results:
        attempts = attempts
    else:
        try:
            emp = Employee.objects.get(user=request.user)
            batches = Batches.objects.filter(faculty=emp)
            attempts = attempts.filter(quiz__batch__in=batches)
        except Employee.DoesNotExist:
            return Response({'results': []})

    batch_id = request.query_params.get('batch_id') or request.query_params.get('batch')
    if batch_id:
        attempts = attempts.filter(quiz__batch_id=batch_id)

    attempts = attempts.select_related('student', 'quiz', 'quiz__batch').order_by('-submitted_at')

    data = []
    for attempt in attempts:
        batch = attempt.quiz.batch
        data.append({
            'id': attempt.id,
            'student_name': f"{attempt.student.first_name} {attempt.student.last_name or ''}".strip(),
            'student_id': attempt.student.student_id,
            'quiz_title': attempt.quiz.title,
            'score': attempt.score,
            'total_marks': attempt.quiz.total_marks,
            'total_questions': attempt.total_questions or attempt.quiz.total_questions,
            'percentage': attempt.percentage,
            'passing_marks': attempt.quiz.passing_marks,
            'is_passed': attempt.is_passed,
            'submitted_at': attempt.submitted_at,
            'batch_id': batch.id if batch else None,
            'batch_number': batch.batch_number if batch else None,
        })

    public_data = []

    if include_public_results and not batch_id:
        public_results = PublicPracticeResult.objects.select_related('public_user').order_by('-completed_at')
        quiz_ids = [result.quiz_id for result in public_results if result.quiz_id]
        quiz_map = {quiz.id: quiz for quiz in Quiz.objects.filter(id__in=quiz_ids)}
        result_usernames = [result.username for result in public_results if result.username]
        student_map = {}
        employee_map = {}
        if result_usernames:
            students = Students.objects.select_related('user').filter(
                Q(student_id__in=result_usernames) |
                Q(email__in=result_usernames) |
                Q(user__username__in=result_usernames)
            )
            student_map = {
                key.lower(): student
                for student in students
                for key in [student.student_id, student.email, student.user.username if student.user else '']
                if key
            }
            employees = Employee.objects.select_related('user').filter(
                Q(email__in=result_usernames) |
                Q(user__username__in=result_usernames)
            )
            employee_map = {
                key.lower(): employee
                for employee in employees
                for key in [employee.email, employee.user.username if employee.user else '']
                if key
            }

        for result in public_results:
            quiz = quiz_map.get(result.quiz_id)
            public_user = result.public_user
            lookup_key = (result.username or '').lower()
            matched_student = student_map.get(lookup_key)
            matched_employee = employee_map.get(lookup_key)
            if matched_student:
                student_name = f"{matched_student.first_name} {matched_student.last_name or ''}".strip()
                student_id = matched_student.student_id
                audience = 'student'
                contact_email = matched_student.email or ''
                contact_mobile = matched_student.mobile_no or ''
            elif matched_employee:
                student_name = f"{matched_employee.first_name} {matched_employee.last_name or ''}".strip()
                student_id = matched_employee.staff_id or result.username
                audience = 'employee'
                contact_email = matched_employee.email or ''
                contact_mobile = matched_employee.mobile_no or ''
            elif public_user:
                student_name = public_user.name or result.username
                student_id = public_user.username
                audience = 'public'
                contact_email = public_user.email or ''
                contact_mobile = public_user.mobile or ''
            else:
                student_name = result.username
                student_id = result.username
                audience = 'public'
                contact_email = ''
                contact_mobile = ''
            public_data.append({
                'id': f'public-{result.id}',
                'source': 'public',
                'audience': audience,
                'name': public_user.name if public_user else student_name,
                'email': contact_email,
                'mobile': contact_mobile,
                'username': result.username,
                'student_name': student_name,
                'student_id': student_id,
                'quiz_title': result.quiz_title,
                'score': result.score,
                'total_marks': result.total_marks,
                'total_questions': result.total_questions,
                'percentage': result.percentage,
                'passing_marks': quiz.passing_marks if quiz else 50,
                'is_passed': result.is_passed,
                'submitted_at': result.completed_at,
                'batch_id': None,
                'batch_number': 'Practice Test',
            })

        oldest_date = timezone.make_aware(datetime.min)
        public_data.sort(key=lambda item: item.get('submitted_at') or oldest_date, reverse=True)
    return Response({
        'results': data,
        'student_results': data,
        'public_results': public_data,
    })
        
# ------------------------------------------------------------------------------
# SESSIONS (UPDATED WORKING VERSION)
# ------------------------------------------------------------------------------

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def batch_sessions(request, batch_id):
    try:
        batch = Batches.objects.get(id=batch_id)
        sync_missing_sessions_from_logsheet(batch)
    except Batches.DoesNotExist:
        pass
    sessions = CourseSession.objects.filter(batch__id=batch_id).order_by('session_number')
    return Response(CourseSessionSerializer(sessions, many=True).data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def student_sessions(request):
    try:
        student = Students.objects.get(user=request.user)
        if not student.assigned_batch:
            return Response([])

        sync_missing_sessions_from_logsheet(student.assigned_batch)
        sessions = CourseSession.objects.filter(batch=student.assigned_batch).order_by('session_number')
        data = []

        for session in sessions:
            progress = Student_Session_Progress.objects.filter(student=student, session=session).first()
            status_obj = StudentSessionStatus.objects.filter(student=student, session=session).first()

            student_status = 'not_started'
            doubt_response = None
            has_response = False
            response_date = None

            # progress takes priority over status_obj
            if progress:
                student_status = progress.student_status
                latest_response = DoubtResponse.objects.filter(doubt=progress).order_by('-created_at').first()
                if latest_response:
                    has_response = True
                    doubt_response = latest_response.message
                    response_date = latest_response.created_at
            elif status_obj:
                student_status = status_obj.student_status

            # Use per-student progress flag, NOT global session.staff_completed
            student_progress_staff_done = progress.staff_completed if progress else False

            # If staff completed this for student but status not updated yet
            if student_progress_staff_done and student_status == 'not_started':
                student_status = 'pending'
                if progress:
                    progress.student_status = 'pending'
                    progress.save()

            data.append({
                'id': session.id,
                'session_number': session.session_number,
                'title': session.title,
                'topics': session.topics,
                'staff_completed': student_progress_staff_done,
                'student_status': student_status,
                'has_response': has_response,
                'doubt_response': doubt_response,
                'response_date': response_date,
                'student_confirmed_at': status_obj.student_confirmed_at if status_obj else None,
                'completed_date': session.completed_date,
            })

        return Response({'results': data})

    except Students.DoesNotExist:
        return Response({'error': 'Student not found'}, status=404)

        
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_batch_sessions_with_logsheet(request, batch_id):
    try:
        batch = Batches.objects.get(id=batch_id)
        sync_missing_sessions_from_logsheet(batch)
        sessions = CourseSession.objects.filter(batch=batch).order_by('session_number')

        # Get active students in this batch
        active_students = Students.objects.filter(assigned_batch=batch)

        sessions_data = []
        for session in sessions:
            session_dict = CourseSessionSerializer(session).data

            if active_students.exists():
                # staff_completed = True only if ALL active students
                # have a progress record with staff_completed=True
                staff_done_count = Student_Session_Progress.objects.filter(
                    session=session,
                    student__in=active_students,
                    staff_completed=True
                ).count()
                session_dict['staff_completed'] = (staff_done_count == active_students.count() and active_students.count() > 0)
            else:
                # No active students — reset to False so staff sees fresh
                session_dict['staff_completed'] = False
                if session.staff_completed:
                    session.staff_completed = False
                    session.completed_date = None
                    session.save()

            sessions_data.append(session_dict)

        return Response({
            'batch': BatchSerializer(batch).data,
            'logsheet_url': (
    batch.course_logsheet.url
    if batch.course_logsheet
    else batch.course_name.course_logsheet.url
    if batch.course_name and batch.course_name.course_logsheet
    else None
),
            'sessions': sessions_data,
        })
    except Batches.DoesNotExist:
        return Response({'error': 'Batch not found'}, status=404)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def staff_mark_session_complete(request, session_id):
    """Staff marks a session as completed - creates pending status for students"""
    try:
        emp = Employee.objects.get(user=request.user)
        session = CourseSession.objects.get(id=session_id)

        session.staff_completed = True
        session.completed_date = timezone.now()
        session.save()

        students = Students.objects.filter(assigned_batch=session.batch)

        for student in students:
            status_obj, created = StudentSessionStatus.objects.get_or_create(
                student=student, 
                session=session
            )
            status_obj.staff_completed = True
            status_obj.staff_completed_at = timezone.now()
            status_obj.status = 'pending'
            status_obj.save()
            
            progress, _ = Student_Session_Progress.objects.get_or_create(
                student=student, session=session
            )
            progress.staff_completed = True
            progress.staff_completed_at = timezone.now()
            progress.student_status = 'pending'
            progress.save()

            if student.user:
                SessionNotification.objects.create(
                    session=session,
                    from_user=request.user,
                    to_user=student.user,
                    notification_type='session_completed',
                    message=f"Session {session.session_number}: '{session.title}' has been completed by your trainer. Please confirm or raise a doubt.",
                    title=f"Session {session.session_number} Completed",
                    requires_action=True
                )

        return Response({
            'success': True,
            'message': f'Session {session.session_number} marked complete. {students.count()} students notified.',
            'session': CourseSessionSerializer(session).data,
        })
    except (Employee.DoesNotExist, CourseSession.DoesNotExist):
        return Response({'error': 'Not found'}, status=404)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def staff_unmark_session(request, session_id):
    try:
        session = CourseSession.objects.get(id=session_id)
        session.staff_completed = False
        session.completed_date = None
        session.save()
        StudentSessionStatus.objects.filter(session=session).update(
            staff_completed=False, status='pending'
        )
        Student_Session_Progress.objects.filter(session=session).update(
            staff_completed=False,
            staff_completed_at=None,
            student_status='pending',
        )
        return Response({
            'success': True,
            'message': 'Session unmarked.',
            'session': CourseSessionSerializer(session).data,
        })
    except CourseSession.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)



@api_view(['POST'])
@permission_classes([IsAuthenticated])
def student_mark_completed(request):
    try:
        student = Students.objects.get(user=request.user)
        session = CourseSession.objects.get(id=request.data.get('session_id'))

        # Update StudentSessionStatus
        status_obj, _ = StudentSessionStatus.objects.get_or_create(
            student=student, session=session
        )
        
        progress, _ = Student_Session_Progress.objects.get_or_create(
            student=student, session=session
        )
        trainer_completed = status_obj.staff_completed or progress.staff_completed or session.staff_completed
        if not trainer_completed:
            return Response({'error': 'Trainer has not completed this session yet.'}, status=400)

        # Mark as completed by student
        status_obj.staff_completed = True
        status_obj.staff_completed_at = status_obj.staff_completed_at or progress.staff_completed_at or session.completed_date or timezone.now()
        status_obj.student_status = 'completed'
        status_obj.student_confirmed_at = timezone.now()
        status_obj.save()

        # Update Student_Session_Progress
        progress.staff_completed = True
        progress.staff_completed_at = progress.staff_completed_at or status_obj.staff_completed_at or session.completed_date or timezone.now()
        progress.completed = True
        progress.completed_date = timezone.now()
        progress.student_status = 'completed'
        progress.student_confirmed_at = timezone.now()
        progress.save()

        # ========== CHECK IF ALL SESSIONS ARE COMPLETED ==========
        total_sessions = CourseSession.objects.filter(batch=student.assigned_batch).count()
        completed_sessions = Student_Session_Progress.objects.filter(
            student=student,
            session__batch=student.assigned_batch,  # ? only current batch
            completed=True
        ).count()

        is_course_completed = (total_sessions > 0 and completed_sessions == total_sessions)

        if is_course_completed:
            return Response({
                'success': True, 
                'message': 'All sessions confirmed. Your mentor will complete or request reassignment.',
                'course_completed': True
            })
        
        return Response({
            'success': True, 
            'message': 'Session marked as completed!',
            'course_completed': False
        })
    
    except (Students.DoesNotExist, CourseSession.DoesNotExist):
        return Response({'error': 'Not found'}, status=404)
    except Exception as e:
        print(f"Error in student_mark_completed: {e}")
        import traceback
        traceback.print_exc()
        return Response({'error': str(e)}, status=400)
    
    
    
@api_view(['POST'])
@permission_classes([IsAuthenticated])
def student_raise_doubt(request):
    try:
        student = Students.objects.get(user=request.user)
        session = CourseSession.objects.get(id=request.data.get('session_id'))
        doubt_text = request.data.get('doubt_text', '').strip()

        if not doubt_text:
            return Response({'error': 'Please describe your doubt.'}, status=400)

        status_obj, _ = StudentSessionStatus.objects.get_or_create(
            student=student, session=session
        )
        status_obj.status = 'doubt'
        status_obj.save()

        progress, _ = Student_Session_Progress.objects.get_or_create(
            student=student, session=session
        )
        progress.has_doubt = True
        progress.doubt_description = doubt_text
        progress.doubt_raised_at = timezone.now()
        progress.doubt_resolved = False
        progress.student_status = 'doubt'
        progress.save()

        emp = session.batch.faculty
        if emp and emp.user:
            SessionNotification.objects.create(
                session=session,
                from_user=request.user,
                to_user=emp.user,
                notification_type='doubt_raised',
                message=f"{student.first_name} {student.last_name or ''} raised a doubt on Session {session.session_number}: '{session.title}'. Doubt: {doubt_text}",
                title=f"Doubt Raised — Session {session.session_number}",
                requires_action=True
            )

        return Response({'success': True, 'message': 'Doubt raised. Your trainer has been notified.'})
    
    except (Students.DoesNotExist, CourseSession.DoesNotExist):
        return Response({'error': 'Not found'}, status=404)



@api_view(['POST'])
@permission_classes([IsAuthenticated])
def staff_reply_doubt(request, progress_id):
    try:
        emp = Employee.objects.get(user=request.user)
        progress = Student_Session_Progress.objects.get(id=progress_id)
        reply_text = request.data.get('reply', '').strip()

        if not reply_text:
            return Response({'error': 'Reply cannot be empty.'}, status=400)

        # Create doubt response
        doubt_response = DoubtResponse.objects.create(
            doubt=progress,
            staff=emp,
            message=reply_text
        )
        
        print(f"? Doubt response saved with ID: {doubt_response.id}")

        # CRITICAL: Update progress - mark doubt as resolved
        progress.doubt_resolved = True
        progress.doubt_resolved_at = timezone.now()
        # IMPORTANT: Change status to 'pending' so student sees the response
        # DO NOT set completed = True
        # DO NOT set student_status = 'completed'
        progress.student_status = 'pending'  # ? This should be 'pending', NOT 'completed'
        progress.save()

        # Update StudentSessionStatus
        status_obj, _ = StudentSessionStatus.objects.get_or_create(
            student=progress.student, 
            session=progress.session
        )
        status_obj.student_status = 'pending'  # ? This should be 'pending'
        status_obj.save()

        # Send notification to student
        if progress.student.user:
            notification = SessionNotification.objects.create(
                session=progress.session,
                from_user=request.user,
                to_user=progress.student.user,
                notification_type='doubt_resolved',
                title=f"Doubt Resolved - Session {progress.session.session_number}",
                message=f"Your doubt has been answered: {reply_text}\n\nPlease review and click 'Mark as Completed'.",
                requires_action=True,
                is_read=False
            )
            print(f"? Notification sent to student: {notification.id}")

        return Response({
            'success': True, 
            'message': 'Reply sent. Student must manually click Mark as Completed.'
        })
        
    except Employee.DoesNotExist:
        return Response({'error': 'Staff not found'}, status=404)
    except Student_Session_Progress.DoesNotExist:
        return Response({'error': 'Doubt not found'}, status=404)
    except Exception as e:
        print(f"Error in staff_reply_doubt: {e}")
        import traceback
        traceback.print_exc()
        return Response({'error': str(e)}, status=500)
    

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_staff_doubts_detail(request):
    try:
        emp = Employee.objects.get(user=request.user)
        my_batches = Batches.objects.filter(faculty=emp)
        doubts = Student_Session_Progress.objects.filter(
            session__batch__in=my_batches,
            has_doubt=True
        ).select_related('student', 'session').order_by('-doubt_raised_at')

        data = []
        for d in doubts:
            replies = DoubtResponse.objects.filter(doubt=d).values('message', 'created_at', 'staff__first_name')
            data.append({
                'id': d.id,
                'student_name': f"{d.student.first_name} {d.student.last_name or ''}",
                'student_id': d.student.student_id,
                'session_number': d.session.session_number,
                'session_title': d.session.title,
                'doubt_text': d.doubt_description,
                'raised_at': d.doubt_raised_at.strftime('%d %b %Y, %H:%M') if d.doubt_raised_at else '',
                'is_resolved': d.doubt_resolved,
                'replies': list(replies),
            })
        return Response(data)
    except Employee.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_student_notifications(request):
    try:
        sync_key = 'notifications_long_absence_sync'
        if not cache.get(sync_key):
            _sync_long_absence_notifications()
            cache.set(sync_key, True, 300)
        notifs = SessionNotification.objects.filter(to_user=request.user)
        employee = Employee.objects.filter(user=request.user).first()
        if employee:
            notifs = notifs.exclude(notification_type='session_completed')
        notifs = notifs.order_by('-created_at')[:50]
        data = []
        for n in notifs:
            message = n.message
            if n.notification_type == 'batch_duration_exceeded':
                batch = _batch_from_duration_notification(n)
                if batch:
                    message = _batch_duration_exceeded_message(batch, batch.faculty)
                    if n.message != message:
                        n.message = message
                        n.save(update_fields=['message'])
            data.append({
                'id': n.id,
                'type': n.notification_type,
                'title': n.title or '',
                'message': message,
                'is_read': n.is_read,
                'requires_action': n.requires_action,
                'created_at': n.created_at.strftime('%d %b %Y, %H:%M'),
                'session_id': n.session.id if n.session else None,
                'action_url': _notification_action_url(n, request.user),
            })
        return Response(data)
    except Exception as e:
        return Response({'error': str(e)}, status=500)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def mark_notification_read(request, notif_id):
    try:
        notif = SessionNotification.objects.get(id=notif_id, to_user=request.user)
        notif.is_read = True
        notif.save()
        return Response({'success': True})
    except SessionNotification.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def mark_all_notifications_read(request):
    updated = SessionNotification.objects.filter(to_user=request.user, is_read=False).update(is_read=True)
    return Response({'success': True, 'updated': updated})


# -- COMPLETED STUDENTS --------------------------------------------------------

class CompletedStudentListView(generics.ListAPIView):
    queryset = CompletedStudent.objects.all().order_by('-completion_date')
    serializer_class = CompletedStudentSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        qs = super().get_queryset().filter(completion_type='full')
        user = self.request.user

        normalize_param = lambda value: re.sub(r'\s+', ' ', value or '').strip().lower()
        branch_param = normalize_param(self.request.query_params.get('branch', ''))
        batch_param = normalize_param(self.request.query_params.get('batch', ''))
        course_param = normalize_param(self.request.query_params.get('course', ''))
        search_param = self.request.query_params.get('search', '').strip()
        date_from = parse_date(self.request.query_params.get('dateFrom') or self.request.query_params.get('date_from') or '')
        date_to = parse_date(self.request.query_params.get('dateTo') or self.request.query_params.get('date_to') or '')

        if branch_param:
            qs = qs.annotate(_branch_norm=Lower(Trim('branch'))).filter(_branch_norm=branch_param)
        if batch_param:
            qs = qs.annotate(_batch_norm=Lower(Trim('batch_number'))).filter(_batch_norm=batch_param)
        if course_param:
            qs = qs.annotate(
                _course_name_norm=Lower(Trim('course_name')),
                _course_norm=Lower(Trim('course')),
            ).filter(Q(_course_name_norm=course_param) | Q(_course_norm=course_param))
        if date_from:
            qs = qs.filter(completion_date__date__gte=date_from)
        if date_to:
            qs = qs.filter(completion_date__date__lte=date_to)
        if search_param:
            qs = qs.filter(
                Q(first_name__icontains=search_param) |
                Q(last_name__icontains=search_param) |
                Q(student_id__icontains=search_param) |
                Q(course_name__icontains=search_param) |
                Q(course__icontains=search_param) |
                Q(batch_number__icontains=search_param) |
                Q(branch__icontains=search_param) |
                Q(faculty_name__icontains=search_param)
            )

        if user.is_superuser or user.is_staff:
            return qs

        emp = Employee.objects.filter(user=user).first()
        if emp and emp.designation.lower() == 'counselor':
            if emp.branch:
                return qs.annotate(_counselor_branch_norm=Lower(Trim('branch'))).filter(
                    _counselor_branch_norm=normalize_param(emp.branch)
                )
            return qs.none()
        if emp:
            batch_numbers = Batches.objects.filter(faculty=emp).values_list('batch_number', flat=True)
            return qs.filter(
                Q(graduated_from_trainer=emp) |
                Q(batch_number__in=batch_numbers)
            ).distinct()

        student = Students.objects.filter(user=user).first()
        if student:
            return qs.filter(original_student_id=student.student_id)

        return qs.none()

    def list(self, request, *args, **kwargs):
        try:
            return super().list(request, *args, **kwargs)
        except Exception as e:
            logger.exception("Unhandled exception in CompletedStudentListView.list")
            return Response({
                'error': 'Completed students retrieval failed',
                'details': str(e),
                'view': 'CompletedStudentListView.list'
            }, status=500)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def completed_students_pdf(request):
    qs = CompletedStudentListView()
    qs.request = request
    students_qs = qs.get_queryset()

    normalize_param = lambda value: re.sub(r'\s+', ' ', value or '').strip().lower()
    branch = normalize_param(request.query_params.get('branch'))
    batch = normalize_param(request.query_params.get('batch'))
    course = normalize_param(request.query_params.get('course'))
    trainer = request.query_params.get('trainer')
    search = request.query_params.get('search')
    date_from = parse_date(request.query_params.get('dateFrom') or request.query_params.get('date_from') or '')
    date_to = parse_date(request.query_params.get('dateTo') or request.query_params.get('date_to') or '')

    if branch:
        students_qs = students_qs.annotate(_pdf_branch_norm=Lower(Trim('branch'))).filter(_pdf_branch_norm=branch)
    if batch:
        students_qs = students_qs.annotate(_pdf_batch_norm=Lower(Trim('batch_number'))).filter(_pdf_batch_norm=batch)
    if course:
        students_qs = students_qs.annotate(
            _pdf_course_name_norm=Lower(Trim('course_name')),
            _pdf_course_norm=Lower(Trim('course')),
        ).filter(Q(_pdf_course_name_norm=course) | Q(_pdf_course_norm=course))
    if trainer:
        students_qs = students_qs.filter(Q(faculty_name__iexact=trainer) | Q(graduated_from_trainer__first_name__icontains=trainer) | Q(graduated_from_trainer__last_name__icontains=trainer))
    if date_from:
        students_qs = students_qs.filter(completion_date__date__gte=date_from)
    if date_to:
        students_qs = students_qs.filter(completion_date__date__lte=date_to)
    if search:
        students_qs = students_qs.filter(
            Q(first_name__icontains=search) |
            Q(last_name__icontains=search) |
            Q(student_id__icontains=search) |
            Q(course_name__icontains=search) |
            Q(course__icontains=search) |
            Q(batch_number__icontains=search) |
            Q(branch__icontains=search) |
            Q(faculty_name__icontains=search)
        )

    rows = []
    for item in students_qs.order_by('-completion_date'):
        rows.append({
            'student': f"{item.first_name} {item.last_name or ''}".strip(),
            'student_id': item.student_id,
            'branch': item.branch,
            'batch': item.batch_number,
            'course': item.course_name or item.course,
            'trainer': item.faculty_name,
            'sessions': f"{item.completed_sessions_count}/{item.total_sessions_count}",
            'completion_date': item.completion_date,
            'attendance': f"{item.attendance_percentage}%",
            'avg_score': f"{item.average_test_score}%",
        })

    return build_monitoring_pdf_response(
        'IIE Completed Students',
        'Completed students report',
        [
            ('Student', 'student'),
            ('Student ID', 'student_id'),
            ('Branch', 'branch'),
            ('Batch', 'batch'),
            ('Course', 'course'),
            ('Trainer', 'trainer'),
            ('Sessions', 'sessions'),
            ('Completed', 'completion_date'),
            ('Attendance', 'attendance'),
            ('Avg Score', 'avg_score'),
        ],
        rows,
        'completed_students_report.pdf',
    )


# -- COMPLETION REQUESTS -------------------------------------------------------

def _completed_student_defaults(student, trainer, completed_sessions, total_sessions):
    att_total = StudentAttendance.objects.filter(student=student).count()
    att_present = StudentAttendance.objects.filter(student=student, status='Present').count()
    att_pct = round((att_present / att_total * 100) if att_total > 0 else 0, 1)

    test_results_qs = TestResult.objects.filter(student=student)
    avg_score = round(
        sum(t.percentage for t in test_results_qs) / test_results_qs.count(), 1
    ) if test_results_qs.count() > 0 else 0

    batch = student.assigned_batch
    trainer_name = f"{trainer.first_name} {trainer.last_name or ''}".strip() if trainer else 'N/A'

    return {
        'student_id': student.student_id,
        'email': student.email,
        'first_name': student.first_name,
        'last_name': student.last_name or '',
        'mobile_no': student.mobile_no,
        'date_of_birth': student.date_of_birth,
        'city': student.city,
        'state': student.state,
        'qualification': student.qualification,
        'course': student.course,
        'gender': student.gender,
        'branch': student.branch,
        'batch_number': batch.batch_number if batch else 'N/A',
        'batch_id': str(batch.id) if batch else 'N/A',
        'batch_start_date': batch.start_date if batch else timezone.now().date(),
        'batch_end_date': batch.end_date if batch else timezone.now().date(),
        'faculty_name': trainer_name,
        'course_name': student.course,
        'completed_sessions_count': completed_sessions,
        'total_sessions_count': total_sessions,
        'attendance_percentage': att_pct,
        'average_test_score': avg_score,
        'completion_type': 'full',
        'completion_percentage': 100,
    }


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def mark_student_completed_by_admin_or_counselor(request, student_id):
    try:
        student = Students.objects.select_related('assigned_batch', 'assigned_staff').get(id=student_id)
    except Students.DoesNotExist:
        return Response({'error': 'Student not found'}, status=404)

    actor_employee = Employee.objects.filter(user=request.user).first()
    is_admin = is_admin_user(request.user) or request.user.is_staff or request.user.is_superuser
    is_counselor = bool(actor_employee and (actor_employee.designation or '').strip().lower() == 'counselor')

    if not is_admin and not is_counselor:
        return Response({'error': 'Only admin or counselor can mark students completed.'}, status=403)

    if is_counselor and not is_admin:
        student_branch = (student.branch or '').strip().lower()
        counselor_branch = (actor_employee.branch or '').strip().lower()
        if student_branch != counselor_branch:
            return Response({'error': 'You can only complete students from your branch.'}, status=403)

    already_completed = CompletedStudent.objects.filter(completion_type='full').filter(
        Q(original_student_id=str(student.id)) | Q(original_student_id=student.student_id)
    ).first()
    if already_completed:
        return Response({
            'success': True,
            'message': 'Student is already completed.',
            'completed_student_id': already_completed.id,
        })

    total_sessions = CourseSession.objects.filter(batch=student.assigned_batch).count() if student.assigned_batch_id else 0
    completed_sessions = Student_Session_Progress.objects.filter(
        student=student,
        completed=True,
        student_status='completed',
    ).count()
    trainer = student.assigned_staff or actor_employee
    defaults = _completed_student_defaults(student, trainer, completed_sessions, total_sessions)

    with transaction.atomic():
        completed_student, _ = CompletedStudent.objects.get_or_create(
            original_student_id=str(student.id),
            defaults=defaults,
        )
        for key, value in defaults.items():
            setattr(completed_student, key, value)
        completed_student.graduated_from_trainer = student.assigned_staff
        completed_student.save()

        old_batch = student.assigned_batch
        student.assigned_batch = None
        student.assigned_staff = None
        student.save()

        Student_Session_Progress.objects.filter(student=student).delete()
        StudentSessionStatus.objects.filter(student=student).delete()
        DoubtResponse.objects.filter(doubt__student=student).delete()

    return Response({
        'success': True,
        'message': f'{student.first_name} moved to completed students.',
        'completed_student_id': completed_student.id,
        'previous_batch': old_batch.batch_number if old_batch else None,
    })


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def complete_student_by_trainer(request, student_id):
    try:
        trainer = Employee.objects.get(user=request.user)
        student = Students.objects.select_related('assigned_batch', 'assigned_staff').get(id=student_id)
    except Employee.DoesNotExist:
        return Response({'error': 'Employee not found'}, status=404)
    except Students.DoesNotExist:
        return Response({'error': 'Student not found'}, status=404)

    if student.assigned_staff_id != trainer.id and not (request.user.is_staff or request.user.is_superuser):
        return Response({'error': 'Only the assigned mentor can complete this student.'}, status=403)

    if not student.assigned_batch:
        return Response({'error': 'Student is not assigned to any batch.'}, status=400)

    total_sessions = CourseSession.objects.filter(batch=student.assigned_batch).count()
    student_completed = Student_Session_Progress.objects.filter(
        student=student,
        session__batch=student.assigned_batch,
        completed=True,
        student_status='completed',
    ).count()
    mentor_completed = Student_Session_Progress.objects.filter(
        student=student,
        session__batch=student.assigned_batch,
        staff_completed=True,
    ).count()

    if total_sessions == 0:
        return Response({'error': 'No sessions found for this batch.'}, status=400)
    if mentor_completed < total_sessions or student_completed < total_sessions:
        return Response({
            'error': 'Complete all mentor sessions and wait for student confirmation before moving to completed.'
        }, status=400)

    defaults = _completed_student_defaults(student, trainer, student_completed, total_sessions)

    with transaction.atomic():
        try:
            completed_student, created = CompletedStudent.objects.get_or_create(
                original_student_id=str(student.id),
                graduated_from_trainer=trainer,
                defaults=defaults,
            )
        except DatabaseError:
            logger.exception("Database error creating CompletedStudent; retrying without optional completion fields")
            defaults.pop('completion_type', None)
            defaults.pop('completion_percentage', None)
            completed_student, created = CompletedStudent.objects.get_or_create(
                original_student_id=str(student.id),
                graduated_from_trainer=trainer,
                defaults=defaults,
            )

        if not created:
            for key, value in defaults.items():
                setattr(completed_student, key, value)
            completed_student.graduated_from_trainer = trainer
            completed_student.save()

        old_batch = student.assigned_batch
        student.assigned_batch = None
        student.assigned_staff = None
        student.save()

        Student_Session_Progress.objects.filter(student=student).delete()
        StudentSessionStatus.objects.filter(student=student).delete()
        DoubtResponse.objects.filter(doubt__student=student).delete()

    return Response({
        'success': True,
        'message': f"{student.first_name} moved to completed students.",
        'completed_student_id': completed_student.id,
        'previous_batch': old_batch.batch_number if old_batch else None,
    })


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def request_completion(request, student_id):
    try:
        emp = Employee.objects.get(user=request.user)
        student = Students.objects.get(id=student_id)
        counselor = Employee.objects.get(id=request.data.get('counselor_id'))
        req = SessionCompletionRequest.objects.create(
            student=student, batch=student.assigned_batch,
            trainer=emp, counselor=counselor,
            topics_covered=request.data.get('topics_covered', ''),
            sessions_completed=request.data.get('sessions_completed', 0),
            total_sessions=request.data.get('total_sessions', 0),
            message=request.data.get('message', ''),
        )
        return Response({'message': 'Request submitted.', 'id': req.id}, status=201)
    except (Employee.DoesNotExist, Students.DoesNotExist):
        return Response({'error': 'Not found'}, status=404)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def counselor_pending_requests(request):
    try:
        counselor = Employee.objects.get(user=request.user)
        reqs = SessionCompletionRequest.objects.filter(
            counselor=counselor, 
            status='pending'
        ).select_related('student', 'batch', 'batch__course_name', 'trainer')
        
        data = []
        for req in reqs:
            # If batch is None, try to get from student's assigned_batch
            batch = req.batch
            if not batch and req.student.assigned_batch:
                batch = req.student.assigned_batch
                # Also check if that batch has a course_name
                if batch and not hasattr(batch, 'course_name_obj'):
                    batch.course_name_obj = batch.course_name
            
            data.append({
                'id': req.id,
                'student_pk': req.student.id,
                'student_name': f"{req.student.first_name} {req.student.last_name or ''}".strip(),
                'student_id': req.student.student_id,
                'trainer_id': req.trainer.id if req.trainer else None,
                'trainer_name': f"{req.trainer.first_name} {req.trainer.last_name or ''}".strip(),
                'batch_id': batch.id if batch else None,
                # Get batch info with fallbacks:
                'batch_number': batch.batch_number if batch else req.student.assigned_batch.batch_number if req.student.assigned_batch else 'N/A',
                'course_name': batch.course_name.course_name if batch and batch.course_name else req.student.course if req.student.course else 'N/A',
                'course_id': batch.course_name.id if batch and batch.course_name else None,
                'course_type': batch.course_type if batch else 'N/A',
                'batch_timing': batch.batch_timing if batch else 'N/A',
                'branch': batch.branch if batch else req.student.branch,
                'counselor_id': req.counselor.id if req.counselor else None,
                'counselor_branch': req.counselor.branch if req.counselor else req.student.branch,
                'sessions_completed': req.sessions_completed,
                'total_sessions': req.total_sessions,
                'topics_covered': req.topics_covered,
                'message': req.message,
                'created_at': req.created_at,
                'status': req.status,
            })
        
        return Response(data)
        
    except Employee.DoesNotExist:
        return Response({'error': 'Counselor not found'}, status=404)


# In api_views.py
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def counselor_approved_requests(request):
    try:
        counselor = Employee.objects.get(user=request.user)
        # Get both approved and reassigned requests
        reqs = SessionCompletionRequest.objects.filter(
            counselor=counselor
        ).exclude(status='pending').order_by('-reviewed_at')
        
        data = []
        for req in reqs:
            # For reassigned requests, show the new trainer name
            trainer_name = ""
            if req.status == 'reassigned' and req.new_trainer:
                trainer_name = f"{req.new_trainer.first_name} {req.new_trainer.last_name or ''}".strip()
            else:
                trainer_name = f"{req.trainer.first_name} {req.trainer.last_name or ''}".strip()
            
            data.append({
                'id': req.id,
                'student_name': f"{req.student.first_name} {req.student.last_name or ''}".strip(),
                'trainer_name': trainer_name,
                'old_trainer_name': f"{req.trainer.first_name} {req.trainer.last_name or ''}".strip() if req.status == 'reassigned' else None,
                'sessions_completed': req.sessions_completed,
                'total_sessions': req.total_sessions,
                'topics_covered': req.topics_covered,
                'message': req.message,
                'status': req.status,
                'counselor_notes': req.counselor_notes,
                'created_at': req.created_at,
                'reviewed_at': req.reviewed_at,
                'reassigned_at': req.reassigned_at if hasattr(req, 'reassigned_at') else None,
            })
        return Response(data)
    except Employee.DoesNotExist:
        return Response({'error': 'Counselor not found'}, status=404)

@api_view(['PATCH'])
@permission_classes([IsAuthenticated])
def process_completion_request(request, pk):
    try:
        req = SessionCompletionRequest.objects.get(id=pk)
        req.status = request.data.get('status', req.status)
        req.counselor_notes = request.data.get('notes', req.counselor_notes)
        req.reviewed_at = timezone.now()
        req.reviewed_by = request.user
        req.save()
        return Response(SessionCompletionRequestSerializer(req).data)
    except SessionCompletionRequest.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_batch_students(request, batch_id):
    students = Students.objects.filter(assigned_batch__id=batch_id).select_related('assigned_batch', 'assigned_staff')
    return Response([_student_detail_payload(student) for student in students])


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_trainers(request):
    branch = request.query_params.get('branch')
    qs = Employee.objects.filter(designation__in=['trainer', 'mentor'])
    if branch:
        qs = qs.filter(branch=branch)
    return Response(EmployeeSerializer(qs, many=True).data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def trainer_batches(request, trainer_id):
    batches = Batches.objects.filter(faculty__id=trainer_id)
    return Response(BatchSerializer(batches, many=True).data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def counselor_student_details(request):
    try:
        counselor = Employee.objects.get(user=request.user)
        students = Students.objects.filter(branch=counselor.branch).select_related('assigned_batch', 'assigned_staff')
    except Employee.DoesNotExist:
        students = Students.objects.all().select_related('assigned_batch', 'assigned_staff')
    return Response(StudentSerializer(students, many=True).data)


def _student_detail_payload(student):
    att_qs = list(StudentAttendance.objects.filter(student=student).select_related('batch', 'staff').order_by('-date')[:50])
    total_att = len(att_qs)
    present_att = len([a for a in att_qs if str(a.status).lower() == 'present'])
    test_qs = list(TestResult.objects.filter(student=student).select_related('test').order_by('-submitted_at'))
    quiz_qs = list(QuizAttempt.objects.filter(student=student).select_related('quiz').order_by('-submitted_at'))
    assigned_quiz_qs = Quiz.objects.filter(
        batch=student.assigned_batch,
        is_published=True
    ).prefetch_related('questions').order_by('-created_at') if student.assigned_batch_id else Quiz.objects.none()
    leave_qs = list(StudentLeaveApplication.objects.filter(student=student).order_by('-applied_at'))
    progress_qs = Student_Session_Progress.objects.filter(student=student)
    total_sessions = progress_qs.count()
    sessions_completed = progress_qs.filter(completed=True).count()
    base = StudentSerializer(student).data
    base.update({
        'attendance_records': [{
            'date': str(a.date),
            'status': a.status,
            'batch_number': a.batch.batch_number if a.batch else '',
            'remarks': a.remarks or '-',
            'marked_by': f"{a.staff.first_name} {a.staff.last_name}" if a.staff else '-',
        } for a in att_qs],
        'attendance_total': total_att,
        'attendance_present': present_att,
        'attendance_percentage': round((present_att / total_att * 100) if total_att else 0, 1),
        'test_results': [{
            'test_name': t.test.title if t.test else '',
            'score': t.score,
            'percentage': t.percentage,
            'submitted_at': str(t.submitted_at)[:10] if t.submitted_at else '',
        } for t in test_qs],
        'test_count': len(test_qs),
        'average_score': round(sum(t.percentage for t in test_qs) / len(test_qs), 1) if test_qs else 0,
        'quiz_results': [{
            'quiz_title': a.quiz.title if a.quiz else '—',
            'score': a.score or 0,
            'total_marks': a.quiz.total_marks if a.quiz else 0,
            'percentage': float(a.percentage or 0),
            'passing_marks': a.quiz.passing_marks if a.quiz else 50,
            'submitted_at': str(a.submitted_at)[:10] if a.submitted_at else '',
        } for a in quiz_qs],
        'quiz_count': assigned_quiz_qs.count(),
        'quiz_attempt_count': len(quiz_qs),
        'assigned_quizzes': [{
            'id': q.id,
            'quiz_title': q.title,
            'total_questions': q.questions.count(),
            'passing_marks': q.passing_marks,
            'duration_minutes': q.duration_minutes,
            'status': 'completed' if any(a.quiz_id == q.id for a in quiz_qs) else 'assigned',
        } for q in assigned_quiz_qs],
        'leave_requests': [{
            'leave_type': l.leave_type,
            'start_date': str(l.start_date),
            'end_date': str(l.end_date),
            'status': l.status,
            'reason': l.reason,
        } for l in leave_qs],
        'leaves_count': len(leave_qs),
        'sessions_completed': sessions_completed,
        'total_sessions': total_sessions,
        'progress_percentage': round((sessions_completed / total_sessions * 100) if total_sessions else 0, 1),
        'doubts_count': progress_qs.filter(has_doubt=True).count(),
        'resolved_doubts': progress_qs.filter(doubt_resolved=True).count(),
    })
    return base


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def counselor_student_detail(request, student_id):
    try:
        counselor = Employee.objects.get(user=request.user, designation__iexact='counselor')
        student = Students.objects.select_related('assigned_batch', 'assigned_staff').get(id=student_id, branch=counselor.branch)
        return Response(_student_detail_payload(student))
    except Employee.DoesNotExist:
        return Response({'error': 'Counselor access required.'}, status=403)
    except Students.DoesNotExist:
        return Response({'error': 'Student not found'}, status=404)



@api_view(['POST'])
@permission_classes([IsAuthenticated])
def counselor_reassign_student(request, request_id):
    """Counselor reassigns a completion request/student to staff and a batch."""
    try:
        counselor = Employee.objects.get(user=request.user)
        completion_req = SessionCompletionRequest.objects.get(id=request_id, counselor=counselor)
        target_counselor_id = request.data.get('target_counselor_id') or counselor.id
        new_trainer_id = request.data.get('new_trainer_id')
        batch_mode = request.data.get('batch_mode', 'existing')
        batch_id = request.data.get('batch_id')
        counselor_notes = request.data.get('notes', '')

        try:
            target_counselor = Employee.objects.get(id=target_counselor_id, designation__iexact='counselor')
        except Employee.DoesNotExist:
            return Response({'error': 'Selected counselor not found'}, status=404)

        if not new_trainer_id:
            return Response({'error': 'Please select a staff member'}, status=400)

        try:
            new_trainer = Employee.objects.get(id=new_trainer_id)
        except Employee.DoesNotExist:
            return Response({'error': 'Selected staff not found'}, status=404)

        if (new_trainer.designation or '').lower() not in ('trainer', 'mentor'):
            return Response({'error': 'Selected employee must be a staff/mentor'}, status=400)
        if new_trainer.branch != target_counselor.branch:
            return Response({'error': 'Selected staff must be from the selected counselor branch'}, status=400)

        student = completion_req.student
        source_batch = completion_req.batch or student.assigned_batch
        if not source_batch:
            return Response({'error': 'Original batch not found for this request'}, status=400)

        if batch_mode not in ('new', 'existing'):
            return Response({'error': 'Invalid batch option'}, status=400)

        def copy_student_progress_to_batch(target_batch, copy_batch_sessions=False):
            source_sessions = CourseSession.objects.filter(batch=source_batch).order_by('session_number')
            source_by_number = {session.session_number: session for session in source_sessions}

            if copy_batch_sessions:
                for source_session in source_sessions:
                    CourseSession.objects.get_or_create(
                        batch=target_batch,
                        session_number=source_session.session_number,
                        defaults={
                            'title': source_session.title,
                            'topics': source_session.topics,
                            'staff_completed': source_session.staff_completed,
                            'session_enabled': source_session.session_enabled,
                            'completed_date': source_session.completed_date,
                        }
                    )

            target_sessions = CourseSession.objects.filter(batch=target_batch).order_by('session_number')
            for target_session in target_sessions:
                source_session = source_by_number.get(target_session.session_number)
                if not source_session:
                    continue

                old_progress = Student_Session_Progress.objects.filter(
                    student=student,
                    session=source_session
                ).first()
                old_status = StudentSessionStatus.objects.filter(
                    student=student,
                    session=source_session
                ).first()
                if not old_progress and not old_status and not source_session.staff_completed:
                    continue

                staff_done = bool(
                    (old_progress and old_progress.staff_completed) or
                    (old_status and old_status.staff_completed) or
                    source_session.staff_completed
                )
                student_status = (
                    old_progress.student_status if old_progress else
                    old_status.student_status if old_status else
                    'pending' if staff_done else 'not_started'
                )
                completed = bool(old_progress and old_progress.completed) or student_status == 'completed'

                progress, _ = Student_Session_Progress.objects.get_or_create(
                    student=student,
                    session=target_session
                )
                progress.staff_completed = staff_done
                progress.staff_completed_at = (
                    old_progress.staff_completed_at if old_progress else
                    old_status.staff_completed_at if old_status else
                    source_session.completed_date
                )
                progress.completed = completed
                progress.completed_date = old_progress.completed_date if old_progress else None
                progress.student_status = student_status
                progress.student_confirmed_at = (
                    old_progress.student_confirmed_at if old_progress else
                    old_status.student_confirmed_at if old_status else None
                )
                progress.save()

                status_obj, _ = StudentSessionStatus.objects.get_or_create(
                    student=student,
                    session=target_session
                )
                status_obj.staff_completed = staff_done
                status_obj.staff_completed_at = progress.staff_completed_at
                status_obj.student_status = 'completed' if completed else ('pending' if staff_done else 'pending')
                status_obj.student_confirmed_at = progress.student_confirmed_at
                status_obj.save()

        with transaction.atomic():
            if batch_mode == 'new':
                new_batch = Batches.objects.create(
                    batch_number=generate_batch_number(new_trainer.branch),
                    course_type=source_batch.course_type,
                    course_name=source_batch.course_name,
                    faculty=new_trainer,
                    start_date=source_batch.start_date,
                    end_date=source_batch.end_date,
                    batch_timing=source_batch.batch_timing,
                    branch=new_trainer.branch,
                )
                copy_student_progress_to_batch(new_batch, copy_batch_sessions=True)
                target_batch = new_batch
            else:
                if not batch_id:
                    return Response({'error': 'Please select a batch'}, status=400)
                try:
                    target_batch = Batches.objects.get(id=batch_id, faculty=new_trainer)
                except Batches.DoesNotExist:
                    return Response({'error': 'Selected batch not found for this staff'}, status=404)
                copy_student_progress_to_batch(target_batch, copy_batch_sessions=False)

            student.previous_trainer = completion_req.trainer
            student.previous_batch = source_batch
            student.is_transferred = True
            student.branch = target_batch.branch
            student.assigned_staff = new_trainer
            student.assigned_batch = target_batch
            student.transfer_date = timezone.now()
            student.save()

            completion_req.new_trainer = new_trainer
            completion_req.counselor_notes = counselor_notes
            completion_req.status = 'reassigned'
            completion_req.reviewed_at = timezone.now()
            completion_req.reviewed_by = request.user
            completion_req.reassigned_at = timezone.now()
            completion_req.save()

        student.refresh_from_db()

        return Response({
            'success': True,
            'message': 'Student reassigned successfully',
            'details': {
                'student_name': f"{student.first_name} {student.last_name or ''}".strip(),
                'student_id': student.student_id,
                'receiving_counselor': f"{target_counselor.first_name} {target_counselor.last_name or ''}".strip(),
                'receiving_staff': f"{new_trainer.first_name} {new_trainer.last_name or ''}".strip(),
                'receiving_batch': target_batch.batch_number,
                'receiving_branch': target_batch.branch,
                'batch_created': batch_mode == 'new',
            }
        })

    except SessionCompletionRequest.DoesNotExist:
        return Response({'error': 'Request not found'}, status=404)
    except Employee.DoesNotExist:
        return Response({'error': 'Employee not found'}, status=404)
    except Exception as e:
        import traceback
        traceback.print_exc()
        return Response({'error': str(e)}, status=500)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def quiz_staff_results_view(request):
    try:
        emp = Employee.objects.get(user=request.user)
        batches = Batches.objects.filter(faculty=emp)
        attempts = QuizAttempt.objects.filter(quiz__batch__in=batches, is_completed=True).order_by('-submitted_at')
        return Response(QuizAttemptSerializer(attempts, many=True).data)
    except Employee.DoesNotExist:
        return Response([])


# -- ADMIN: BRANCH-WISE ATTENDANCE OVERVIEW ------------------------------------

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_employee_monitoring(request):
    if not is_admin_user(request.user):
        return Response({'error': 'Admin access required.'}, status=status.HTTP_403_FORBIDDEN)

    params = request.query_params
    qs = UserActivity.objects.filter(user_type='employee').select_related('user', 'employee')
    qs = prepare_activity_monitoring_queryset(qs)
    qs = apply_activity_date_filters(qs, params)

    branch = params.get('branch')
    designation = params.get('designation')
    search = params.get('search')

    if branch:
        qs = qs.filter(employee__branch=branch)
    if designation:
        qs = qs.filter(employee__designation=designation)
    if search:
        qs = qs.filter(
            Q(employee__first_name__icontains=search) |
            Q(employee__last_name__icontains=search) |
            Q(employee__email__icontains=search) |
            Q(employee__staff_id__icontains=search) |
            Q(user__username__icontains=search)
        )

    records = [format_employee_activity(activity) for activity in qs.order_by('-login_time')]
    return Response({
        'results': records,
        'filters': {
            'branches': list(Employee.objects.exclude(branch__isnull=True).exclude(branch='').values_list('branch', flat=True).distinct().order_by('branch')),
            'designations': list(Employee.objects.exclude(designation__isnull=True).exclude(designation='').values_list('designation', flat=True).distinct().order_by('designation')),
        },
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_employee_monitoring_pdf(request):
    if not is_admin_user(request.user):
        return Response({'error': 'Admin access required.'}, status=status.HTTP_403_FORBIDDEN)

    params = request.query_params
    qs = UserActivity.objects.filter(user_type='employee').select_related('user', 'employee')
    qs = prepare_activity_monitoring_queryset(qs)
    qs = apply_activity_date_filters(qs, params)

    branch = params.get('branch')
    designation = params.get('designation')
    search = params.get('search')

    if branch:
        qs = qs.filter(employee__branch=branch)
    if designation:
        qs = qs.filter(employee__designation=designation)
    if search:
        qs = qs.filter(
            Q(employee__first_name__icontains=search) |
            Q(employee__last_name__icontains=search) |
            Q(employee__email__icontains=search) |
            Q(employee__staff_id__icontains=search) |
            Q(user__username__icontains=search)
        )

    records = [format_employee_activity(activity) for activity in qs.order_by('-login_time')]
    columns = [
        ('Name', 'name'),
        ('Email / Staff ID', 'email'),
        ('Designation', 'designation'),
        ('Branch', 'branch'),
        ('Login Time', 'login_time'),
        ('Logout Time', 'logout_time'),
        ('Last Seen', 'last_seen'),
    ]
    for record in records:
        if record.get('staff_id'):
            record['email'] = f"{record.get('email') or '-'} / {record['staff_id']}"
        if not record.get('logout_time'):
            record['logout_time'] = 'Still active'

    return build_monitoring_pdf_response(
        'Employee Monitoring Report',
        'Login, logout, and last-seen activity for employees and counselors',
        columns,
        records,
        'employee_monitoring_report.pdf',
    )


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_student_monitoring(request):
    if not is_admin_user(request.user):
        return Response({'error': 'Admin access required.'}, status=status.HTTP_403_FORBIDDEN)

    params = request.query_params
    qs = UserActivity.objects.filter(user_type='student').select_related('user', 'student', 'student__assigned_staff')
    qs = prepare_activity_monitoring_queryset(qs)
    qs = apply_activity_date_filters(qs, params)

    branch = params.get('branch')
    staff = params.get('staff')
    search = params.get('search')

    if branch:
        qs = qs.filter(student__branch=branch)
    if staff and str(staff).isdigit():
        qs = qs.filter(student__assigned_staff_id=staff)
    if search:
        qs = qs.filter(
            Q(student__first_name__icontains=search) |
            Q(student__last_name__icontains=search) |
            Q(student__email__icontains=search) |
            Q(student__student_id__icontains=search) |
            Q(user__username__icontains=search)
        )

    staff_members = Employee.objects.order_by('first_name', 'last_name')
    records = [format_student_activity(activity) for activity in qs.order_by('-login_time')]
    return Response({
        'results': records,
        'filters': {
            'branches': list(Students.objects.exclude(branch__isnull=True).exclude(branch='').values_list('branch', flat=True).distinct().order_by('branch')),
            'staff': [
                {
                    'id': staff_member.id,
                    'name': f"{staff_member.first_name} {staff_member.last_name or ''}".strip(),
                    'branch': staff_member.branch,
                    'designation': staff_member.designation,
                }
                for staff_member in staff_members
            ],
        },
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_student_monitoring_pdf(request):
    if not is_admin_user(request.user):
        return Response({'error': 'Admin access required.'}, status=status.HTTP_403_FORBIDDEN)

    params = request.query_params
    qs = UserActivity.objects.filter(user_type='student').select_related('user', 'student', 'student__assigned_staff')
    qs = prepare_activity_monitoring_queryset(qs)
    qs = apply_activity_date_filters(qs, params)

    branch = params.get('branch')
    staff = params.get('staff')
    search = params.get('search')

    if branch:
        qs = qs.filter(student__branch=branch)
    if staff and str(staff).isdigit():
        qs = qs.filter(student__assigned_staff_id=staff)
    if search:
        qs = qs.filter(
            Q(student__first_name__icontains=search) |
            Q(student__last_name__icontains=search) |
            Q(student__email__icontains=search) |
            Q(student__student_id__icontains=search) |
            Q(user__username__icontains=search)
        )

    records = [format_student_activity(activity) for activity in qs.order_by('-login_time')]
    columns = [
        ('Student', 'name'),
        ('Email / Student ID', 'email'),
        ('Branch', 'branch'),
        ('Staff Name', 'staff_name'),
        ('Login Time', 'login_time'),
        ('Logout Time', 'logout_time'),
        ('Last Seen', 'last_seen'),
    ]
    for record in records:
        if record.get('student_id'):
            record['email'] = f"{record.get('email') or '-'} / {record['student_id']}"
        if not record.get('logout_time'):
            record['logout_time'] = 'Still active'

    return build_monitoring_pdf_response(
        'Student Monitoring Report',
        'Login, logout, and last-seen activity for students',
        columns,
        records,
        'student_monitoring_report.pdf',
    )


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def mentor_student_monitoring(request):
    try:
        mentor = Employee.objects.get(user=request.user)
    except Employee.DoesNotExist:
        return Response({'error': 'Mentor access required.'}, status=status.HTTP_403_FORBIDDEN)

    if mentor.designation.lower() == 'counselor':
        return Response({'error': 'Mentor access required.'}, status=status.HTTP_403_FORBIDDEN)

    params = request.query_params
    qs = UserActivity.objects.filter(
        user_type='student',
        student__assigned_staff=mentor,
    ).select_related('user', 'student', 'student__assigned_staff')
    qs = prepare_activity_monitoring_queryset(qs)
    qs = apply_activity_date_filters(qs, params)

    search = params.get('search')
    if search:
        qs = qs.filter(
            Q(student__first_name__icontains=search) |
            Q(student__last_name__icontains=search) |
            Q(student__email__icontains=search) |
            Q(student__student_id__icontains=search) |
            Q(user__username__icontains=search)
        )

    records = [format_student_activity(activity) for activity in qs.order_by('-login_time')]
    return Response({
        'results': records,
        'filters': {
            'branches': [mentor.branch] if mentor.branch else [],
        },
    })


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def mentor_student_monitoring_pdf(request):
    try:
        mentor = Employee.objects.get(user=request.user)
    except Employee.DoesNotExist:
        return Response({'error': 'Mentor access required.'}, status=status.HTTP_403_FORBIDDEN)

    if mentor.designation.lower() == 'counselor':
        return Response({'error': 'Mentor access required.'}, status=status.HTTP_403_FORBIDDEN)

    params = request.query_params
    qs = UserActivity.objects.filter(
        user_type='student',
        student__assigned_staff=mentor,
    ).select_related('user', 'student', 'student__assigned_staff')
    qs = prepare_activity_monitoring_queryset(qs)
    qs = apply_activity_date_filters(qs, params)

    search = params.get('search')
    if search:
        qs = qs.filter(
            Q(student__first_name__icontains=search) |
            Q(student__last_name__icontains=search) |
            Q(student__email__icontains=search) |
            Q(student__student_id__icontains=search) |
            Q(user__username__icontains=search)
        )

    records = [format_student_activity(activity) for activity in qs.order_by('-login_time')]
    columns = [
        ('Student', 'name'),
        ('Email / Student ID', 'email'),
        ('Branch', 'branch'),
        ('Login Time', 'login_time'),
        ('Logout Time', 'logout_time'),
        ('Last Seen', 'last_seen'),
    ]
    for record in records:
        if record.get('student_id'):
            record['email'] = f"{record.get('email') or '-'} / {record['student_id']}"
        if not record.get('logout_time'):
            record['logout_time'] = 'Still active'

    return build_monitoring_pdf_response(
        'Student Login Records Report',
        'Login, logout, and last-seen activity for assigned students',
        columns,
        records,
        'student_login_records_report.pdf',
    )


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_branch_attendance(request):
    if not is_admin_user(request.user):
        return Response({'error': 'Admin access required.'}, status=403)
    from django.db.models import Q, Avg
    branch = request.query_params.get('branch')
    staff_id = request.query_params.get('staff_id')

    branches = list(Employee.objects.values_list('branch', flat=True).distinct().order_by('branch'))
    branch_stats = {}
    for b in branches:
        branch_stats[b] = {
            'staff_count': Employee.objects.filter(branch=b).count(),
            'staff_with_batches': Employee.objects.filter(branch=b, batches__isnull=False).distinct().count(),
        }

    if not branch and not staff_id:
        total_staff = Employee.objects.count()
        staff_with_batches = Employee.objects.filter(batches__isnull=False).distinct().count()
        return Response({
            'view_mode': 'branches',
            'branches': branches,
            'branch_stats': branch_stats,
            'total_staff': total_staff,
            'staff_with_batches': staff_with_batches,
            'total_branches': len(branches),
            'total_students': Students.objects.count(),
            'total_batches': Batches.objects.count(),
            'total_attendance': StudentAttendance.objects.count(),
        })

    if branch and not staff_id:
        staff_list = Employee.objects.filter(branch=branch).order_by('first_name')
        attendance_data = []
        for staff in staff_list:
            batches = Batches.objects.filter(faculty=staff)
            if batches.exists():
                batch_data = []
                for batch in batches:
                    recent = list(StudentAttendance.objects.filter(batch=batch).select_related('student').order_by('-date')[:10])
                    batch_data.append({'batch': BatchSerializer(batch).data, 'attendance_count': len(recent)})
                attendance_data.append({
                    'staff': EmployeeSerializer(staff).data,
                    'batches': batch_data,
                    'batch_count': batches.count(),
                    'attendance_count': StudentAttendance.objects.filter(staff=staff).count(),
                })
        return Response({
            'view_mode': 'branch_staff',
            'branches': branches,
            'branch_stats': branch_stats,
            'branch_name': branch,
            'attendance_data': attendance_data,
            'staff_in_branch': staff_list.count(),
            'staff_with_batches': len(attendance_data),
        })

    if staff_id:
        try:
            staff = Employee.objects.get(id=staff_id)
            batches = Batches.objects.filter(faculty=staff)
            students = Students.objects.filter(
                Q(assigned_staff=staff) | Q(assigned_batch__faculty=staff)
            ).distinct()

            student_details = []
            for student in students:
                # -- Attendance --------------------------------------------
                att_qs = list(StudentAttendance.objects.filter(student=student).select_related('batch', 'staff').order_by('-date')[:50])
                total_att = len(att_qs)
                present_att = len([a for a in att_qs if a.status == 'Present'])
                att_pct = round((present_att / total_att * 100) if total_att > 0 else 0, 1)

                # -- Leave -------------------------------------------------
                leave_qs = list(StudentLeaveApplication.objects.filter(student=student).order_by('-applied_at'))

                # -- Support -----------------------------------------------
                support_qs = list(StudentSupportRequest.objects.filter(student=student.user).order_by('-created_at'))

                # -- Test Results ------------------------------------------
                test_qs = list(TestResult.objects.filter(student=student).select_related('test').order_by('-submitted_at'))

                # -- Quiz Results ------------------------------------------
                quiz_qs = list(QuizAttempt.objects.filter(student=student).select_related('quiz').order_by('-submitted_at'))

                # -- Session Progress --------------------------------------
                progress_qs = Student_Session_Progress.objects.filter(student=student)
                total_sessions = progress_qs.count()
                sessions_completed = progress_qs.filter(completed=True).count()
                progress_pct = round((sessions_completed / total_sessions * 100) if total_sessions > 0 else 0, 1)
                doubts_count = progress_qs.filter(has_doubt=True).count()
                resolved_doubts = progress_qs.filter(doubt_resolved=True).count()
                last_progress = progress_qs.filter(completed=True).order_by('-completed_date').first()
                last_activity = last_progress.completed_date if last_progress else None
                # ---------------------------------------------------------

                student_details.append({
                    'student': StudentSerializer(student).data,
                    'attendance_records': [{
                        'date': str(a.date),
                        'status': a.status,
                        'batch_number': a.batch.batch_number if a.batch else '',
                        'remarks': a.remarks or '-',
                        'marked_by': f"{a.staff.first_name} {a.staff.last_name}" if a.staff else '-'
                    } for a in att_qs],
                    'attendance_total': total_att,
                    'attendance_present': present_att,
                    'attendance_percentage': att_pct,
                    'test_results': [{
                        'test_name': t.test.title if t.test else '',
                        'score': t.score,
                        'percentage': t.percentage,
                        'submitted_at': str(t.submitted_at)[:10] if t.submitted_at else ''
                    } for t in test_qs],
                    'test_count': len(test_qs),
                    'average_score': round(sum(t.percentage for t in test_qs) / len(test_qs), 1) if test_qs else 0,
                    'leave_requests': [{
                        'leave_type': l.leave_type,
                        'start_date': str(l.start_date),
                        'end_date': str(l.end_date),
                        'status': l.status,
                        'reason': l.reason
                    } for l in leave_qs],
                    'leaves_count': len(leave_qs),
                    'support_requests': [{
                        'message': s.message,
                        'status': s.status,
                        'created_at': str(s.created_at)[:10] if s.created_at else ''
                    } for s in support_qs],
                    'support_count': len(support_qs),
                    # -- Session progress fields ---------------------------
                    'sessions_completed': sessions_completed,
                    'total_sessions': total_sessions,
                    'progress_percentage': progress_pct,
                    'doubts_count': doubts_count,
                    'resolved_doubts': resolved_doubts,
                    'last_activity': last_activity,
                    # -- Quiz Results --------------------------------------
                    'quiz_results': [{
                        'quiz_title': a.quiz.title if a.quiz else '—',
                        'score': a.score or 0,
                        'total_marks': a.quiz.total_marks if a.quiz else 0,
                        'percentage': float(a.percentage or 0),
                        'passing_marks': a.quiz.passing_marks if a.quiz else 50,
                        'submitted_at': str(a.submitted_at)[:10] if a.submitted_at else '',
                    } for a in quiz_qs],
                    'quiz_count': len(quiz_qs),
                })

            all_att = StudentAttendance.objects.filter(student__in=students)
            total_att_all = all_att.count()
            present_att_all = all_att.filter(status='Present').count()

            return Response({
                'view_mode': 'staff_details',
                'branches': branches,
                'branch_stats': branch_stats,
                'branch_name': staff.branch,
                'staff': EmployeeSerializer(staff).data,
                'student_details': student_details,
                'total_batches': batches.count(),
                'total_students': students.count(),
                'total_sessions': CourseSession.objects.filter(batch__in=batches).count(),
                'overall_attendance_total': total_att_all,
                'overall_attendance_present': present_att_all,
                'overall_attendance_percentage': round((present_att_all / total_att_all * 100) if total_att_all > 0 else 0, 1),
            })
        except Employee.DoesNotExist:
            pass

    return Response({'view_mode': 'branches', 'branches': branches, 'branch_stats': branch_stats})


def _tracking_branch_values(branch):
    branch = (branch or '').strip()
    if branch in ('kuniyamuthur', 'kunniyamuthur'):
        return ['kuniyamuthur', 'kunniyamuthur']
    return [branch]


def _tracking_login_usage(staff):
    week_start = timezone.now() - timedelta(days=7)
    login_count = UserActivity.objects.filter(
        employee=staff,
        user_type='employee',
        login_time__gte=week_start,
    ).count()
    login_target = 7
    return {
        'login_usage_count': login_count,
        'login_usage_target': login_target,
        'login_usage_percentage': round(min(100, (login_count / login_target * 100) if login_target else 0), 1),
    }


def _is_counselor_staff(staff):
    return (getattr(staff, 'designation', '') or '').strip().lower() == 'counselor'


def _tracking_percent_of_target(value, target):
    return round(min(100, (value / target * 100) if target else 0), 1)


def _tracking_counselor_metrics(staff):
    now = timezone.now()
    week_since = now - timedelta(days=7)
    month_since = now - timedelta(days=30)
    branch_values = _tracking_branch_values(staff.branch)

    branch_students = Students.objects.filter(branch__in=branch_values)
    branch_batches = Batches.objects.filter(branch__in=branch_values)
    assigned_students = branch_students.filter(Q(assigned_staff__isnull=False) | Q(assigned_batch__isnull=False))
    fee_payments = FeePayment.objects.filter(student__branch__in=branch_values)
    fee_transactions = FeeTransaction.objects.filter(
        Q(collected_by=staff.user) | Q(fee_payment__student__branch__in=branch_values)
    ).distinct()
    login_usage = _tracking_login_usage(staff)

    week_students_added = branch_students.filter(created_at__gte=week_since).count()
    month_students_added = branch_students.filter(created_at__gte=month_since).count()
    week_batches_added = branch_batches.filter(created_at__gte=week_since).count()
    month_batches_added = branch_batches.filter(created_at__gte=month_since).count()
    week_students_assigned = assigned_students.filter(updated_at__gte=week_since).count()
    month_students_assigned = assigned_students.filter(updated_at__gte=month_since).count()
    week_fee_managed = fee_transactions.filter(paid_at__gte=week_since).count()
    month_fee_managed = fee_transactions.filter(paid_at__gte=month_since).count()

    performance_graph = {
        'students_added': _tracking_percent_of_target(month_students_added, 30),
        'batches_added': _tracking_percent_of_target(month_batches_added, 6),
        'students_assigned': _tracking_percent_of_target(month_students_assigned, 30),
        'fee_management': _tracking_percent_of_target(month_fee_managed, 30),
        'login_usage': login_usage['login_usage_percentage'],
    }
    activity_score = round(sum(performance_graph.values()) / len(performance_graph), 1)

    return {
        'branch_students': branch_students,
        'branch_batches': branch_batches,
        'assigned_students': assigned_students,
        'fee_payments': fee_payments,
        'fee_transactions': fee_transactions,
        'login_usage': login_usage,
        'week_since': week_since,
        'month_since': month_since,
        'week_students_added': week_students_added,
        'month_students_added': month_students_added,
        'week_batches_added': week_batches_added,
        'month_batches_added': month_batches_added,
        'week_students_assigned': week_students_assigned,
        'month_students_assigned': month_students_assigned,
        'week_fee_managed': week_fee_managed,
        'month_fee_managed': month_fee_managed,
        'performance_graph': performance_graph,
        'activity_score': min(activity_score, 100),
    }


def _tracking_batch_completion(batch, staff, batch_students=None):
    sessions = list(CourseSession.objects.filter(batch=batch).order_by('session_number'))
    session_ids = [session.id for session in sessions]
    total_sessions = len(session_ids)
    if batch_students is None:
        batch_students = Students.objects.filter(assigned_batch=batch)
    student_ids = list(batch_students.values_list('id', flat=True))
    student_count = len(student_ids)
    if not total_sessions or not student_count:
        return {'completed_sessions': 0, 'total_sessions': total_sessions, 'percentage': 0}

    staff_done_session_ids = set(CourseSession.objects.filter(id__in=session_ids, staff_completed=True).values_list('id', flat=True))
    staff_done_session_ids.update(DailySessionCompletion.objects.filter(
        faculty=staff,
        session_id__in=session_ids,
        completed=True,
    ).values_list('session_id', flat=True))

    staff_progress_counts = {
        row['session_id']: row['done_count']
        for row in Student_Session_Progress.objects.filter(
            session_id__in=session_ids,
            student_id__in=student_ids,
            staff_completed=True,
        ).values('session_id').annotate(done_count=Count('student_id', distinct=True))
    }
    staff_done_session_ids.update(
        session_id for session_id, count in staff_progress_counts.items() if count == student_count
    )

    student_completed_counts = {
        row['session_id']: row['done_count']
        for row in Student_Session_Progress.objects.filter(
            session_id__in=session_ids,
            student_id__in=student_ids,
        ).filter(
            Q(completed=True) | Q(student_status='completed')
        ).values('session_id').annotate(done_count=Count('student_id', distinct=True))
    }

    completed_sessions = 0
    for session_id in session_ids:
        if session_id not in staff_done_session_ids:
            continue
        students_done = student_completed_counts.get(session_id, 0)
        # Older session rows only stored staff_completed. Use that data until students start confirming sessions.
        if students_done == 0 and staff_progress_counts.get(session_id, 0) == student_count:
            students_done = staff_progress_counts.get(session_id, 0)
        if students_done == student_count:
            completed_sessions += 1

    return {
        'completed_sessions': completed_sessions,
        'total_sessions': total_sessions,
        'percentage': round((completed_sessions / total_sessions * 100) if total_sessions else 0, 1),
    }


def _tracking_staff_card_summary(staff):
    if _is_counselor_staff(staff):
        metrics = _tracking_counselor_metrics(staff)
        login_qs = UserActivity.objects.filter(employee=staff, user_type='employee').order_by('-login_time')
        last_login = login_qs.first()
        return {
            'staff': EmployeeSerializer(staff).data,
            'tracking_type': 'counselor',
            'batch_count': metrics['branch_batches'].count(),
            'student_count': metrics['branch_students'].count(),
            'assigned_students_count': metrics['assigned_students'].count(),
            'completed_students_count': CompletedStudent.objects.filter(branch__in=_tracking_branch_values(staff.branch), completion_type='full').count(),
            'attendance_marked_count': 0,
            'attendance_days_count': 0,
            'sessions_completed': 0,
            'total_sessions': 0,
            'session_completion_percentage': 0,
            'materials_uploaded': 0,
            'materials_assigned': 0,
            'material_batch_count': 0,
            'tests_created': 0,
            'tests_assigned': 0,
            'quizzes_created': 0,
            'quiz_attempts': 0,
            'new_batches_count': metrics['month_batches_added'],
            'new_students_count': metrics['month_students_added'],
            'week_students_added': metrics['week_students_added'],
            'month_students_added': metrics['month_students_added'],
            'week_batches_added': metrics['week_batches_added'],
            'month_batches_added': metrics['month_batches_added'],
            'week_students_assigned': metrics['week_students_assigned'],
            'month_students_assigned': metrics['month_students_assigned'],
            'week_fee_managed': metrics['week_fee_managed'],
            'month_fee_managed': metrics['month_fee_managed'],
            'fee_records_count': metrics['fee_payments'].count(),
            'fee_transactions_count': metrics['fee_transactions'].count(),
            'login_count': login_qs.count(),
            'login_usage_count': metrics['login_usage']['login_usage_count'],
            'login_usage_target': metrics['login_usage']['login_usage_target'],
            'last_login': last_login.login_time if last_login else None,
            'last_seen': last_login.last_seen if last_login else None,
            'performance_graph': metrics['performance_graph'],
            'activity_score': metrics['activity_score'],
        }

    batches = Batches.objects.filter(faculty=staff)
    batch_count = batches.count()
    students = Students.objects.filter(Q(assigned_staff=staff) | Q(assigned_batch__faculty=staff)).distinct()
    sessions_total = CourseSession.objects.filter(batch__in=batches).count()
    sessions_completed = max(
        DailySessionCompletion.objects.filter(faculty=staff, session__batch__in=batches, completed=True).values('session_id').distinct().count(),
        CourseSession.objects.filter(batch__in=batches, staff_completed=True).count(),
        Student_Session_Progress.objects.filter(session__batch__in=batches, staff_completed=True).values('session_id').distinct().count(),
    )
    login_qs = UserActivity.objects.filter(employee=staff, user_type='employee').order_by('-login_time')
    last_login = login_qs.first()
    login_usage = _tracking_login_usage(staff)
    materials_uploaded = StudyMaterial.objects.filter(uploaded_by=staff).count()
    quizzes_created = Quiz.objects.filter(created_by=staff).count()
    batch_completion_percentage = round((sessions_completed / sessions_total * 100) if sessions_total else 0, 1)
    score_parts = [
        batch_completion_percentage,
        min((quizzes_created / max(batch_count, 1)) * 100, 100),
        min((materials_uploaded / max(batch_count, 1)) * 100, 100),
        login_usage['login_usage_percentage'],
    ]
    activity_score = round(sum(score_parts) / len(score_parts), 1)
    return {
        'staff': EmployeeSerializer(staff).data,
        'batch_count': batch_count,
        'student_count': students.count(),
        'completed_students_count': CompletedStudent.objects.filter(graduated_from_trainer=staff, completion_type='full').count(),
        'attendance_marked_count': StudentAttendance.objects.filter(staff=staff).count(),
        'attendance_days_count': StudentAttendance.objects.filter(staff=staff).values('date').distinct().count(),
        'sessions_completed': sessions_completed,
        'total_sessions': sessions_total,
        'session_completion_percentage': batch_completion_percentage,
        'materials_uploaded': materials_uploaded,
        'materials_assigned': 0,
        'material_batch_count': 0,
        'tests_created': QuizTest.objects.filter(created_by=staff).count(),
        'tests_assigned': 0,
        'quizzes_created': quizzes_created,
        'quiz_attempts': 0,
        'new_batches_count': batches.filter(created_at__gte=timezone.now() - timedelta(days=7)).count(),
        'new_students_count': students.filter(created_at__gte=timezone.now() - timedelta(days=7)).count(),
        'login_count': login_qs.count(),
        'login_usage_count': login_usage['login_usage_count'],
        'login_usage_target': login_usage['login_usage_target'],
        'last_login': last_login.login_time if last_login else None,
        'last_seen': last_login.last_seen if last_login else None,
        'performance_graph': {},
        'activity_score': min(activity_score, 100),
    }


def _tracking_staff_summary(staff):
    if _is_counselor_staff(staff):
        metrics = _tracking_counselor_metrics(staff)
        login_qs = UserActivity.objects.filter(employee=staff, user_type='employee').order_by('-login_time')
        last_login = login_qs.first()
        return {
            'staff': EmployeeSerializer(staff).data,
            'tracking_type': 'counselor',
            'batch_count': metrics['branch_batches'].count(),
            'completed_batch_count': 0,
            'student_count': metrics['branch_students'].count(),
            'assigned_students_count': metrics['assigned_students'].count(),
            'completed_students_count': CompletedStudent.objects.filter(branch__in=_tracking_branch_values(staff.branch), completion_type='full').count(),
            'attendance_marked_count': 0,
            'attendance_days_count': 0,
            'sessions_completed': 0,
            'total_sessions': 0,
            'session_completion_percentage': 0,
            'materials_uploaded': 0,
            'materials_assigned': 0,
            'material_batch_count': 0,
            'tests_created': 0,
            'tests_assigned': 0,
            'quizzes_created': 0,
            'quiz_attempts': 0,
            'new_batches_count': metrics['month_batches_added'],
            'new_students_count': metrics['month_students_added'],
            'week_students_added': metrics['week_students_added'],
            'month_students_added': metrics['month_students_added'],
            'week_batches_added': metrics['week_batches_added'],
            'month_batches_added': metrics['month_batches_added'],
            'week_students_assigned': metrics['week_students_assigned'],
            'month_students_assigned': metrics['month_students_assigned'],
            'week_fee_managed': metrics['week_fee_managed'],
            'month_fee_managed': metrics['month_fee_managed'],
            'fee_records_count': metrics['fee_payments'].count(),
            'fee_transactions_count': metrics['fee_transactions'].count(),
            'login_count': login_qs.count(),
            'login_usage_count': metrics['login_usage']['login_usage_count'],
            'login_usage_target': metrics['login_usage']['login_usage_target'],
            'last_login': last_login.login_time if last_login else None,
            'last_seen': last_login.last_seen if last_login else None,
            'performance_graph': metrics['performance_graph'],
            'activity_score': metrics['activity_score'],
        }

    new_since = timezone.now() - timedelta(days=7)
    batches = Batches.objects.filter(faculty=staff)
    batch_count = batches.count()
    students = Students.objects.filter(
        Q(assigned_staff=staff) | Q(assigned_batch__faculty=staff)
    ).distinct()
    attendance_total = StudentAttendance.objects.filter(staff=staff).count()
    attendance_dates = StudentAttendance.objects.filter(staff=staff).values('date').distinct().count()
    batch_completion_stats = [_tracking_batch_completion(batch, staff) for batch in batches]
    sessions_total = sum(item['total_sessions'] for item in batch_completion_stats)
    sessions_completed = sum(item['completed_sessions'] for item in batch_completion_stats)
    completed_batch_count = sum(
        1 for item in batch_completion_stats
        if item['total_sessions'] > 0 and item['completed_sessions'] == item['total_sessions']
    )
    login_qs = UserActivity.objects.filter(employee=staff, user_type='employee').order_by('-login_time')
    last_login = login_qs.first()
    login_usage = _tracking_login_usage(staff)
    materials_uploaded = StudyMaterial.objects.filter(uploaded_by=staff).count()
    material_assignments = StudyMaterialAssignment.objects.filter(assigned_by=staff)
    materials_assigned = material_assignments.count()
    material_batch_count = material_assignments.values('batch_id').distinct().count()
    quizzes_created = Quiz.objects.filter(created_by=staff).count()
    quiz_attempts = QuizAttempt.objects.filter(quiz__batch__in=batches, is_completed=True).count()
    completed_students = CompletedStudent.objects.filter(graduated_from_trainer=staff, completion_type='full').count()

    def percent_of_target(value, target):
        return round(min(100, (value / target * 100) if target else 0), 1)

    batch_completion_percentage = round((completed_batch_count / batch_count * 100) if batch_count else 0, 1)
    performance_graph = {
        'batch_completion': batch_completion_percentage,
        'quiz_upload': percent_of_target(quizzes_created, max(batch_count, 1)),
        'material_upload': percent_of_target(materials_uploaded, max(batch_count, 1)),
        'login_usage': login_usage['login_usage_percentage'],
    }
    activity_score = round(sum(performance_graph.values()) / len(performance_graph), 1)

    return {
        'staff': EmployeeSerializer(staff).data,
        'batch_count': batch_count,
        'completed_batch_count': completed_batch_count,
        'student_count': students.count(),
        'completed_students_count': completed_students,
        'attendance_marked_count': attendance_total,
        'attendance_days_count': attendance_dates,
        'sessions_completed': sessions_completed,
        'total_sessions': sessions_total,
        'session_completion_percentage': batch_completion_percentage,
        'materials_uploaded': materials_uploaded,
        'materials_assigned': materials_assigned,
        'material_batch_count': material_batch_count,
        'tests_created': 0,
        'tests_assigned': 0,
        'quizzes_created': quizzes_created,
        'quiz_attempts': quiz_attempts,
        'new_batches_count': batches.filter(created_at__gte=new_since).count(),
        'new_students_count': students.filter(created_at__gte=new_since).count(),
        'login_count': login_qs.count(),
        'login_usage_count': login_usage['login_usage_count'],
        'login_usage_target': login_usage['login_usage_target'],
        'last_login': last_login.login_time if last_login else None,
        'last_seen': last_login.last_seen if last_login else None,
        'performance_graph': performance_graph,
        'activity_score': min(activity_score, 100),
    }


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_employee_tracking(request):
    if not is_admin_user(request.user):
        return Response({'error': 'Admin access required.'}, status=403)

    branch = request.query_params.get('branch')
    staff_id = request.query_params.get('staff_id')

    raw_branches = Employee.objects.exclude(branch__isnull=True).exclude(branch='').values_list('branch', flat=True).distinct()
    branch_order = ['100ft', 'hopes', 'kuniyamuthur']

    def canonical_branch(value):
        value = (value or '').strip()
        if value == 'kunniyamuthur':
            return 'kuniyamuthur'
        return value

    branches = sorted(
        {canonical_branch(item) for item in raw_branches if canonical_branch(item)},
        key=lambda value: (branch_order.index(value) if value in branch_order else len(branch_order), value),
    )

    branch_cards = []
    for item in branches:
        values = _tracking_branch_values(item)
        staff_qs = Employee.objects.filter(branch__in=values)
        branch_batches = Batches.objects.filter(branch__in=values)
        branch_students = get_active_students_queryset().filter(branch__in=values).distinct()
        branch_attendance = StudentAttendance.objects.filter(staff__branch__in=values).count()
        branch_daily_sessions = DailySessionCompletion.objects.filter(
            faculty__branch__in=values,
            completed=True,
        ).values('session_id').distinct().count()
        branch_direct_sessions = CourseSession.objects.filter(batch__branch__in=values, staff_completed=True).count()
        branch_progress_sessions = Student_Session_Progress.objects.filter(
            session__batch__branch__in=values,
            staff_completed=True,
        ).values('session_id').distinct().count()
        branch_sessions_completed = max(branch_daily_sessions, branch_direct_sessions, branch_progress_sessions)
        branch_material_uploads = StudyMaterial.objects.filter(uploaded_by__branch__in=values).count()
        branch_material_assignments = StudyMaterialAssignment.objects.filter(assigned_by__branch__in=values).count()
        branch_quizzes_created = Quiz.objects.filter(created_by__branch__in=values).count()
        branch_login_days = UserActivity.objects.filter(
            user_type='employee',
            employee__branch__in=values,
            login_time__gte=timezone.now() - timedelta(days=7),
        ).count()
        branch_score_parts = [
            min(branch_sessions_completed * 2, 30),
            min((branch_material_uploads + branch_material_assignments) * 3, 25),
            min(branch_quizzes_created * 4, 25),
            min(branch_login_days * 3, 20),
        ]
        branch_cards.append({
            'branch': item,
            'staff_count': staff_qs.count(),
            'student_count': branch_students.count(),
            'batch_count': branch_batches.count(),
            'attendance_marked_count': branch_attendance,
            'sessions_completed': branch_sessions_completed,
            'materials_count': branch_material_uploads + branch_material_assignments,
            'tests_count': 0,
            'quizzes_count': branch_quizzes_created,
            'login_count': UserActivity.objects.filter(user_type='employee', employee__branch__in=values).count(),
            'activity_score': min(round(sum(branch_score_parts), 1), 100),
        })

    if not branch and not staff_id:
        return Response({
            'view_mode': 'branches',
            'branches': branches,
            'branch_cards': branch_cards,
            'totals': {
                'branches': len(branches),
                'staff': Employee.objects.count(),
                'students': Students.objects.count(),
                'batches': Batches.objects.count(),
            },
        })

    if branch and not staff_id:
        staff_qs = Employee.objects.filter(branch__in=_tracking_branch_values(branch)).order_by('first_name', 'last_name')
        staff_summaries = [_tracking_staff_card_summary(staff) for staff in staff_qs]
        return Response({
            'view_mode': 'branch_staff',
            'branch': canonical_branch(branch),
            'branches': branches,
            'branch_cards': branch_cards,
            'staff': staff_summaries,
            'totals': {
                'staff': staff_qs.count(),
                'students': sum(item['student_count'] for item in staff_summaries),
                'batches': sum(item['batch_count'] for item in staff_summaries),
                'completed_students': sum(item['completed_students_count'] for item in staff_summaries),
            },
        })

    if staff_id:
        try:
            staff = Employee.objects.get(id=staff_id)
        except Employee.DoesNotExist:
            return Response({'error': 'Staff not found'}, status=404)

        summary = _tracking_staff_summary(staff)
        new_since = timezone.now() - timedelta(days=7)
        if summary.get('tracking_type') == 'counselor':
            branch_values = _tracking_branch_values(staff.branch)
            month_since = timezone.now() - timedelta(days=30)
            students = Students.objects.filter(branch__in=branch_values).select_related('assigned_batch', 'assigned_staff').order_by('first_name', 'last_name')
            last_month_students = students.filter(created_at__gte=month_since).order_by('-created_at')
            fee_transactions_qs = FeeTransaction.objects.filter(
                Q(collected_by=staff.user) | Q(fee_payment__student__branch__in=branch_values)
            ).select_related('fee_payment__student', 'fee_payment__batch').distinct().order_by('-paid_at')
            login_records = UserActivity.objects.filter(employee=staff, user_type='employee').order_by('-login_time')[:20]
            weekly_login_records = UserActivity.objects.filter(
                employee=staff,
                user_type='employee',
                login_time__gte=timezone.now() - timedelta(days=7),
            ).order_by('-login_time')
            student_rows = [{
                'id': student.id,
                'student_id': student.student_id,
                'name': f"{student.first_name} {student.last_name or ''}".strip(),
                'email': student.email,
                'mobile_no': student.mobile_no,
                'batch': student.assigned_batch.batch_number if student.assigned_batch else '',
                'assigned_staff': f"{student.assigned_staff.first_name} {student.assigned_staff.last_name or ''}".strip() if student.assigned_staff else '',
                'assigned': bool(student.assigned_batch_id or student.assigned_staff_id),
                'is_new': student.created_at >= new_since,
                'created_at': student.created_at,
                'updated_at': student.updated_at,
            } for student in students[:100]]
            last_month_student_rows = [{
                'id': student.id,
                'student_id': student.student_id,
                'name': f"{student.first_name} {student.last_name or ''}".strip(),
                'email': student.email,
                'mobile_no': student.mobile_no,
                'batch': student.assigned_batch.batch_number if student.assigned_batch else '',
                'assigned_staff': f"{student.assigned_staff.first_name} {student.assigned_staff.last_name or ''}".strip() if student.assigned_staff else '',
                'created_at': student.created_at,
            } for student in last_month_students[:100]]

            return Response({
                'view_mode': 'staff_detail',
                'branch': canonical_branch(staff.branch),
                'branches': branches,
                'branch_cards': branch_cards,
                'staff': summary,
                'charts': {
                    'students_added': summary['performance_graph']['students_added'],
                    'batches_added': summary['performance_graph']['batches_added'],
                    'students_assigned': summary['performance_graph']['students_assigned'],
                    'fee_management': summary['performance_graph']['fee_management'],
                    'login_usage': summary['performance_graph']['login_usage'],
                    'activity_score': summary['activity_score'],
                    'content_total': 0,
                },
                'highlights': {
                    'new_batches': summary['new_batches_count'],
                    'new_students': summary['new_students_count'],
                    'new_since': month_since,
                    'last_month_students_count': summary['month_students_added'],
                    'week_students_added': summary['week_students_added'],
                    'month_students_added': summary['month_students_added'],
                    'week_batches_added': summary['week_batches_added'],
                    'month_batches_added': summary['month_batches_added'],
                    'week_students_assigned': summary['week_students_assigned'],
                    'month_students_assigned': summary['month_students_assigned'],
                    'week_fee_managed': summary['week_fee_managed'],
                    'month_fee_managed': summary['month_fee_managed'],
                },
                'batches': [],
                'batch_details': [],
                'students': student_rows,
                'last_month_added_students': last_month_student_rows,
                'fee_transactions': [{
                    'student': f"{item.fee_payment.student.first_name} {item.fee_payment.student.last_name or ''}".strip() if item.fee_payment and item.fee_payment.student else '',
                    'student_id': item.fee_payment.student.student_id if item.fee_payment and item.fee_payment.student else '',
                    'batch': item.fee_payment.batch.batch_number if item.fee_payment and item.fee_payment.batch else '',
                    'amount': item.amount,
                    'payment_mode': item.payment_mode,
                    'paid_at': item.paid_at,
                    'bill_generated': item.bill_generated,
                } for item in fee_transactions_qs[:30]],
                'recent_attendance': [],
                'session_completions': [],
                'materials': [],
                'tests': [],
                'quizzes': [],
                'login_records': [{
                    'login_time': item.login_time,
                    'logout_time': item.logout_time,
                    'last_seen': item.last_seen,
                } for item in login_records],
                'weekly_login_records': [{
                    'login_time': item.login_time,
                    'logout_time': item.logout_time,
                    'last_seen': item.last_seen,
                } for item in weekly_login_records],
            })

        batches = Batches.objects.filter(faculty=staff).select_related('course_name').order_by('-created_at')
        students = Students.objects.filter(
            Q(assigned_staff=staff) | Q(assigned_batch__faculty=staff)
        ).select_related('assigned_batch').distinct().order_by('first_name', 'last_name')

        attendance_qs = StudentAttendance.objects.filter(staff=staff).select_related('student', 'batch').order_by('-date', '-created_at')
        attendance_total = attendance_qs.count()
        attendance_present = attendance_qs.filter(status='Present').count()
        session_completions = DailySessionCompletion.objects.filter(
            faculty=staff,
            session__batch__in=batches,
        ).select_related('session', 'session__batch').order_by('-completion_date')[:20]
        login_records = UserActivity.objects.filter(employee=staff, user_type='employee').order_by('-login_time')[:20]
        materials = StudyMaterial.objects.filter(uploaded_by=staff).select_related('batch').order_by('-uploaded_at')[:15]
        tests = QuizTest.objects.filter(created_by=staff).order_by('-created_at')[:15]
        quizzes = Quiz.objects.filter(created_by=staff).select_related('batch').order_by('-created_at')[:15]
        weekly_login_records = UserActivity.objects.filter(
            employee=staff,
            user_type='employee',
            login_time__gte=timezone.now() - timedelta(days=7),
        ).order_by('-login_time')

        student_rows = []
        for student in students[:80]:
            student_att = StudentAttendance.objects.filter(student=student, batch=student.assigned_batch)
            student_present = student_att.filter(status='Present').count()
            progress_qs = Student_Session_Progress.objects.filter(student=student, session__batch=student.assigned_batch)
            total_progress = progress_qs.count()
            done_progress = progress_qs.filter(Q(completed=True) | Q(student_status='completed')).count()
            if done_progress == 0:
                done_progress = progress_qs.filter(staff_completed=True).count()
            student_rows.append({
                'id': student.id,
                'student_id': student.student_id,
                'name': f"{student.first_name} {student.last_name or ''}".strip(),
                'email': student.email,
                'mobile_no': student.mobile_no,
                'batch': student.assigned_batch.batch_number if student.assigned_batch else '',
                'attendance_total': student_att.count(),
                'attendance_present': student_present,
                'attendance_percentage': round((student_present / student_att.count() * 100) if student_att.count() else 0, 1),
                'sessions_completed': done_progress,
                'total_sessions': total_progress,
                'session_percentage': round((done_progress / total_progress * 100) if total_progress else 0, 1),
                'is_new': student.created_at >= new_since,
                'created_at': student.created_at,
            })

        batch_details = []
        for batch in batches:
            batch_students = Students.objects.filter(
                Q(assigned_staff=staff) | Q(assigned_batch=batch),
                assigned_batch=batch,
            ).distinct().order_by('first_name', 'last_name')
            completion = _tracking_batch_completion(batch, staff, batch_students)
            total_sessions = completion['total_sessions']
            completed_sessions = completion['completed_sessions']
            batch_attendance = StudentAttendance.objects.filter(batch=batch, staff=staff)
            batch_attendance_total = batch_attendance.count()
            batch_attendance_present = batch_attendance.filter(status='Present').count()
            batch_material_uploads = StudyMaterial.objects.filter(uploaded_by=staff, batch=batch)
            batch_materials = StudyMaterialAssignment.objects.filter(assigned_by=staff, batch=batch)
            batch_material_ids = set(batch_material_uploads.values_list('id', flat=True))
            batch_material_ids.update(
                StudyMaterialAssignment.objects.filter(
                    material__uploaded_by=staff,
                    batch=batch,
                ).values_list('material_id', flat=True)
            )
            batch_uploaded_materials = StudyMaterial.objects.filter(id__in=batch_material_ids)
            batch_tests = AssignedTest.objects.filter(batch=batch, test__created_by=staff)
            batch_quizzes = Quiz.objects.filter(batch=batch, created_by=staff)

            batch_student_rows = []
            for student in batch_students:
                student_att = StudentAttendance.objects.filter(student=student, batch=batch)
                student_att_total = student_att.count()
                student_att_present = student_att.filter(status='Present').count()
                student_progress = Student_Session_Progress.objects.filter(student=student, session__batch=batch)
                student_sessions_total = student_progress.count() or total_sessions
                student_sessions_done = student_progress.filter(Q(completed=True) | Q(student_status='completed')).count()
                if student_sessions_done == 0:
                    student_sessions_done = student_progress.filter(staff_completed=True).count()
                batch_student_rows.append({
                    'id': student.id,
                    'student_id': student.student_id,
                    'name': f"{student.first_name} {student.last_name or ''}".strip(),
                    'mobile_no': student.mobile_no,
                    'email': student.email,
                    'created_at': student.created_at,
                    'is_new': student.created_at >= new_since,
                    'attendance_total': student_att_total,
                    'attendance_present': student_att_present,
                    'attendance_percentage': round((student_att_present / student_att_total * 100) if student_att_total else 0, 1),
                    'sessions_completed': student_sessions_done,
                    'total_sessions': student_sessions_total,
                    'session_percentage': round((student_sessions_done / student_sessions_total * 100) if student_sessions_total else 0, 1),
                })

            batch_details.append({
                'id': batch.id,
                'batch_number': batch.batch_number,
                'course': batch.course_name.course_name if batch.course_name else '',
                'course_type': batch.course_type,
                'timing': batch.batch_timing,
                'branch': batch.branch,
                'start_date': batch.start_date,
                'end_date': batch.end_date,
                'created_at': batch.created_at,
                'is_new': batch.created_at >= new_since,
                'student_count': batch_students.count(),
                'new_student_count': batch_students.filter(created_at__gte=new_since).count(),
                'total_sessions': total_sessions,
                'sessions_completed': completed_sessions,
                'session_percentage': completion['percentage'],
                'attendance_total': batch_attendance_total,
                'attendance_present': batch_attendance_present,
                'attendance_percentage': round((batch_attendance_present / batch_attendance_total * 100) if batch_attendance_total else 0, 1),
                'materials_uploaded': batch_uploaded_materials.count(),
                'materials_assigned': batch_materials.count(),
                'tests_uploaded': batch_tests.values('test_id').distinct().count(),
                'tests_assigned': batch_tests.count(),
                'quizzes_uploaded': batch_quizzes.count(),
                'content_upload_total': batch_uploaded_materials.count() + batch_quizzes.count(),
                'recent_uploads': (
                    [{
                        'kind': 'Material',
                        'title': item.title,
                        'date': item.uploaded_at,
                    } for item in batch_uploaded_materials.order_by('-uploaded_at')[:5]] +
                    [{
                        'kind': 'Test',
                        'title': item.test.title if item.test else '',
                        'date': item.assigned_date,
                    } for item in batch_tests.select_related('test').order_by('-assigned_date')[:5]] +
                    [{
                        'kind': 'Quiz',
                        'title': item.title,
                        'date': item.created_at,
                    } for item in batch_quizzes.order_by('-created_at')[:5]]
                )[:8],
                'students': batch_student_rows,
            })

        return Response({
            'view_mode': 'staff_detail',
            'branch': canonical_branch(staff.branch),
            'branches': branches,
            'branch_cards': branch_cards,
            'staff': summary,
            'charts': {
                'batch_completion': summary['performance_graph']['batch_completion'],
                'quiz_upload': summary['performance_graph']['quiz_upload'],
                'material_upload': summary['performance_graph']['material_upload'],
                'login_usage': summary['performance_graph']['login_usage'],
                'activity_score': summary['activity_score'],
                'content_total': summary['materials_uploaded'] + summary['quizzes_created'],
            },
            'highlights': {
                'new_batches': summary['new_batches_count'],
                'new_students': summary['new_students_count'],
                'new_since': new_since,
            },
            'batches': [{
                'id': batch.id,
                'batch_number': batch.batch_number,
                'course': batch.course_name.course_name if batch.course_name else '',
                'course_type': batch.course_type,
                'timing': batch.batch_timing,
                'start_date': batch.start_date,
                'end_date': batch.end_date,
                'student_count': Students.objects.filter(assigned_batch=batch).count(),
                'session_count': CourseSession.objects.filter(batch=batch).count(),
                'is_new': batch.created_at >= new_since,
            } for batch in batches],
            'batch_details': batch_details,
            'students': student_rows,
            'recent_attendance': [{
                'date': row.date,
                'student': f"{row.student.first_name} {row.student.last_name or ''}".strip() if row.student else '',
                'student_id': row.student.student_id if row.student else '',
                'batch': row.batch.batch_number if row.batch else '',
                'status': row.status,
                'remarks': row.remarks or '',
            } for row in attendance_qs[:30]],
            'session_completions': [{
                'date': item.completion_date,
                'session': item.session.title if item.session else '',
                'session_number': item.session.session_number if item.session else '',
                'batch': item.session.batch.batch_number if item.session and item.session.batch else '',
                'completed': item.completed,
                'topics_covered': item.topics_covered or '',
            } for item in session_completions],
            'materials': [{
                'title': item.title,
                'batch': item.batch.batch_number if item.batch else 'Library',
                'uploaded_at': item.uploaded_at,
                'is_library': item.is_library,
            } for item in materials],
            'tests': [{
                'title': item.title,
                'created_at': item.created_at,
                'assigned_count': AssignedTest.objects.filter(test=item).count(),
            } for item in tests],
            'quizzes': [{
                'title': item.title,
                'batch': item.batch.batch_number if item.batch else '',
                'created_at': item.created_at,
                'published': item.is_published,
                'attempts': item.attempts.filter(is_completed=True).count(),
            } for item in quizzes],
            'login_records': [{
                'login_time': item.login_time,
                'logout_time': item.logout_time,
                'last_seen': item.last_seen,
            } for item in login_records],
            'weekly_login_records': [{
                'login_time': item.login_time,
                'logout_time': item.logout_time,
                'last_seen': item.last_seen,
            } for item in weekly_login_records],
        })

    return Response({'error': 'Invalid tracking request'}, status=400)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_employee_tracking_pdf(request):
    if not is_admin_user(request.user):
        return Response({'error': 'Admin access required.'}, status=403)

    staff_id = request.query_params.get('staff_id')
    if not staff_id:
        return Response({'error': 'staff_id is required'}, status=400)

    try:
        staff = Employee.objects.get(id=staff_id)
    except Employee.DoesNotExist:
        return Response({'error': 'Staff not found'}, status=404)

    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.graphics.shapes import Circle, Drawing, Rect, String, Wedge
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    def pdf_text(value):
        if value is None or value == '':
            return '-'
        if hasattr(value, 'isoformat'):
            try:
                value = timezone.localtime(value)
            except Exception:
                pass
            return value.strftime('%d %b %Y, %I:%M %p')
        return str(value)

    def display_branch(value):
        labels = {
            '100ft': '100ft',
            'hopes': 'Hopes',
            'kuniyamuthur': 'Kuniyamuthur',
            'kunniyamuthur': 'Kuniyamuthur',
        }
        return labels.get(value, value or '-')

    summary = _tracking_staff_summary(staff)
    is_counselor_report = summary.get('tracking_type') == 'counselor'
    batches = (
        Batches.objects.filter(branch__in=_tracking_branch_values(staff.branch)).select_related('course_name').order_by('batch_number')
        if is_counselor_report
        else Batches.objects.filter(faculty=staff).select_related('course_name').order_by('batch_number')
    )
    attendance_qs = StudentAttendance.objects.filter(staff=staff)
    attendance_total = attendance_qs.count()
    attendance_present = attendance_qs.filter(status__iexact='Present').count()
    attendance_absent = attendance_qs.filter(status__iexact='Absent').count()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle('StaffReportTitle', parent=styles['Title'], fontName='Helvetica-Bold', fontSize=18, alignment=TA_CENTER, textColor=colors.HexColor('#0f172a'), spaceAfter=4)
    subtitle_style = ParagraphStyle('StaffReportSubtitle', parent=styles['Normal'], fontSize=9, alignment=TA_CENTER, textColor=colors.HexColor('#64748b'), spaceAfter=8)
    section_style = ParagraphStyle('StaffReportSection', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=11, textColor=colors.HexColor('#0f172a'), spaceBefore=10, spaceAfter=6)
    cell_style = ParagraphStyle('StaffReportCell', parent=styles['Normal'], fontSize=7.5, leading=9.5, textColor=colors.HexColor('#0f172a'))
    header_style = ParagraphStyle('StaffReportHeader', parent=cell_style, fontName='Helvetica-Bold', textColor=colors.white)

    def make_table(columns, rows, widths=None):
        data = [[Paragraph(escape(label), header_style) for label, _ in columns]]
        for row in rows:
            data.append([Paragraph(escape(pdf_text(row.get(key))), cell_style) for _, key in columns])
        if len(data) == 1:
            data.append([Paragraph('No records found', cell_style)] + [''] * (len(columns) - 1))
        table = Table(data, colWidths=widths, repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1572e8')),
            ('GRID', (0, 0), (-1, -1), 0.35, colors.HexColor('#d9e2ec')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        return table

    def progress_bar(value, color_hex):
        pct = max(0, min(100, float(value or 0)))
        width = 58 * mm
        height = 8 * mm
        drawing = Drawing(width, height)
        drawing.add(Rect(0, 2 * mm, width, 4 * mm, fillColor=colors.HexColor('#e2e8f0'), strokeColor=None, rx=2 * mm, ry=2 * mm))
        drawing.add(Rect(0, 2 * mm, width * pct / 100, 4 * mm, fillColor=colors.HexColor(color_hex), strokeColor=None, rx=2 * mm, ry=2 * mm))
        drawing.add(String(width + 4, 2 * mm, f'{pct:g}%', fontName='Helvetica-Bold', fontSize=8, fillColor=colors.HexColor('#0f172a')))
        return drawing

    def donut_chart(value, color_hex, label='', size=30 * mm):
        pct = max(0, min(100, float(value or 0)))
        drawing = Drawing(size, size + 9 * mm)
        center = size / 2
        radius = size * 0.42
        inner_radius = size * 0.27
        drawing.add(Circle(center, center + 6 * mm, radius, fillColor=colors.HexColor('#e2e8f0'), strokeColor=None))
        if pct > 0:
            arc_pct = min(pct, 99.9)
            drawing.add(Wedge(center, center + 6 * mm, radius, 90, 90 - (arc_pct * 3.6), fillColor=colors.HexColor(color_hex), strokeColor=None))
        drawing.add(Circle(center, center + 6 * mm, inner_radius, fillColor=colors.white, strokeColor=None))
        drawing.add(String(center, center + 4.5 * mm, f'{pct:g}%', textAnchor='middle', fontName='Helvetica-Bold', fontSize=10, fillColor=colors.HexColor(color_hex)))
        if label:
            drawing.add(String(center, 1.5 * mm, label, textAnchor='middle', fontName='Helvetica-Bold', fontSize=6.7, fillColor=colors.HexColor('#475569')))
        return drawing

    def metric_icon(color_hex):
        drawing = Drawing(18 * mm, 18 * mm)
        drawing.add(Rect(3 * mm, 3 * mm, 12 * mm, 12 * mm, fillColor=colors.HexColor(color_hex), strokeColor=None, rx=3 * mm, ry=3 * mm))
        drawing.add(Circle(9 * mm, 9 * mm, 2.4 * mm, fillColor=colors.white, strokeColor=None))
        return drawing

    def metric_card(label, value, subtext, color_hex):
        return Table(
            [[
                [
                    Paragraph(escape(str(value)), ParagraphStyle('MetricValue', parent=cell_style, fontName='Helvetica-Bold', fontSize=16, leading=18, textColor=colors.HexColor('#0f172a'))),
                    Paragraph(escape(label), ParagraphStyle('MetricLabel', parent=cell_style, fontName='Helvetica-Bold', fontSize=7.5, textColor=colors.HexColor('#64748b'))),
                    Paragraph(escape(subtext), ParagraphStyle('MetricSub', parent=cell_style, fontSize=6.5, textColor=colors.HexColor('#94a3b8'))),
                ],
                metric_icon(color_hex),
            ]],
            colWidths=[36 * mm, 31 * mm],
            rowHeights=[24 * mm],
            style=TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#f8fafc')),
                ('BOX', (0, 0), (-1, -1), 0.45, colors.HexColor('#d9e2ec')),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('LEFTPADDING', (0, 0), (-1, -1), 7),
                ('RIGHTPADDING', (0, 0), (-1, -1), 7),
                ('TOPPADDING', (0, 0), (-1, -1), 6),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ]),
        )

    def make_staff_hero():
        staff_title = f"{staff_name}"
        staff_meta = f"{staff.designation or 'Staff'} | {display_branch(staff.branch)} | Last login {pdf_text(summary.get('last_login'))}"
        badges = f"{summary['student_count']} students handled    {summary['batch_count']} batches    {summary['new_students_count']} new students"
        return Table(
            [[
                [
                    Paragraph(escape(staff_title), ParagraphStyle('HeroName', parent=cell_style, fontName='Helvetica-Bold', fontSize=17, leading=19, textColor=colors.white)),
                    Paragraph(escape(staff_meta), ParagraphStyle('HeroMeta', parent=cell_style, fontName='Helvetica-Bold', fontSize=7.5, textColor=colors.HexColor('#dbeafe'))),
                    Paragraph(escape(badges), ParagraphStyle('HeroBadges', parent=cell_style, fontName='Helvetica-Bold', fontSize=7, textColor=colors.white)),
                ],
                donut_chart(summary['performance_graph']['login_usage'], '#f59e0b', 'Login usage', 36 * mm),
                progress_bar(summary['performance_graph']['login_usage'], '#f59e0b'),
            ]],
            colWidths=[150 * mm, 42 * mm, 76 * mm],
            rowHeights=[33 * mm],
            style=TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#123047')),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('BOX', (0, 0), (-1, -1), 0.45, colors.HexColor('#123047')),
                ('LEFTPADDING', (0, 0), (-1, -1), 8),
                ('RIGHTPADDING', (0, 0), (-1, -1), 8),
                ('TOPPADDING', (0, 0), (-1, -1), 6),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ]),
        )

    def make_performance_table(rows):
        data = [[
            Paragraph('Metric', header_style),
            Paragraph('Count', header_style),
            Paragraph('Graph Percentage', header_style),
        ]]
        for row in rows:
            data.append([
                Paragraph(escape(pdf_text(row['metric'])), cell_style),
                Paragraph(escape(pdf_text(row['count'])), cell_style),
                progress_bar(row['percentage'], row['color']),
            ])
        table = Table(data, colWidths=[52 * mm, 44 * mm, 82 * mm], repeatRows=1)
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1572e8')),
            ('GRID', (0, 0), (-1, -1), 0.35, colors.HexColor('#d9e2ec')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f8fafc')]),
            ('LEFTPADDING', (0, 0), (-1, -1), 5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 5),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))
        return table

    def make_batch_header(batch_row):
        return Table(
            [[
                [
                    Paragraph(escape(batch_row['batch']), ParagraphStyle('BatchTitle', parent=cell_style, fontName='Helvetica-Bold', fontSize=9.5, textColor=colors.HexColor('#0f172a'))),
                    Paragraph(escape(batch_row['meta']), ParagraphStyle('BatchMeta', parent=cell_style, fontSize=7, textColor=colors.HexColor('#64748b'))),
                ],
                Paragraph(escape(f"{batch_row['students']} students    {batch_row['new_students']} new    {batch_row['uploads']} uploads    {batch_row['sessions']} sessions done"), ParagraphStyle('BatchBadges', parent=cell_style, fontName='Helvetica-Bold', fontSize=7, alignment=TA_CENTER, textColor=colors.HexColor('#0f766e'))),
                donut_chart(batch_row['session_percentage'], '#059669', 'Sessions', 30 * mm),
            ]],
            colWidths=[100 * mm, 118 * mm, 42 * mm],
            style=TableStyle([
                ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#ecfdf5')),
                ('BOX', (0, 0), (-1, -1), 0.45, colors.HexColor('#bbf7d0')),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('LEFTPADDING', (0, 0), (-1, -1), 6),
                ('RIGHTPADDING', (0, 0), (-1, -1), 6),
                ('TOPPADDING', (0, 0), (-1, -1), 5),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
            ]),
        )

    generated_at = timezone.localtime(timezone.now()).strftime('%d %b %Y, %I:%M %p')
    staff_name = f"{staff.first_name or ''} {staff.last_name or ''}".strip() or staff.username or 'Staff'
    logo_path = Path(settings.BASE_DIR) / 'connect' / 'assets' / 'IIE.png'
    header_left = Image(str(logo_path), width=28 * mm, height=18 * mm) if logo_path.exists() else Paragraph('<b>IIE</b>', cell_style)
    header = Table(
        [[
            header_left,
            [
                Paragraph('Staff Activity Report', title_style),
                Paragraph(escape(f"{staff_name} | {staff.designation or 'Staff'} | {display_branch(staff.branch)} | Generated: {generated_at}"), subtitle_style),
            ],
            '',
        ]],
        colWidths=[35 * mm, None, 35 * mm],
    )
    header.setStyle(TableStyle([
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), 0),
        ('RIGHTPADDING', (0, 0), (-1, -1), 0),
        ('TOPPADDING', (0, 0), (-1, -1), 0),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
    ]))

    batch_rows = []
    for batch in batches:
        completion = _tracking_batch_completion(batch, staff)
        batch_students = Students.objects.filter(assigned_batch=batch)
        total_sessions = completion['total_sessions']
        direct_material_ids = StudyMaterial.objects.filter(uploaded_by=staff, batch=batch).values_list('id', flat=True)
        assigned_material_ids = StudyMaterialAssignment.objects.filter(
            Q(assigned_by=staff) | Q(material__uploaded_by=staff),
            batch=batch,
        ).values_list('material_id', flat=True)
        material_count = StudyMaterial.objects.filter(id__in=set(list(direct_material_ids) + list(assigned_material_ids))).count()
        batch_attendance_total = StudentAttendance.objects.filter(staff=staff, batch=batch).count()
        batch_attendance_present = StudentAttendance.objects.filter(staff=staff, batch=batch, status__iexact='Present').count()
        quiz_count = Quiz.objects.filter(created_by=staff, batch=batch).count()
        student_rows = []
        for student in batch_students.order_by('first_name', 'last_name'):
            student_att_qs = StudentAttendance.objects.filter(staff=staff, batch=batch, student=student)
            student_att_total = student_att_qs.count()
            student_att_present = student_att_qs.filter(status__iexact='Present').count()
            student_sessions_done = Student_Session_Progress.objects.filter(
                student=student,
                session__batch=batch,
            ).filter(
                Q(completed=True) | Q(student_status='completed')
            ).values('session_id').distinct().count()
            student_rows.append({
                'student': f"{student.first_name} {student.last_name or ''}".strip(),
                'student_id': student.student_id,
                'attendance': f"{student_att_present}/{student_att_total} ({round((student_att_present / student_att_total * 100) if student_att_total else 0, 1)}%)",
                'sessions': f"{student_sessions_done}/{total_sessions} ({round((student_sessions_done / total_sessions * 100) if total_sessions else 0, 1)}%)",
                'created': student.created_at.strftime('%d/%m/%Y') if student.created_at else '-',
            })
        batch_rows.append({
            'batch': batch.batch_number,
            'course': batch.course_name.course_name if batch.course_name else '-',
            'meta': f"{batch.course_name.course_name if batch.course_name else '-'} | {batch.course_type or '-'} | {batch.batch_timing or '-'}",
            'students': batch_students.count(),
            'new_students': batch_students.filter(created_at__gte=timezone.now() - timedelta(days=7)).count(),
            'sessions': f"{completion['completed_sessions']}/{completion['total_sessions']}",
            'session_percentage': completion['percentage'],
            'attendance': f"{batch_attendance_present}/{batch_attendance_total}",
            'attendance_percentage': round((batch_attendance_present / batch_attendance_total * 100) if batch_attendance_total else 0, 1),
            'materials': material_count,
            'quizzes': quiz_count,
            'uploads': material_count + quiz_count,
            'student_rows': student_rows,
        })

    weekly_login_records = UserActivity.objects.filter(
        employee=staff,
        user_type='employee',
        login_time__gte=timezone.now() - timedelta(days=7),
    ).order_by('-login_time')

    def format_duration(start, end):
        if not start:
            return '-'
        session_end = end or timezone.now()
        delta = session_end - start
        total_minutes = max(0, int(delta.total_seconds() // 60))
        hours, minutes = divmod(total_minutes, 60)
        if hours:
            return f"{hours}h {minutes}m"
        return f"{minutes}m"

    login_history_rows = [{
        'login_time': item.login_time,
        'logout_time': item.logout_time or 'Still active',
        'last_seen': item.last_seen,
        'duration': format_duration(item.login_time, item.logout_time or item.last_seen),
    } for item in weekly_login_records]

    if is_counselor_report:
        performance_rows = [
            {'metric': 'Students Added', 'count': summary['month_students_added'], 'percentage': summary['performance_graph']['students_added'], 'color': '#0891b2'},
            {'metric': 'Batches Added', 'count': summary['month_batches_added'], 'percentage': summary['performance_graph']['batches_added'], 'color': '#7c3aed'},
            {'metric': 'Students Assigned', 'count': summary['month_students_assigned'], 'percentage': summary['performance_graph']['students_assigned'], 'color': '#059669'},
            {'metric': 'Fee Managed', 'count': summary['month_fee_managed'], 'percentage': summary['performance_graph']['fee_management'], 'color': '#dc2626'},
            {'metric': 'Login Usage', 'count': f"{summary['login_usage_count']}/{summary['login_usage_target']} logins", 'percentage': summary['performance_graph']['login_usage'], 'color': '#ca8a04'},
        ]
    else:
        performance_rows = [
            {'metric': 'Batch Completion', 'count': f"{summary['completed_batch_count']}/{summary['batch_count']} batches", 'percentage': summary['performance_graph']['batch_completion'], 'color': '#059669'},
            {'metric': 'Quiz Upload', 'count': summary['quizzes_created'], 'percentage': summary['performance_graph']['quiz_upload'], 'color': '#7c3aed'},
            {'metric': 'Material Upload', 'count': summary['materials_uploaded'], 'percentage': summary['performance_graph']['material_upload'], 'color': '#0891b2'},
            {'metric': 'Login Usage', 'count': f"{summary['login_usage_count']}/{summary['login_usage_target']} logins", 'percentage': summary['performance_graph']['login_usage'], 'color': '#ca8a04'},
        ]

    story = [
        header,
        Spacer(1, 4),
        make_staff_hero(),
        Spacer(1, 6),
        Table(
            [[
                metric_card('Students', summary['student_count'], 'students handled', '#2563eb'),
                metric_card('Batches', summary['batch_count'], 'active batches', '#059669'),
                metric_card('Materials', summary['materials_uploaded'], 'uploads count', '#0891b2'),
                metric_card('Quizzes', summary['quizzes_created'], 'upload count', '#7c3aed'),
            ]],
            colWidths=[67 * mm, 67 * mm, 67 * mm, 67 * mm],
            style=TableStyle([
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('LEFTPADDING', (0, 0), (-1, -1), 2),
                ('RIGHTPADDING', (0, 0), (-1, -1), 2),
            ]),
        ),
        Paragraph('Performance Graph', section_style),
        Table(
            [[
                make_performance_table(performance_rows),
                donut_chart(summary['activity_score'], '#0f766e', 'Overall Performance', 42 * mm),
            ]],
            colWidths=[190 * mm, 70 * mm],
            style=TableStyle([
                ('BOX', (0, 0), (-1, -1), 0.45, colors.HexColor('#d9e2ec')),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('LEFTPADDING', (0, 0), (-1, -1), 8),
                ('RIGHTPADDING', (0, 0), (-1, -1), 8),
                ('TOPPADDING', (0, 0), (-1, -1), 8),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
            ]),
        ),
        Spacer(1, 5),
        Paragraph('Batch-wise Session and Student Attendance Progress', section_style),
    ]

    for batch_row in batch_rows:
        story.extend([
            make_batch_header(batch_row),
            Table(
                [[
                    metric_card('Materials Uploaded', batch_row['materials'], 'batch upload count', '#0891b2'),
                    metric_card('Quizzes Uploaded', batch_row['quizzes'], 'batch upload count', '#7c3aed'),
                ]],
                colWidths=[130 * mm, 130 * mm],
                style=TableStyle([
                    ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                    ('LEFTPADDING', (0, 0), (-1, -1), 2),
                    ('RIGHTPADDING', (0, 0), (-1, -1), 2),
                    ('TOPPADDING', (0, 0), (-1, -1), 5),
                    ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
                ]),
            ),
            make_performance_table([
                {'metric': 'Batch session progress', 'count': batch_row['sessions'], 'percentage': batch_row['session_percentage'], 'color': '#059669'},
                {'metric': 'Batch attendance progress', 'count': batch_row['attendance'], 'percentage': batch_row['attendance_percentage'], 'color': '#2563eb'},
            ]),
            make_table(
                [('Student', 'student'), ('Student ID', 'student_id'), ('Attendance Progress', 'attendance'), ('Session Progress', 'sessions'), ('Created', 'created')],
                batch_row['student_rows'],
                widths=[62 * mm, 40 * mm, 55 * mm, 55 * mm, 35 * mm],
            ),
            Spacer(1, 6),
        ])

    story.extend([
        Paragraph('Login History (Last 7 Days)', section_style),
        make_table(
            [('Login Time', 'login_time'), ('Logout Time', 'logout_time'), ('Last Seen', 'last_seen'), ('Session Duration', 'duration')],
            login_history_rows,
            widths=[70 * mm, 70 * mm, 70 * mm, 45 * mm],
        ),
    ])

    doc.build(story)
    buffer.seek(0)
    filename = f"staff_activity_{staff_name.replace(' ', '_') or staff.id}.pdf"
    response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
    response['Content-Disposition'] = f'inline; filename="{filename}"'
    return response

# -- ADMIN: STUDY MATERIALS OVERVIEW ------------------------------------------

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_materials_overview(request):
    branch = request.query_params.get('branch')
    staff_id = request.query_params.get('staff_id')  # ADD THIS
    branches = list(Employee.objects.values_list('branch', flat=True).distinct())

    if branch and staff_id:
        # -- LEVEL 3: Materials for a specific staff member --
        try:
            staff = Employee.objects.get(id=staff_id, branch=branch)
        except Employee.DoesNotExist:
            return Response({'branches': branches, 'materials': [], 'staff': None})

        materials = StudyMaterial.objects.filter(
            uploaded_by=staff
        ).select_related('batch', 'uploaded_by').prefetch_related('assignments__batch').order_by('-uploaded_at')

        materials_data = []
        for m in materials:
            materials_data.append(StudyMaterialSerializer(m, context={'request': request}).data)

        return Response({
            'branches': branches,
            'selected_branch': branch,
            'staff': {
                'id': staff.id,
                'first_name': staff.first_name,
                'last_name': staff.last_name or '',
                'designation': staff.designation,
                'email': staff.email,
            },
            'materials': materials_data,
        })

    if branch:
        # -- LEVEL 2: Staff list for a branch --
        staff_members = Employee.objects.filter(
            branch=branch,
            designation__in=['mentor', 'trainer']
        ).order_by('first_name')

        staff_data = []
        for staff in staff_members:
            material_count = StudyMaterial.objects.filter(uploaded_by=staff).count()
            staff_data.append({
                'id': staff.id,
                'first_name': staff.first_name,
                'last_name': staff.last_name or '',
                'designation': staff.designation,
                'email': staff.email,
                'branch': staff.branch,
                'material_count': material_count,
            })

        return Response({
            'branches': branches,
            'selected_branch': branch,
            'staff_members': staff_data,
            'materials': [],
        })

    # -- LEVEL 1: Branch list --
    return Response({
        'branches': branches,
        'selected_branch': None,
        'staff_members': [],
        'materials': [],
    })

# -- ADMIN: SUPPORT REQUESTS OVERVIEW -----------------------------------------

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_support_overview(request):
    support_type = request.query_params.get('type', 'staff')
    status_filter = request.query_params.get('status')

    if support_type == 'staff':
        qs = SupportRequest.objects.select_related('staff').order_by('-created_at')
        if status_filter:
            qs = qs.filter(status=status_filter)
        return Response({'requests': SupportRequestSerializer(qs, many=True).data, 'type': 'staff'})
    elif support_type == 'student':
        qs = StudentSupportRequest.objects.select_related('student').order_by('-created_at')
        if status_filter:
            qs = qs.filter(status=status_filter)
        return Response({'requests': StudentSupportRequestSerializer(qs, many=True).data, 'type': 'student'})
    elif support_type == 'counselor':
        qs = CounselorSupportRequest.objects.select_related('counselor').order_by('-created_at')
        if status_filter:
            qs = qs.filter(status=status_filter)
        return Response({'requests': CounselorSupportRequestSerializer(qs, many=True).data, 'type': 'counselor'})

    return Response({'requests': [], 'type': support_type})


@api_view(['PATCH'])
@permission_classes([IsAuthenticated])
def update_support_status(request, pk):
    support_type = request.data.get('type', 'staff')
    new_status = request.data.get('status')
    try:
        if support_type == 'staff':
            obj = SupportRequest.objects.get(id=pk)
            obj.status = new_status
            obj.save()
            return Response({'success': True, 'status': obj.status})
        elif support_type == 'student':
            obj = StudentSupportRequest.objects.get(id=pk)
            obj.status = new_status
            obj.save()
            return Response({'success': True, 'status': obj.status})
        elif support_type == 'counselor':
            obj = CounselorSupportRequest.objects.get(id=pk)
            obj.status = new_status
            obj.save()
            return Response({'success': True, 'status': obj.status})
    except Exception as e:
        return Response({'error': str(e)}, status=404)


# -- ADMIN: SEPARATE SUPPORT ENDPOINTS ----------------------------------------

class AdminStaffSupportListView(APIView):
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        requests = SupportRequest.objects.select_related('staff').order_by('-created_at')
        data = []
        for req in requests:
            try:
                emp = Employee.objects.get(user=req.staff)
                if emp.designation.lower() != 'counselor':
                    data.append({
                        'id': req.id,
                        'staff_name': f"{emp.first_name} {emp.last_name or ''}",
                        'message': req.message,
                        'status': req.status,
                        'created_at': req.created_at,
                    })
            except Employee.DoesNotExist:
                data.append({
                    'id': req.id,
                    'staff_name': req.staff.username,
                    'message': req.message,
                    'status': req.status,
                    'created_at': req.created_at,
                })
        return Response({'results': data})
    
    def patch(self, request, id):
        try:
            req = SupportRequest.objects.get(id=id)
            req.status = request.data.get('status', req.status)
            req.save()
            return Response({'success': True, 'status': req.status})
        except SupportRequest.DoesNotExist:
            return Response({'error': 'Not found'}, status=404)


class AdminStaffSupportHistoryView(APIView):
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        requests = SupportRequest.objects.filter(status='resolved').select_related('staff').order_by('-created_at')
        data = []
        for req in requests:
            try:
                emp = Employee.objects.get(user=req.staff)
                if emp.designation.lower() != 'counselor':
                    data.append({
                        'id': req.id,
                        'staff_name': f"{emp.first_name} {emp.last_name or ''}",
                        'message': req.message,
                        'status': req.status,
                        'created_at': req.created_at,
                    })
            except Employee.DoesNotExist:
                pass
        return Response({'results': data})


class AdminCounselorSupportListView(APIView):
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        requests = CounselorSupportRequest.objects.select_related('counselor').order_by('-created_at')
        data = []
        for req in requests:
            try:
                emp = Employee.objects.get(user=req.counselor)
                data.append({
                    'id': req.id,
                    'counselor_name': f"{emp.first_name} {emp.last_name or ''}",
                    'message': req.message,
                    'status': req.status,
                    'created_at': req.created_at,
                })
            except Employee.DoesNotExist:
                data.append({
                    'id': req.id,
                    'counselor_name': req.counselor.username,
                    'message': req.message,
                    'status': req.status,
                    'created_at': req.created_at,
                })
        return Response({'results': data})
    
    def patch(self, request, id):
        try:
            req = CounselorSupportRequest.objects.get(id=id)
            req.status = request.data.get('status', req.status)
            req.save()
            return Response({'success': True, 'status': req.status})
        except CounselorSupportRequest.DoesNotExist:
            return Response({'error': 'Not found'}, status=404)


class AdminCounselorSupportHistoryView(APIView):
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        requests = CounselorSupportRequest.objects.filter(status='resolved').select_related('counselor').order_by('-created_at')
        data = []
        for req in requests:
            try:
                emp = Employee.objects.get(user=req.counselor)
                data.append({
                    'id': req.id,
                    'counselor_name': f"{emp.first_name} {emp.last_name or ''}",
                    'message': req.message,
                    'status': req.status,
                    'created_at': req.created_at,
                })
            except Employee.DoesNotExist:
                pass
        return Response({'results': data})


class AdminStudentSupportListView(APIView):
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        requests = StudentSupportRequest.objects.select_related('student').order_by('-created_at')
        data = []
        for req in requests:
            try:
                student = Students.objects.get(user=req.student)
                data.append({
                    'id': req.id,
                    'student_name': f"{student.first_name} {student.last_name or ''}",
                    'student_id': student.student_id,
                    'message': req.message,
                    'status': req.status,
                    'created_at': req.created_at,
                })
            except Students.DoesNotExist:
                data.append({
                    'id': req.id,
                    'student_name': req.student.username,
                    'message': req.message,
                    'status': req.status,
                    'created_at': req.created_at,
                })
        return Response({'results': data})
    
    def patch(self, request, id):
        try:
            req = StudentSupportRequest.objects.get(id=id)
            req.status = request.data.get('status', req.status)
            req.save()
            return Response({'success': True, 'status': req.status})
        except StudentSupportRequest.DoesNotExist:
            return Response({'error': 'Not found'}, status=404)


class AdminStudentSupportHistoryView(APIView):
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        requests = StudentSupportRequest.objects.filter(status='resolved').select_related('student').order_by('-created_at')
        data = []
        for req in requests:
            try:
                student = Students.objects.get(user=req.student)
                data.append({
                    'id': req.id,
                    'student_name': f"{student.first_name} {student.last_name or ''}",
                    'student_id': student.student_id,
                    'message': req.message,
                    'status': req.status,
                    'created_at': req.created_at,
                })
            except Students.DoesNotExist:
                pass
        return Response({'results': data})


# -- ADMIN LEAVE MANAGEMENT ----------------------------------------------------

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_staff_leave(request):
    status_filter = request.query_params.get('status', '')
    history = request.query_params.get('history', 'false') == 'true'
    qs = StaffLeaveRequest.objects.select_related('staff').exclude(
        staff__designation__iexact='counselor'
    ).order_by('-applied_at')
    if status_filter:
        qs = qs.filter(status=status_filter)
    elif not history:
        qs = qs.filter(status='Pending')
    if history:
        qs = StaffLeaveRequest.objects.select_related('staff').exclude(
            staff__designation__iexact='counselor'
        ).exclude(status='Pending').order_by('-applied_at')
    data = []
    for l in qs:
        data.append({
            'id': l.id,
            'staff_name': f"{l.staff.first_name} {l.staff.last_name or ''}",
            'staff_designation': l.staff.designation,
            'staff_branch': l.staff.branch,
            'leave_type': l.leave_type,
            'start_date': str(l.start_date),
            'end_date': str(l.end_date),
            'no_of_days': l.no_of_days,
            'contact_info': l.contact_info or '',
            'reason': l.reason,
            'status': l.status,
            'applied_at': l.applied_at.strftime('%d %b %Y, %H:%M') if l.applied_at else '',
        })
    return Response(data)


@api_view(['PATCH'])
@permission_classes([IsAuthenticated])
def admin_process_staff_leave(request, pk):
    try:
        leave = StaffLeaveRequest.objects.get(id=pk)
        action = request.data.get('action')
        leave.status = 'Approved' if action == 'Accept' else 'Rejected'
        leave.save()
        return Response({'success': True, 'status': leave.status})
    except StaffLeaveRequest.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_counselor_leave(request):
    history = request.query_params.get('history', 'false') == 'true'
    qs = CounselorLeaveRequest.objects.select_related('counselor').order_by('-applied_at')
    if history:
        qs = qs.exclude(status='Pending')
    else:
        qs = qs.filter(status='Pending')
    data = []
    for l in qs:
        data.append({
            'id': l.id,
            'staff_name': f"{l.counselor.first_name} {l.counselor.last_name or ''}",
            'staff_designation': l.counselor.designation,
            'staff_branch': l.counselor.branch,
            'leave_type': l.leave_type,
            'start_date': str(l.start_date),
            'end_date': str(l.end_date),
            'no_of_days': l.no_of_days,
            'contact_info': l.contact_info or '',
            'reason': l.reason,
            'status': l.status,
            'applied_at': l.applied_at.strftime('%d %b %Y, %H:%M') if l.applied_at else '',
        })
    return Response(data)


@api_view(['PATCH'])
@permission_classes([IsAuthenticated])
def admin_process_counselor_leave(request, pk):
    try:
        leave = CounselorLeaveRequest.objects.get(id=pk)
        action = request.data.get('action')
        leave.status = 'Approved' if action == 'Accept' else 'Rejected'
        leave.save()
        return Response({'success': True, 'status': leave.status})
    except CounselorLeaveRequest.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)


# -- PDF EXTRACTION -----------------------------------------------------------

def extract_pdf_text(logsheet_path):
    """
    Extract text from PDF using available libraries.
    Tries pdfplumber first, falls back to PyPDF2.
    """
    full_text = ""
    
    # Try pdfplumber first (better for tables)
    try:
        import pdfplumber
        with pdfplumber.open(logsheet_path) as pdf:
            for page in pdf.pages:
                # Try to extract tables first
                tables = page.extract_tables()
                if tables:
                    for table in tables:
                        for row in table:
                            full_text += " | ".join([str(cell or '') for cell in row]) + "\n"
                
                # Also extract regular text
                page_text = page.extract_text()
                if page_text:
                    full_text += page_text + "\n"
        
        if full_text.strip():
            return full_text
    except ImportError:
        pass  # pdfplumber not installed, will use PyPDF2
    except Exception as e:
        print(f"[EXTRACT] pdfplumber extraction failed: {e}, trying PyPDF2...")
    
    # Fallback to PyPDF2
    try:
        pdf_reader = PyPDF2.PdfReader(logsheet_path)
        for page in pdf_reader.pages:
            page_text = page.extract_text() or ''
            full_text += page_text + "\n"
    except Exception as e:
        print(f"[EXTRACT] PyPDF2 extraction failed: {e}")
        return ""
    
    return full_text


def is_valid_pdf_file(logsheet_path):
    try:
        with open(logsheet_path, 'rb') as pdf_file:
            return pdf_file.read(5) == b'%PDF-'
    except OSError:
        return False


def _clean_session_text(value):
    return re.sub(r'\s+', ' ', str(value or '')).strip(' |:-\t\r\n')


def _safe_session_payload(session_number, content):
    cleaned_content = _clean_session_text(content) or 'Course Content'
    title_prefix = f"Session {session_number}: "
    title_max_length = CourseSession._meta.get_field('title').max_length or 255
    available_title_length = max(title_max_length - len(title_prefix), 20)

    title_candidate = cleaned_content
    if len(title_candidate) > available_title_length:
        title_candidate = title_candidate[:available_title_length - 3].rstrip() + '...'

    return {
        'session_number': session_number,
        'title': f"{title_prefix}{title_candidate}",
        'topics': cleaned_content[:1000],
    }


def _assign_pdf_session_number(original_number, used_numbers, last_number):
    """
    Keep the PDF number when it is usable, otherwise preserve the extra PDF row
    with the next available number. This avoids silently dropping duplicate or
    repeated session labels from logsheets.
    """
    if original_number not in used_numbers and original_number > 0:
        return original_number

    next_number = max(last_number, max(used_numbers) if used_numbers else 0) + 1
    while next_number in used_numbers:
        next_number += 1
    return next_number


def _parse_session_markers_anywhere(full_text, batch_id=None, debug=False):
    marker_pattern = re.compile(
        r'\b(?:Session|Module|Day|Chapter|Unit|Lesson)\s*[-:\s]*([0-9]{1,3})\s*[:.)-]?\s*',
        re.IGNORECASE
    )
    matches = list(marker_pattern.finditer(full_text))
    if not matches:
        return []

    sessions = []
    used_numbers = set()
    last_number = 0

    for index, match in enumerate(matches):
        original_number = int(match.group(1))
        content_start = match.end()
        content_end = matches[index + 1].start() if index + 1 < len(matches) else len(full_text)
        content = full_text[content_start:content_end]

        content = re.sub(r'\bPage\s+\d+\b', ' ', content, flags=re.IGNORECASE)
        content = _clean_session_text(content)
        if not content or content.lower() in {'session', 'module', 'day', 'chapter', 'unit', 'lesson'}:
            continue

        session_number = _assign_pdf_session_number(original_number, used_numbers, last_number)
        used_numbers.add(session_number)
        last_number = max(last_number, session_number)
        sessions.append(_safe_session_payload(session_number, content))

    if debug:
        print(
            f"[EXTRACT] Batch {batch_id}: Detected {len(matches)} session markers, "
            f"kept {len(sessions)} sessions"
        )

    return sessions


def _parse_numbered_rows(lines, batch_id=None, debug=False):
    sessions = []
    used_numbers = set()
    last_number = 0
    numbered_patterns = [
        re.compile(r'^\s*(\d{1,3})\s*\|\s*(.+)$'),
        re.compile(r'^\s*(\d{1,3})\s*[-.:)]\s+(.+)$'),
        re.compile(r'^\s*(\d{1,3})\s{2,}(.+)$'),
    ]

    for line in lines:
        for pattern in numbered_patterns:
            match = pattern.match(line)
            if not match:
                continue
            original_number = int(match.group(1))
            content = _clean_session_text(match.group(2))
            if len(content) < 3 or content.lower().startswith(('page', 'table', 'figure', 's.no')):
                break
            session_number = _assign_pdf_session_number(original_number, used_numbers, last_number)
            used_numbers.add(session_number)
            last_number = max(last_number, session_number)
            sessions.append(_safe_session_payload(session_number, content))
            break

    if debug:
        print(f"[EXTRACT] Batch {batch_id}: Numbered-row fallback kept {len(sessions)} sessions")

    return sessions


def _normalize_session_match_text(value):
    value = re.sub(r'^session\s+\d+\s*:\s*', '', str(value or ''), flags=re.IGNORECASE)
    value = re.sub(r'[^a-z0-9]+', ' ', value.lower())
    return re.sub(r'\s+', ' ', value).strip()


def _session_title_similarity(left, right):
    from difflib import SequenceMatcher

    left_norm = _normalize_session_match_text(left)
    right_norm = _normalize_session_match_text(right)
    if not left_norm or not right_norm:
        return 0

    if left_norm in right_norm or right_norm in left_norm:
        return 1

    return SequenceMatcher(None, left_norm[:220], right_norm[:220]).ratio()


def parse_sessions_from_text(full_text, batch_id=None, debug=False):
    """
    Parse session information from extracted PDF text.
    Supports multiple formats: Session 1, Session-1, Module 1, Day 1, 1. Topic, tables, etc.
    
    Args:
        full_text: Raw text extracted from PDF
        batch_id: Batch ID for debug logging
        debug: Enable debug logging
        
    Returns:
        List of dicts with session_number, title, topics
    """
    if not full_text.strip():
        return []
    
    # Log first 2000 characters for debugging
    if debug:
        preview = full_text[:2000].replace('\n', '\\n')
        print(f"[EXTRACT] Batch {batch_id}: First 2000 chars:\n{preview}\n")
    
    normalized_text = re.sub(r'\r\n?', '\n', full_text)
    sessions = _parse_session_markers_anywhere(normalized_text, batch_id=batch_id, debug=debug)
    if sessions:
        return sorted(sessions, key=lambda item: item['session_number'])

    if debug:
        print(f"[EXTRACT] Batch {batch_id}: No explicit session markers found, using numbered-row fallback")

    lines = [line.strip() for line in normalized_text.splitlines() if line.strip()]
    return sorted(_parse_numbered_rows(lines, batch_id=batch_id, debug=debug), key=lambda item: item['session_number'])


def extract_sessions_from_logsheet(batch, debug=False, prefer_course_logsheet=False):
    """
    Extract sessions from batch logsheet.
    If batch logsheet file is missing, fallback to course logsheet.
    """
    try:
        candidates = []

        def add_batch_logsheet_candidate():
            if batch.course_logsheet:
                batch_path = getattr(batch.course_logsheet, "path", None)

                if debug:
                    print(f"[EXTRACT] Batch {batch.id}: Checking batch.course_logsheet")
                    print(f"[EXTRACT] Batch {batch.id}: Batch field={batch.course_logsheet.name}")
                    print(f"[EXTRACT] Batch {batch.id}: Batch path={batch_path}")
                    print(f"[EXTRACT] Batch {batch.id}: Batch file exists={os.path.exists(batch_path) if batch_path else False}")

                if batch_path and os.path.exists(batch_path):
                    candidates.append((batch.course_logsheet, "batch.course_logsheet"))

        def add_course_logsheet_candidate():
            if not (batch.course_name and batch.course_name.course_logsheet):
                return
            course_path = getattr(batch.course_name.course_logsheet, "path", None)

            if debug:
                print(f"[EXTRACT] Batch {batch.id}: Checking course.course_logsheet fallback")
                print(f"[EXTRACT] Batch {batch.id}: Course field={batch.course_name.course_logsheet.name}")
                print(f"[EXTRACT] Batch {batch.id}: Course path={course_path}")
                print(f"[EXTRACT] Batch {batch.id}: Course file exists={os.path.exists(course_path) if course_path else False}")

            if course_path and os.path.exists(course_path):
                candidates.append((batch.course_name.course_logsheet, "course.course_logsheet"))

        if prefer_course_logsheet:
            add_course_logsheet_candidate()
            add_batch_logsheet_candidate()
        else:
            add_batch_logsheet_candidate()
            add_course_logsheet_candidate()

        if not candidates:
            if debug:
                print(f"[EXTRACT] Batch {batch.id}: No valid existing logsheet file found")
            batch._logsheet_extract_error = 'No uploaded logsheet file was found on disk. Please upload the PDF logsheet again.'
            return []

        last_error = ''
        for pdf_field, source in candidates:
            logsheet_path = getattr(pdf_field, "path", None)

            if debug:
                print(f"[EXTRACT] Batch {batch.id}: Selected Source={source}")
                print(f"[EXTRACT] Batch {batch.id}: Selected Field name={pdf_field.name}")
                print(f"[EXTRACT] Batch {batch.id}: Selected Path={logsheet_path}")

            if not logsheet_path or not is_valid_pdf_file(logsheet_path):
                last_error = f'{pdf_field.name} is not a valid PDF file. Please re-upload the original PDF logsheet.'
                if debug:
                    print(f"[EXTRACT] Batch {batch.id}: {last_error}")
                continue

            full_text = extract_pdf_text(logsheet_path)

            if not full_text.strip():
                last_error = f'Could not read text from {pdf_field.name}. The PDF may be scanned, image-only, or corrupted.'
                if debug:
                    print(f"[EXTRACT] Batch {batch.id}: No text extracted from PDF")
                continue

            if debug:
                print(f"[EXTRACT] Batch {batch.id}: Extracted {len(full_text)} chars from PDF")
                print(f"[EXTRACT] Batch {batch.id}: First 1000 chars:\n{full_text[:1000]}")

            sessions_data = parse_sessions_from_text(full_text, batch_id=batch.id, debug=debug)
            if sessions_data:
                return sessions_data

            last_error = f'No session rows could be detected in {pdf_field.name}. Please check the logsheet format.'

        batch._logsheet_extract_error = last_error or 'Could not extract sessions from the uploaded logsheet.'
        return []

    except Exception as e:
        print(f"[EXTRACT] ERROR Batch {batch.id}: {type(e).__name__}: {str(e)}")
        import traceback
        traceback.print_exc()
        return []


def sync_missing_sessions_from_logsheet(batch, debug=False, prefer_course_logsheet=False):
    """
    Non-destructively align DB sessions with the PDF logsheet.
    Existing progress rows are preserved by updating CourseSession rows in-place.
    """
    sessions_data = extract_sessions_from_logsheet(
        batch,
        debug=debug,
        prefer_course_logsheet=prefer_course_logsheet
    )
    if not sessions_data:
        return 0

    title_max_length = CourseSession._meta.get_field('title').max_length or 255
    for item in sessions_data:
        title = item.get('title', '')
        if len(title) > title_max_length:
            title = title[:title_max_length - 3].rstrip() + '...'
        item['title'] = title

    existing_sessions = list(CourseSession.objects.filter(batch=batch).order_by('session_number'))
    if not existing_sessions:
        new_sessions = [
            CourseSession.objects.create(
                batch=batch,
                session_number=item.get('session_number'),
                title=item.get('title', ''),
                topics=item.get('topics', ''),
                staff_completed=False,
                session_enabled=True
            )
            for item in sessions_data
        ]
        _ensure_student_session_rows(batch, new_sessions)
        return len(new_sessions)

    parsed_count = len(sessions_data)
    existing_count = len(existing_sessions)

    if parsed_count <= existing_count:
        existing_by_number = {session.session_number: session for session in existing_sessions}
        existing_numbers = set(existing_by_number.keys())
        new_sessions = []
        for item in sessions_data:
            session_number = item.get('session_number')
            existing_session = existing_by_number.get(session_number)
            if existing_session:
                update_fields = []
                if existing_session.title != item.get('title', ''):
                    existing_session.title = item.get('title', '')
                    update_fields.append('title')
                if (existing_session.topics or '') != (item.get('topics', '') or ''):
                    existing_session.topics = item.get('topics', '')
                    update_fields.append('topics')
                if update_fields:
                    existing_session.save(update_fields=update_fields)
                continue
            session = CourseSession.objects.create(
                batch=batch,
                session_number=session_number,
                title=item.get('title', ''),
                topics=item.get('topics', ''),
                staff_completed=False,
                session_enabled=True
            )
            new_sessions.append(session)
            existing_numbers.add(session_number)

        if new_sessions:
            _ensure_student_session_rows(batch, new_sessions)
        return len(new_sessions)

    new_sessions = []
    updated_count = 0

    with transaction.atomic():
        temp_sessions = list(CourseSession.objects.filter(batch=batch).order_by('session_number'))
        for session in temp_sessions:
            session.session_number = -session.id
            session.save(update_fields=['session_number'])

        unmatched_sessions = temp_sessions[:]
        used_existing = set()

        for item in sessions_data:
            best_session = None
            best_score = 0
            for session in unmatched_sessions:
                if session.id in used_existing:
                    continue
                score = _session_title_similarity(item.get('title', ''), session.title)
                if score > best_score:
                    best_score = score
                    best_session = session

            if best_session and best_score >= 0.55:
                best_session.session_number = item.get('session_number')
                best_session.title = item.get('title', '')
                best_session.topics = item.get('topics', '')
                best_session.save(update_fields=['session_number', 'title', 'topics'])
                used_existing.add(best_session.id)
                updated_count += 1
                continue

            session = CourseSession.objects.create(
                batch=batch,
                session_number=item.get('session_number'),
                title=item.get('title', ''),
                topics=item.get('topics', ''),
                staff_completed=False,
                session_enabled=True
            )
            new_sessions.append(session)

        next_number = parsed_count + 1
        for session in unmatched_sessions:
            if session.id in used_existing:
                continue
            while CourseSession.objects.filter(batch=batch, session_number=next_number).exists():
                next_number += 1
            session.session_number = next_number
            session.save(update_fields=['session_number'])
            next_number += 1

    if new_sessions:
        _ensure_student_session_rows(batch, new_sessions)

    if debug:
        print(
            f"[SESSION_SYNC] Batch {batch.id}: added {len(new_sessions)} missing sessions, "
            f"realigned {updated_count} existing sessions"
        )

    return len(new_sessions)


def _ensure_student_session_rows(batch, sessions):
    students = Students.objects.filter(assigned_batch=batch)
    for student in students:
        for session in sessions:
            Student_Session_Progress.objects.get_or_create(
                student=student,
                session=session,
                defaults={
                    'completed': False,
                    'staff_completed': False,
                    'student_status': 'not_started',
                    'has_doubt': False,
                    'doubt_resolved': False,
                }
            )
            StudentSessionStatus.objects.get_or_create(
                student=student,
                session=session,
                defaults={
                    'staff_completed': False,
                    'student_status': 'pending',
                }
            )


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def extract_and_store_sessions(request, batch_id):
    """
    Extract sessions from batch/course logsheet and store in database.
    Automatically uses course_logsheet if batch_logsheet is empty.
    Existing staff completion ticks, student session progress, doubts, and
    notifications are preserved.
    """
    try:
        print(f"\n{'='*60}")
        print(f"[SESSION_EXTRACT] START: batch_id={batch_id}")
        print(f"{'='*60}")
        
        batch = Batches.objects.get(id=batch_id)
        print(f"[SESSION_EXTRACT] Batch {batch.id}: {batch.batch_number}")
        print(f"[SESSION_EXTRACT] Batch.course_logsheet: {batch.course_logsheet.name if batch.course_logsheet else 'EMPTY'}")
        print(f"[SESSION_EXTRACT] Course.course_logsheet: {batch.course_name.course_logsheet.name if batch.course_name and batch.course_name.course_logsheet else 'EMPTY'}")
        
        # Validate that a logsheet exists (batch or course)
        if not batch.course_logsheet and (not batch.course_name or not batch.course_name.course_logsheet):
            return Response({
                'error': 'No logsheet uploaded for this batch or course. Please upload a PDF logsheet first.'
            }, status=400)

        sessions_data = extract_sessions_from_logsheet(batch, debug=True)
        
        if not sessions_data:
            specific_error = getattr(batch, '_logsheet_extract_error', '')
            error_msg = (
                (specific_error + ' ') if specific_error else ''
            ) + (
                'Could not extract any sessions from the PDF logsheet. '
                'Supported formats: "Session 1", "Module-1", "Day 1", "1. Topic", "1) Topic", '
                '"1 - Topic", "1: Topic", "Session 61: WEB ARCHITECTURE", or table format with numbers. '
                'Check deployment logs for first 2000 characters of extracted PDF text.'
            )
            print(f"[SESSION_EXTRACT] ERROR: {error_msg}")
            return Response({'error': error_msg}, status=400)
        
        before_count = CourseSession.objects.filter(batch=batch).count()
        added_count = sync_missing_sessions_from_logsheet(batch, debug=True)
        sessions = CourseSession.objects.filter(batch=batch).order_by('session_number')
        students_count = Students.objects.filter(assigned_batch=batch).count()
        
        print(f"\n[SESSION_EXTRACT] COMPLETE:")
        print(f"  ? PDF sessions detected: {len(sessions_data)}")
        print(f"  ? Existing sessions before sync: {before_count}")
        print(f"  ? Missing sessions added: {added_count}")
        print(f"  ? Total sessions now: {sessions.count()}")
        print(f"  ? Staff/student progress preserved")
        print(f"{'='*60}\n")
        
        return Response({
            'success': True,
            'message': f'Session sheet synced safely. {added_count} missing sessions added; existing staff completion ticks were preserved.',
            'sessions': CourseSessionSerializer(sessions, many=True).data,
            'total_sessions': sessions.count(),
            'students_updated': students_count,
            'preserved_progress': True,
        })
        
    except Batches.DoesNotExist:
        print(f"[SESSION_EXTRACT] ERROR: Batch {batch_id} not found")
        return Response({'error': 'Batch not found.'}, status=404)
    except Exception as e:
        print(f"[SESSION_EXTRACT] ERROR: {type(e).__name__}: {str(e)}")
        import traceback
        traceback.print_exc()
        return Response({'error': str(e)}, status=500)

# -- ADMIN: MENTORS LIST -------------------------------------------------------

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_view_mentors(request):
    designation = request.query_params.get('designation')
    branch = request.query_params.get('branch')
    qs = Employee.objects.all().order_by('first_name')
    if designation:
        qs = qs.filter(designation__iexact=designation)
    if branch:
        qs = qs.filter(branch=branch)
    return Response(EmployeeSerializer(qs, many=True).data)


# -- STAFF STUDENT LEAVE MANAGEMENT -------------------------------------------

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def staff_student_leave_requests(request):
    """Staff views leave requests from their assigned students"""
    try:
        staff = Employee.objects.get(user=request.user)
        leaves = StudentLeaveApplication.objects.filter(
            assigned_staff=staff
        ).order_by('-applied_at')
        
        data = []
        for l in leaves:
            data.append({
                'id': l.id,
                'student_name': f"{l.student.first_name} {l.student.last_name or ''}",
                'student_id': l.student.student_id,
                'leave_type': l.leave_type,
                'start_date': l.start_date.strftime('%Y-%m-%d'),
                'end_date': l.end_date.strftime('%Y-%m-%d'),
                'number_of_days': l.number_of_days,
                'reason': l.reason,
                'contact_info': l.contact_info,
                'status': l.status,
                'applied_at': l.applied_at,
            })
        return Response({'results': data})
    except Employee.DoesNotExist:
        return Response({'error': 'Staff not found'}, status=404)


@api_view(['PATCH'])
@permission_classes([IsAuthenticated])
def staff_process_student_leave(request, pk):
    """Staff approves or rejects student leave request"""
    try:
        staff = Employee.objects.get(user=request.user)
        leave = StudentLeaveApplication.objects.get(id=pk, assigned_staff=staff)
        
        new_status = request.data.get('status')
        remarks = request.data.get('remarks', '')
        
        if new_status not in ['approved', 'rejected']:
            return Response({'error': 'Invalid status'}, status=400)
        
        leave.status = new_status
        leave.staff_remarks = remarks
        leave.processed_at = timezone.now()
        leave.save()
        
        # Update student's used leave days if approved
        if new_status == 'approved':
            leave.student.update_used_leave_days()
        
        return Response({'success': True, 'status': leave.status})
    except Employee.DoesNotExist:
        return Response({'error': 'Staff not found'}, status=404)
    except StudentLeaveApplication.DoesNotExist:
        return Response({'error': 'Leave request not found'}, status=404)



@api_view(['GET'])
@permission_classes([IsAuthenticated])
def student_quizzes(request):
    """Get quizzes available for the logged-in student with ALL options"""
    try:
        student = Students.objects.get(user=request.user)
        if not student.assigned_batch:
            return Response({'results': []})
        
        quizzes = Quiz.objects.filter(
            batch=student.assigned_batch,
            is_published=True
        ).distinct().order_by('-created_at')
        
        data = []
        for quiz in quizzes:
            attempts = QuizAttempt.objects.filter(student=student, quiz=quiz)
            completed_attempts = attempts.filter(is_completed=True)
            best_score = attempts.filter(is_completed=True).order_by('-percentage').first()
            is_completed = completed_attempts.exists()
            
            questions = quiz.questions.all().order_by('question_number')
            questions_data = []
            total_marks = 0
            for q in questions:
                questions_data.append({
                    'id': q.id,
                    'question_text': q.question_text,
                    'option_a': q.option_a or '',
                    'option_b': q.option_b or '',
                    'option_c': q.option_c or '',
                    'option_d': q.option_d or '',
                    'marks': q.marks,
                })
                total_marks += q.marks
            
            # If no questions, use default
            if total_marks == 0 and questions.count() > 0:
                total_marks = questions.count()
            
            data.append({
                'id': quiz.id,
                'title': quiz.title,
                'description': quiz.description,
                'total_questions': questions.count(),
                'total_marks': total_marks,  # Send calculated total_marks
                'duration_minutes': quiz.duration_minutes,
                'passing_marks': quiz.passing_marks,
                'max_attempts': quiz.max_attempts,
                'user_attempts': attempts.count(),
                'best_score': best_score.percentage if best_score else None,
                'status': 'completed' if is_completed else 'available',
                'last_attempt_id': best_score.id if best_score else None,
                'completed_at': best_score.submitted_at if best_score else None,
                'questions': questions_data
            })
        
        return Response({'results': data})
    except Students.DoesNotExist:
        return Response({'error': 'Student not found'}, status=404)

@api_view(['POST'])
@permission_classes([IsAuthenticated])
def student_take_quiz(request, quiz_id):
    """Student submits quiz answers"""
    try:
        student = Students.objects.get(user=request.user)
        quiz = Quiz.objects.get(id=quiz_id, is_published=True)
        
        assigned_to_batch = bool(student.assigned_batch and quiz.batch_id == student.assigned_batch_id)

        if not assigned_to_batch:
            return Response({'error': 'This quiz is not assigned to your batch'}, status=400)
        
        # Student quizzes are one-time only in the mobile app.
        existing_attempts = QuizAttempt.objects.filter(student=student, quiz=quiz).count()
        if existing_attempts > 0:
            return Response({'error': 'You have already attended this quiz. Only one attempt allowed.'}, status=400)

        # Check attempt limit
        if existing_attempts >= quiz.max_attempts and quiz.max_attempts > 0:
            return Response({'error': f'Maximum attempts ({quiz.max_attempts}) reached'}, status=400)
        
        answers_data = request.data.get('answers', {})
        questions = quiz.questions.all()
        
        # Debug prints
        print("=" * 60)
        print(f"QUIZ: {quiz.title}")
        print(f"Total questions in DB: {questions.count()}")
        
        # Calculate total_marks from questions
        total_marks = 0
        for q in questions:
            total_marks += q.marks
            print(f"Question {q.id}: marks = {q.marks}")
        
        print(f"Calculated TOTAL_MARKS: {total_marks}")
        print(f"Quiz.total_marks in DB: {quiz.total_marks}")
        print(f"Received answers: {answers_data}")
        print("=" * 60)
        
        # If total_marks is still 0, set default
        if total_marks == 0:
            total_marks = questions.count()  # Each question 1 mark
            print(f"WARNING: total_marks was 0, set to {total_marks}")
        
        # Calculate score
        score = 0
        correct_count = 0
        
        for question in questions:
            selected = answers_data.get(str(question.id), '')
            
            # Normalize for comparison
            selected_normalized = str(selected).strip().upper()
            correct_normalized = str(question.correct_answer).strip().upper()
            
            is_correct = selected_normalized == correct_normalized
            
            if is_correct:
                score += question.marks
                correct_count += 1
                print(f"Q{question.id}: CORRECT (+{question.marks})")
            else:
                print(f"Q{question.id}: WRONG (selected={selected_normalized}, correct={correct_normalized})")
        
        # Calculate percentage
        if total_marks > 0:
            percentage = round((score / total_marks) * 100, 1)
        else:
            percentage = 0
            
        is_passed = percentage >= quiz.passing_marks
        
        print(f"\nFINAL: score={score}, total_marks={total_marks}, percentage={percentage}%")
        print("=" * 60)
        
        # Save attempt
        attempt = QuizAttempt.objects.create(
            quiz=quiz,
            student=student,
            attempt_number=existing_attempts + 1,
            score=score,
            percentage=percentage,
            is_passed=is_passed,
            is_completed=True,
            submitted_at=timezone.now()
        )
        
        # Also update quiz.total_marks for future reference
        if quiz.total_marks != total_marks:
            quiz.total_marks = total_marks
            quiz.total_questions = questions.count()
            quiz.save()
            print(f"Updated quiz.total_marks to {total_marks}")
        
        # Save answers
        for question in questions:
            selected = answers_data.get(str(question.id), '')
            selected_normalized = str(selected).strip().upper()
            correct_normalized = str(question.correct_answer).strip().upper()
            is_correct = selected_normalized == correct_normalized
            marks_obtained = question.marks if is_correct else 0
            
            QuizAnswer.objects.create(
                attempt=attempt,
                question=question,
                selected_answer=selected,
                is_correct=is_correct,
                marks_obtained=marks_obtained
            )

        quiz_recipients = []
        if quiz.created_by and quiz.created_by.user:
            quiz_recipients.append(quiz.created_by.user)
        if quiz.batch and quiz.batch.faculty and quiz.batch.faculty.user:
            quiz_recipients.append(quiz.batch.faculty.user)
        if student.assigned_staff and student.assigned_staff.user:
            quiz_recipients.append(student.assigned_staff.user)
        for recipient in set(quiz_recipients):
            _queue_user_notification(
                recipient,
                request.user,
                'quiz_result',
                'Quiz Result Submitted',
                f"Quiz Result: {student.first_name} {student.last_name or ''} completed {quiz.title} with {percentage}%.",
                False,
            )
        
        return Response({
            'attempt_id': attempt.id,
            'score': score,
            'total_marks': total_marks,
            'percentage': percentage,
            'is_passed': is_passed,
            'correct_count': correct_count,
            'total_questions': questions.count()
        })
        
    except Students.DoesNotExist:
        return Response({'error': 'Student not found'}, status=404)
    except Quiz.DoesNotExist:
        return Response({'error': 'Quiz not found'}, status=404)
    except Exception as e:
        print(f"Error: {e}")
        import traceback
        traceback.print_exc()
        return Response({'error': str(e)}, status=400)

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def quiz_result_details(request, attempt_id):
    """Get detailed quiz result with answers"""
    try:
        attempt = QuizAttempt.objects.select_related(
            'quiz',
            'quiz__batch',
            'quiz__created_by',
            'student',
            'student__user',
            'student__assigned_staff',
        ).get(id=attempt_id)
        
        # Check permission
        has_access = attempt.student.user == request.user or request.user.is_superuser or request.user.is_staff
        if not has_access:
            try:
                employee = Employee.objects.get(user=request.user)
                has_access = (
                    attempt.quiz.created_by_id == employee.id or
                    (attempt.quiz.batch and attempt.quiz.batch.faculty_id == employee.id) or
                    attempt.student.assigned_staff_id == employee.id
                )
            except Employee.DoesNotExist:
                has_access = False

        if not has_access:
            return Response({'error': 'Access denied'}, status=403)
        
        answers = QuizAnswer.objects.filter(attempt=attempt).select_related('question').order_by('question__question_number')
        
        questions_details = []
        for answer in answers:
            q = answer.question
            # Get selected option text
            selected_option_text = ''
            if answer.selected_answer:
                if answer.selected_answer.upper() == 'A':
                    selected_option_text = q.option_a
                elif answer.selected_answer.upper() == 'B':
                    selected_option_text = q.option_b
                elif answer.selected_answer.upper() == 'C':
                    selected_option_text = q.option_c
                elif answer.selected_answer.upper() == 'D':
                    selected_option_text = q.option_d
            
            # Get correct option text
            correct_option_text = ''
            if q.correct_answer.upper() == 'A':
                correct_option_text = q.option_a
            elif q.correct_answer.upper() == 'B':
                correct_option_text = q.option_b
            elif q.correct_answer.upper() == 'C':
                correct_option_text = q.option_c
            elif q.correct_answer.upper() == 'D':
                correct_option_text = q.option_d
            
            questions_details.append({
                'question_id': q.id,
                'question_number': q.question_number,
                'question_text': q.question_text,
                'selected_answer': answer.selected_answer or '-',
                'selected_option_text': selected_option_text or 'Not answered',
                'correct_answer': q.correct_answer,
                'correct_option_text': correct_option_text or 'N/A',
                'is_correct': answer.is_correct,
                'marks': q.marks,
                'marks_obtained': answer.marks_obtained,
                'explanation': q.explanation or '',
            })
        
        wrong_count = len([q for q in questions_details if not q['is_correct']])
        
        return Response({
            'attempt_id': attempt.id,
            'quiz_title': attempt.quiz.title,
            'student_name': f"{attempt.student.first_name} {attempt.student.last_name or ''}".strip(),
            'student_id': attempt.student.student_id,
            'attempt_number': attempt.attempt_number,
            'score': attempt.score,
            'total_marks': attempt.quiz.total_marks,
            'percentage': attempt.percentage,
            'is_passed': attempt.is_passed,
            'correct_count': sum(1 for a in answers if a.is_correct),
            'wrong_count': wrong_count,
            'total_questions': answers.count(),
            'submitted_at': attempt.submitted_at,
            'questions': questions_details,
        })
    except QuizAttempt.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def student_assigned_tests(request):
    try:
        student = Students.objects.get(user=request.user)
    except Students.DoesNotExist:
        return Response([])

    if not student.assigned_batch:
        return Response([])

    # Only tests assigned to THIS student's batch.
    assigned = AssignedTest.objects.filter(
        batch=student.assigned_batch
    ).select_related('test').order_by('-assigned_date')

    data = []
    assigned_test_ids = set()
    for a in assigned:
        test = a.test
        assigned_test_ids.add(test.id)
        questions = Question.objects.filter(test=test)
        data.append({
            'id': a.id,
            'test_id': test.id,
            'test_title': test.title,
            'test_description': test.description,
            'total_questions': questions.count(),
            'assigned_date': a.assigned_date if hasattr(a, 'assigned_date') else None,
        })

    # Legacy safety: older mentor-created tests sometimes exist without an
    # AssignedTest row, so include mentor-owned tests without duplicates.
    if student.assigned_staff:
        mentor_tests = QuizTest.objects.filter(
            created_by=student.assigned_staff
        ).exclude(id__in=assigned_test_ids).order_by('-created_at')
        for test in mentor_tests:
            questions = Question.objects.filter(test=test)
            if not questions.exists():
                continue
            data.append({
                'id': -test.id,
                'test_id': test.id,
                'test_title': test.title,
                'test_description': test.description,
                'total_questions': questions.count(),
                'assigned_date': test.created_at,
            })

    return Response(data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def student_test_results(request):
    try:
        student = Students.objects.get(user=request.user)
    except Students.DoesNotExist:
        return Response({'results': []})

    # -- Only results for THIS student --
    results = TestResult.objects.filter(
        student=student
    ).select_related('test').order_by('-submitted_at')

    data = []
    for r in results:
        data.append({
            'id': r.id,
            'test_id': r.test.id,
            'test_title': r.test.title,
            'score': r.score,
            'total_questions': r.total_questions,
            'percentage': float(r.percentage),
            'submitted_at': r.submitted_at,
        })

    return Response({'results': data})



@api_view(['POST'])
@permission_classes([IsAuthenticated])
def student_take_test(request, test_id):
    """Student submits test answers - ONE TIME ONLY"""
    try:
        student = Students.objects.get(user=request.user)
        test = QuizTest.objects.get(id=test_id)
        
        # Check if student is in a batch that has this test assigned
        if not student.assigned_batch:
            return Response({'error': 'You are not assigned to any batch'}, status=400)
        
        assigned_exists = AssignedTest.objects.filter(
            test=test,
            batch=student.assigned_batch
        ).exists()
        mentor_owned_test = bool(student.assigned_staff_id and test.created_by_id == student.assigned_staff_id)
        
        if not assigned_exists and not mentor_owned_test:
            return Response({'error': 'This test is not assigned to your batch'}, status=400)
        
        # ? ONE TIME ATTEMPT - Check if already taken
        if TestResult.objects.filter(student=student, test=test).exists():
            return Response({'error': 'You have already taken this test. Only one attempt allowed.'}, status=400)
        
        questions = Question.objects.filter(test=test)
        answers_data = request.data.get('answers', {})
        
        # Calculate score
        score = 0
        total_questions = questions.count()
        
        for question in questions:
            question_id = str(question.id)
            selected_value = answers_data.get(question_id, '')
            
            # Determine what the correct answer should be
            correct_text = ""
            if question.correct_answer in ['A', 'a', '1', question.option1]:
                correct_text = question.option1
            elif question.correct_answer in ['B', 'b', '2', question.option2]:
                correct_text = question.option2
            elif question.correct_answer in ['C', 'c', '3', question.option3]:
                correct_text = question.option3
            elif question.correct_answer in ['D', 'd', '4', question.option4]:
                correct_text = question.option4
            else:
                correct_text = question.correct_answer
            
            # Determine what the student selected
            selected_text = ""
            if selected_value in ['A', 'a', '1', 'option1']:
                selected_text = question.option1
            elif selected_value in ['B', 'b', '2', 'option2']:
                selected_text = question.option2
            elif selected_value in ['C', 'c', '3', 'option3']:
                selected_text = question.option3
            elif selected_value in ['D', 'd', '4', 'option4']:
                selected_text = question.option4
            else:
                selected_text = selected_value
            
            is_correct = selected_text.lower().strip() == correct_text.lower().strip()
            
            if is_correct:
                score += 1
        
        percentage = (score / total_questions * 100) if total_questions > 0 else 0
        is_passed = percentage >= 50
        
        # Save result
        result = TestResult.objects.create(
            student=student,
            test=test,
            score=score,
            total_questions=total_questions,
            percentage=percentage
        )
        
        return Response({
            'score': score,
            'total_questions': total_questions,
            'percentage': percentage,
            'is_passed': is_passed
        })
        
    except Students.DoesNotExist:
        return Response({'error': 'Student not found'}, status=404)
    except QuizTest.DoesNotExist:
        return Response({'error': 'Test not found'}, status=404)
    except Exception as e:
        print(f"Error: {e}")
        return Response({'error': str(e)}, status=400)



@api_view(['GET'])
@permission_classes([IsAuthenticated])
def test_result_details(request, result_id):
    """Get detailed test result with answers"""
    try:
        result = TestResult.objects.get(id=result_id)
        
        # Check if the result belongs to the logged-in student
        if result.student.user != request.user:
            return Response({'error': 'Access denied'}, status=403)
        
        # Get the questions and answers
        test = result.test
        questions = Question.objects.filter(test=test)
        
        # You'll need to store answers in a separate model or parse from result
        # For now, return basic details
        return Response({
            'test_title': test.title,
            'score': result.score,
            'total_questions': result.total_questions,
            'percentage': result.percentage,
            'is_passed': result.percentage >= 50,
            'correct_count': result.score,
            'wrong_count': result.total_questions - result.score,
            'submitted_at': result.submitted_at,
        })
    except TestResult.DoesNotExist:
        return Response({'error': 'Result not found'}, status=404)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def test_questions(request, test_id):
    """Get questions for a specific test"""
    try:
        test = QuizTest.objects.get(id=test_id)
        questions = Question.objects.filter(test=test)
        data = [{
            'id': q.id,
            'question_text': q.question_text,
            'option1': q.option1,
            'option2': q.option2,
            'option3': q.option3,
            'option4': q.option4,
        } for q in questions]
        return Response({'count': len(data), 'questions': data})
    except QuizTest.DoesNotExist:
        return Response({'error': 'Test not found'}, status=404)

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_test_questions(request, test_id):
    """Get questions for a specific test"""
    try:
        test = QuizTest.objects.get(id=test_id)
        questions = Question.objects.filter(test=test)
        
        questions_data = []
        for q in questions:
            # Determine the correct option letter based on stored correct_answer
            correct_option = "A"  # default
            if q.correct_answer in ["1", "A", q.option1]:
                correct_option = "A"
            elif q.correct_answer in ["2", "B", q.option2]:
                correct_option = "B"
            elif q.correct_answer in ["3", "C", q.option3]:
                correct_option = "C"
            elif q.correct_answer in ["4", "D", q.option4]:
                correct_option = "D"
            
            questions_data.append({
                'id': q.id,
                'question_text': q.question_text,
                'option1': q.option1,
                'option2': q.option2,
                'option3': q.option3 or '',
                'option4': q.option4 or '',
            })
        
        return Response({
            'test_id': test.id,
            'test_title': test.title,
            'total_questions': questions.count(),
            'questions': questions_data
        })
    except QuizTest.DoesNotExist:
        return Response({'error': 'Test not found'}, status=404)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def test_results(request):
    try:
        employee = Employee.objects.get(user=request.user)
    except Employee.DoesNotExist:
        return Response({'error': 'Employee not found'}, status=404)

    batch_id = request.query_params.get('batch_id')

    # -- Only show results for students assigned to THIS staff --
    results = TestResult.objects.select_related(
        'student', 'test', 'student__assigned_batch'
    ).filter(
        student__assigned_staff=employee  # only THIS staff's students
    ).order_by('-submitted_at')

    if batch_id:
        results = results.filter(student__assigned_batch_id=batch_id)

    data = []
    for r in results:
        data.append({
            'id': r.id,
            'student_name': f"{r.student.first_name} {r.student.last_name}",
            'student_id': r.student.student_id,
            'batch_number': r.student.assigned_batch.batch_number if r.student.assigned_batch else '—',
            'batch_id': r.student.assigned_batch.id if r.student.assigned_batch else None,
            'test_title': r.test.title,
            'test_id': r.test.id,
            'score': r.score,
            'total_questions': r.total_questions,
            'percentage': float(r.percentage),
            'submitted_at': r.submitted_at,
        })

    return Response(data)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_session_doubt_responses(request, session_id):
    """Get all doubt responses for a session"""
    try:
        student = Students.objects.get(user=request.user)
        session = CourseSession.objects.get(id=session_id)
        
        # Get the student's progress for this session
        progress = Student_Session_Progress.objects.filter(
            student=student, 
            session=session
        ).first()
        
        if not progress:
            return Response({'responses': [], 'has_response': False})
        
        # Get all responses for this doubt
        responses = DoubtResponse.objects.filter(
            doubt=progress
        ).select_related('staff').order_by('created_at')
        
        responses_data = []
        for response in responses:
            responses_data.append({
                'id': response.id,
                'staff_name': f"{response.staff.first_name} {response.staff.last_name or ''}",
                'message': response.message,
                'created_at': response.created_at,
            })
        
        return Response({
            'responses': responses_data,
            'has_response': len(responses_data) > 0,
            'session_title': session.title,
            'session_number': session.session_number
        })
        
    except Students.DoesNotExist:
        return Response({'error': 'Student not found'}, status=404)
    except CourseSession.DoesNotExist:
        return Response({'error': 'Session not found'}, status=404)
    except Exception as e:
        print(f"Error in get_session_doubt_responses: {e}")
        return Response({'error': str(e)}, status=500)

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def get_student_progress_summary(request, student_id):
    """Get summary of student's progress for the current trainer"""
    try:
        trainer = Employee.objects.get(user=request.user)
        student = Students.objects.get(id=student_id, assigned_staff=trainer)
        
        batch = student.assigned_batch
        if not batch:
            return Response({'error': 'Student not assigned to any batch'}, status=400)
        
        total_sessions = CourseSession.objects.filter(batch=batch).count()
        completed_sessions = Student_Session_Progress.objects.filter(
            student=student,
            session__batch=batch,
            completed=True
        ).count()
        
        # Get details of completed sessions
        completed_details = []
        progress_records = Student_Session_Progress.objects.filter(
            student=student,
            session__batch=batch,
            completed=True
        ).select_related('session').order_by('session__session_number')
        
        for p in progress_records:
            completed_details.append({
                'session_number': p.session.session_number,
                'title': p.session.title,
                'completed_date': p.completed_date,
            })
        
        return Response({
            'student_name': f"{student.first_name} {student.last_name or ''}",
            'student_id': student.student_id,
            'batch_number': batch.batch_number,
            'total_sessions': total_sessions,
            'completed_sessions': completed_sessions,
            'remaining_sessions': total_sessions - completed_sessions,
            'completion_percentage': round((completed_sessions / total_sessions * 100), 1) if total_sessions > 0 else 0,
            'completed_sessions_details': completed_details,
            'logsheet_url': (
    batch.course_logsheet.url
    if batch.course_logsheet
    else batch.course_name.course_logsheet.url
    if batch.course_name and batch.course_name.course_logsheet
    else None
),
        })
        
    except Employee.DoesNotExist:
        return Response({'error': 'Trainer not found'}, status=404)
    except Students.DoesNotExist:
        return Response({'error': 'Student not found'}, status=404)



# from django.http import HttpResponse
# import json

# @api_view(['GET'])
# @permission_classes([IsAuthenticated])
# def download_completion_report(request, student_id):
#     """Simple working version - returns text file for testing"""
#     try:
#         # Get the student
#         from .models import CompletedStudent
#         completed_student = CompletedStudent.objects.get(id=student_id)
        
#         # Create a simple text response
#         content = f"""
#         COMPLETION REPORT
#         =================
        
#         Student Name: {completed_student.first_name} {completed_student.last_name or ''}
#         Student ID: {completed_student.student_id}
#         Course: {completed_student.course_name or completed_student.course}
#         Batch: {completed_student.batch_number}
#         Branch: {completed_student.branch}
#         Start Date: {completed_student.batch_start_date}
#         Completion Date: {completed_student.completion_date}
#         Trainer: {completed_student.faculty_name}

#         """
        
#         response = HttpResponse(content, content_type='text/plain')
#         filename = f"completion_report_{completed_student.first_name}.txt"
#         response['Content-Disposition'] = f'attachment; filename="{filename}"'
#         return response
        
#     except CompletedStudent.DoesNotExist:
#         return Response({'error': f'Student with id {student_id} not found'}, status=404)
#     except Exception as e:
#         return Response({'error': str(e)}, status=500)


from django.http import HttpResponse
from io import BytesIO
import re

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def download_completion_report(request, student_id):
    """Generate a styled PDF completion report"""
    try:
        from .models import CompletedStudent
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, Image as RLImage
        from reportlab.lib.enums import TA_CENTER

        s = CompletedStudent.objects.get(id=student_id)
    except CompletedStudent.DoesNotExist:
        return Response({'error': f'Student {student_id} not found'}, status=404)
    except Exception as e:
        return Response({'error': str(e)}, status=500)

    # -- Filename: StudentName_CourseName_Completion_Report.pdf --------------
    BLANK = '\u2014'

    def display(value):
        if value is None or value == '':
            return BLANK
        return str(value)

    student_name = f"{display(s.first_name)} {s.last_name or ''}".strip()
    course_name  = display(s.course_name or s.course or 'Course').strip()

    def safe(t):
        return re.sub(r'[^\w\s-]', '', t).replace(' ', '_')

    filename = f"{safe(student_name)}_{safe(course_name)}_Completion_Report.pdf"
    print(f"[COMPLETION_REPORT] completed_student_id={student_id}")
    print(f"[COMPLETION_REPORT] student_name={student_name}")

    # -- Colours --------------------------------------------------------------
    NAVY  = colors.HexColor('#0f1b2d')
    AMBER = colors.HexColor('#f4a940')
    TEAL  = colors.HexColor('#2ec4b6')
    SAGE  = colors.HexColor('#4caf81')
    SLATE = colors.HexColor('#8099b3')
    LIGHT = colors.HexColor('#f8fafc')
    WHITE = colors.white

    # -- Document -------------------------------------------------------------
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=20*mm, rightMargin=20*mm,
        topMargin=18*mm, bottomMargin=18*mm,
    )
    W = A4[0] - 40*mm

    # -- Styles ---------------------------------------------------------------
    title_style    = ParagraphStyle('T',  fontName='Helvetica-Bold',    fontSize=22, textColor=WHITE, alignment=TA_CENTER)
    subtitle_style = ParagraphStyle('S',  fontName='Helvetica',         fontSize=11, textColor=colors.HexColor('#fcd17a'), alignment=TA_CENTER)
    section_style  = ParagraphStyle('Se', fontName='Helvetica-Bold',    fontSize=11, textColor=NAVY, spaceAfter=6, spaceBefore=4)
    label_style    = ParagraphStyle('L',  fontName='Helvetica-Bold',    fontSize=9,  textColor=SLATE)
    value_style    = ParagraphStyle('V',  fontName='Helvetica',         fontSize=10, textColor=NAVY)
    congrats_style = ParagraphStyle('C',  fontName='Helvetica-Bold',    fontSize=13, textColor=SAGE, alignment=TA_CENTER)
    note_style     = ParagraphStyle('N',  fontName='Helvetica',         fontSize=9,  textColor=SLATE, alignment=TA_CENTER)
    footer_style   = ParagraphStyle('F',  fontName='Helvetica-Oblique', fontSize=8,  textColor=SLATE, alignment=TA_CENTER)

    def row(label, value):
        return [Paragraph(label.upper(), label_style), Paragraph(display(value), value_style)]

    def make_table(data):
        t = Table(data, colWidths=[W * 0.32, W * 0.68])
        t.setStyle(TableStyle([
            ('BACKGROUND',    (0, 0), (0, -1), LIGHT),
            ('ROWBACKGROUNDS',(0, 0), (-1,-1), [WHITE, LIGHT]),
            ('GRID',          (0, 0), (-1,-1), 0.4, colors.HexColor('#e8e6e1')),
            ('TOPPADDING',    (0, 0), (-1,-1), 6),
            ('BOTTOMPADDING', (0, 0), (-1,-1), 6),
            ('LEFTPADDING',   (0, 0), (-1,-1), 10),
            ('RIGHTPADDING',  (0, 0), (-1,-1), 10),
        ]))
        return t

    story = []

    # -- Header banner --------------------------------------------------------
    logo_path = Path(settings.BASE_DIR) / 'connect' / 'assets' / 'IIE.png'
    logo = RLImage(str(logo_path), width=28 * mm, height=18 * mm) if logo_path.exists() else Paragraph('IIE CONNECT', title_style)
    ht = Table(
        [[logo,
          Paragraph('Certificate of Course Completion', subtitle_style)]],
        colWidths=[W * 0.28, W * 0.72]
    )
    ht.setStyle(TableStyle([
        ('BACKGROUND',    (0,0), (-1,-1), NAVY),
        ('TOPPADDING',    (0,0), (-1,-1), 14),
        ('BOTTOMPADDING', (0,0), (-1,-1), 14),
        ('LEFTPADDING',   (0,0), (-1,-1), 16),
        ('RIGHTPADDING',  (0,0), (-1,-1), 16),
        ('VALIGN',        (0,0), (-1,-1), 'MIDDLE'),
    ]))
    story.append(ht)
    story.append(Spacer(1, 10))

    # -- Congratulations strip ------------------------------------------------
    ct = Table([[Paragraph(f'Congratulations, {student_name}!', congrats_style)]], colWidths=[W])
    ct.setStyle(TableStyle([
        ('BACKGROUND',    (0,0), (-1,-1), colors.HexColor('#e8f8f0')),
        ('TOPPADDING',    (0,0), (-1,-1), 10),
        ('BOTTOMPADDING', (0,0), (-1,-1), 10),
        ('LEFTPADDING',   (0,0), (-1,-1), 14),
        ('RIGHTPADDING',  (0,0), (-1,-1), 14),
        ('BOX',           (0,0), (-1,-1), 1, SAGE),
    ]))
    story += [ct, Spacer(1, 4),
              Paragraph('This report confirms successful completion of the course programme.', note_style),
              Spacer(1, 12)]

    # -- Student Information --------------------------------------------------
    story.append(Paragraph('Student Information', section_style))
    story.append(HRFlowable(width=W, thickness=2, color=AMBER, spaceAfter=6))
    story.append(make_table([
        row('Student Name',  student_name),
        row('Student ID',    s.student_id),
        row('Email',         s.email),
        row('Mobile',        s.mobile_no),
        row('Gender',        s.gender),
        row('Date of Birth', s.date_of_birth),
        row('City / State',  f"{display(s.city)} / {display(s.state)}"),
        row('Qualification', s.qualification),
        row('Branch',        s.branch),
    ]))
    story.append(Spacer(1, 14))

    # -- Course & Batch Details -----------------------------------------------
    c_sess = s.completed_sessions_count or 0
    t_sess = s.total_sessions_count or 0
    sess_text = f"{c_sess} / {t_sess}" + (f"  ({round(c_sess/t_sess*100)}%)" if t_sess else '')

    story.append(Paragraph('Course & Batch Details', section_style))
    story.append(HRFlowable(width=W, thickness=2, color=TEAL, spaceAfter=6))
    story.append(make_table([
        row('Course',          course_name),
        row('Batch Number',    s.batch_number),
        row('Batch Start',     s.batch_start_date),
        row('Batch End',       s.batch_end_date),
        row('Sessions Done',   sess_text),
        row('Trainer',         s.faculty_name),
        row('Completion Date', s.completion_date),
    ]))
    story.append(Spacer(1, 18))

    story.append(Paragraph('Performance Summary', section_style))
    story.append(HRFlowable(width=W, thickness=2, color=SAGE, spaceAfter=6))
    story.append(make_table([
        row('Sessions Completed',      sess_text),
        row('Attendance Percentage',   f"{display(s.attendance_percentage)}%"),
        row('Average Test Score',      f"{display(s.average_test_score)}%"),
    ]))
    story.append(Spacer(1, 18))

    # -- Footer ---------------------------------------------------------------
    story.append(HRFlowable(width=W, thickness=1, color=SLATE, spaceAfter=6))
    story.append(Paragraph(
        'This is an auto-generated report from IIE Pulse. For queries contact administration.',
        footer_style
    ))

    # -- Build & return -------------------------------------------------------
    doc.build(story)
    buf.seek(0)

    pdf = buf.getvalue()
    print(f"[COMPLETION_REPORT] generated_pdf_size={len(pdf)}")

    response = HttpResponse(pdf, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    response['Content-Length'] = str(len(pdf))
    return response


@api_view(['PUT', 'PATCH'])
@permission_classes([IsAuthenticated])
def update_announcement(request, pk):
    try:
        announcement = Announcement.objects.get(id=pk)
        serializer = AnnouncementSerializer(announcement, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data)
        return Response(serializer.errors, status=400)
    except Announcement.DoesNotExist:
        return Response({'error': 'Announcement not found'}, status=404)


@api_view(['PUT', 'PATCH'])
@permission_classes([IsAuthenticated])
def update_question(request, question_id):
    try:
        question = Question.objects.get(id=question_id)
        data = request.data
        
        question.question_text = data.get('question_text', question.question_text)
        question.option1 = data.get('option1', question.option1)
        question.option2 = data.get('option2', question.option2)
        question.option3 = data.get('option3', question.option3)
        question.option4 = data.get('option4', question.option4)
        question.correct_answer = data.get('correct_answer', question.correct_answer)
        question.save()
        
        return Response({'message': 'Question updated successfully'})
    except Question.DoesNotExist:
        return Response({'error': 'Question not found'}, status=404)


@api_view(['DELETE'])
@permission_classes([IsAuthenticated])
def delete_question(request, question_id):
    try:
        question = Question.objects.get(id=question_id)
        question.delete()
        return Response({'message': 'Question deleted successfully'})
    except Question.DoesNotExist:
        return Response({'error': 'Question not found'}, status=404)




@api_view(['POST'])
@permission_classes([AllowAny])
def forgot_password(request):
    email = request.data.get('email', '').strip().lower()

    if not email:
        return Response({'error': 'Email is required'}, status=400)

    user = User.objects.filter(Q(username=email) | Q(email=email)).first()

    if not user:
        return Response({'error': 'No account found with this email'}, status=404)

    otp = str(random.randint(100000, 999999))
    cache.set(f'otp_{email}', otp, timeout=600)

    try:
        portal_url = get_portal_url(request)
        send_email_async(
            subject='IIE Pulse — Password Reset OTP',
            message=f"""Dear {user.first_name or user.username},

Your OTP for password reset is:

    {otp}

This OTP is valid for 10 minutes.

If you did not request this, please ignore this email.

Best regards,
IIE Pulse Team
{portal_url}
""",
            recipient_list=[email],
        )
    except Exception:
        logger.exception('Failed to queue password reset email for %s', email)

    return Response({'message': 'OTP sent successfully'})


@api_view(['POST'])
@permission_classes([AllowAny])
def reset_password(request):
    from django.contrib.auth.password_validation import validate_password
    from django.core.exceptions import ValidationError

    email = request.data.get('email', '').strip().lower()
    otp = request.data.get('otp', '').strip()
    new_password = request.data.get('new_password', '').strip()

    if not all([email, otp, new_password]):
        return Response({'error': 'All fields are required'}, status=400)

    cached_otp = cache.get(f'otp_{email}')

    if not cached_otp:
        return Response({'error': 'OTP has expired. Please request a new one'}, status=400)
    if cached_otp != otp:
        return Response({'error': 'Invalid OTP'}, status=400)

    user = User.objects.filter(Q(username=email) | Q(email=email)).first()

    if not user:
        return Response({'error': 'User not found'}, status=404)

    try:
        validate_password(new_password, user)
    except ValidationError as exc:
        return Response({'error': ' '.join(exc.messages)}, status=400)

    user.set_password(new_password)
    user.save()
    cache.delete(f'otp_{email}')

    return Response({'message': 'Password reset successfully'})



# -- Fee Management Views ------------------------------------------------------

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def student_fee_details(request):
    try:
        student_id = request.query_params.get('student_id') or request.query_params.get('student')
        if student_id:
            employee = Employee.objects.filter(user=request.user).first()
            if not (is_admin_user(request.user) or employee):
                return Response({'error': 'Access denied'}, status=403)
            student_qs = Students.objects.all()
            if employee and not is_admin_user(request.user):
                student_qs = student_qs.filter(branch=employee.branch)
            student = student_qs.get(id=student_id)
        else:
            student = Students.objects.get(user=request.user)
        fee = FeePayment.objects.filter(student=student).order_by('-created_at').first()
        if not fee:
            return Response({'fee': None})
        transactions = FeeTransaction.objects.filter(fee_payment=fee).order_by('-paid_at')
        return Response({
            'fee': {
                'id': fee.id,
                'total_fee': float(fee.total_fee),
                'amount_paid': float(fee.amount_paid),
                'balance': float(fee.balance),
                'is_fully_paid': fee.is_fully_paid,
                'batch_number': fee.batch.batch_number,
                'course_name': fee.batch.course_name.course_name if fee.batch.course_name else '—',
            },
            'transactions': [{
                'id': t.id,
                'amount': float(t.amount),
                'payment_mode': t.payment_mode,
                'notes': t.notes or '—',
                'paid_at': t.paid_at,
                'bill_generated': t.bill_generated,
            } for t in transactions]
        })
    except Students.DoesNotExist:
        return Response({'error': 'Student not found'}, status=404)

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_fee_list(request):
    """Admin views all fee records; counselors/staff view their branch."""
    employee = Employee.objects.filter(user=request.user).first()
    if not is_admin_user(request.user) and not employee:
        return Response({'error': 'Access denied.'}, status=403)
    fees = FeePayment.objects.select_related('student', 'batch').order_by('-created_at')
    branch = request.query_params.get('branch')
    if employee and not is_admin_user(request.user):
        branch = employee.branch
    if branch:
        fees = fees.filter(student__branch=branch)

    data = []
    for fee in fees:
        data.append({
            'id': fee.id,
            'student_name': f"{fee.student.first_name} {fee.student.last_name or ''}",
            'student_id': fee.student.student_id,
            'branch': fee.student.branch,
            'batch_number': fee.batch.batch_number,
            'course_name': fee.batch.course_name.course_name if fee.batch.course_name else '—',
            'total_fee': float(fee.total_fee),
            'amount_paid': float(fee.amount_paid),
            'balance': float(fee.balance),
            'is_fully_paid': fee.is_fully_paid,
            'created_at': fee.created_at,
        })
    return Response({'results': data})


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def add_fee_payment(request, fee_id):
    try:
        user = request.user
        fee = FeePayment.objects.get(id=fee_id)

        # -- Allow admin and counselor -----------------------------
        if not (user.is_superuser or user.is_staff):
            try:
                emp = Employee.objects.get(user=user)
                if emp.designation.lower() != 'counselor':
                    return Response({'error': 'Access denied'}, status=403)
                if fee.student.branch != emp.branch:
                    return Response({'error': 'Access denied — different branch'}, status=403)
            except Employee.DoesNotExist:
                return Response({'error': 'Access denied'}, status=403)
        # ---------------------------------------------------------

        amount = float(request.data.get('amount', 0))
        payment_mode = request.data.get('payment_mode', 'cash')
        notes = request.data.get('notes', '')

        if amount <= 0:
            return Response({'error': 'Amount must be greater than 0'}, status=400)
        if amount > float(fee.balance):
            return Response({'error': f'Amount exceeds balance of ?{fee.balance}'}, status=400)

        from decimal import Decimal
        transaction = FeeTransaction.objects.create(
            fee_payment=fee,
            amount=amount,
            payment_mode=payment_mode,
            notes=notes,
            collected_by=request.user,
        )

        fee.refresh_from_db()
        fee.amount_paid = fee.amount_paid + Decimal(str(amount))
        fee.balance = fee.total_fee - fee.amount_paid
        fee.is_fully_paid = fee.balance <= Decimal('0')
        fee.save()

        return Response({
            'success': True,
            'message': f'Payment of ?{amount} recorded successfully!',
            'balance': float(fee.balance),
            'is_fully_paid': fee.is_fully_paid,
        })
    except FeePayment.DoesNotExist:
        return Response({'error': 'Fee record not found'}, status=404)
    except Exception as e:
        import traceback; traceback.print_exc()
        return Response({'error': str(e)}, status=500)

from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import os

# Register a Unicode font
font_path = os.path.join(os.path.dirname(__file__), 'fonts', 'NotoSans-Regular.ttf')
font_bold_path = os.path.join(os.path.dirname(__file__), 'fonts', 'NotoSans-Bold.ttf')

if os.path.exists(font_path):
    pdfmetrics.registerFont(TTFont('NotoSans', font_path))
    pdfmetrics.registerFont(TTFont('NotoSans-Bold', font_bold_path))
    MAIN_FONT = 'NotoSans'
    BOLD_FONT = 'NotoSans-Bold'
else:
    MAIN_FONT = 'Helvetica'
    BOLD_FONT = 'Helvetica-Bold'


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def generate_bill(request, fee_id):
    """Generate PDF bill in IIE format"""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import mm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, Image as RLImage
        from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
        from django.http import HttpResponse
        from io import BytesIO

        # -- Access control ----------------------------------------
        # -- Access control ----------------------------------------
        user = request.user
        if user.is_superuser or user.is_staff:
            fee = FeePayment.objects.select_related('student', 'batch').get(id=fee_id)
        else:
            try:
                emp = Employee.objects.get(user=user)
                if emp.designation.lower() == 'counselor':
                    fee = FeePayment.objects.select_related('student', 'batch').get(
                        id=fee_id, student__branch=emp.branch
                    )
                else:
                    return Response({'error': 'Access denied'}, status=403)
            except Employee.DoesNotExist:
                try:
                    student = Students.objects.get(user=user)
                    fee = FeePayment.objects.select_related('student', 'batch').get(id=fee_id, student=student)
                except Students.DoesNotExist:
                    return Response({'error': 'Access denied'}, status=403)

        transactions = FeeTransaction.objects.filter(fee_payment=fee).order_by('paid_at')

        # -- Colors ------------------------------------------------
        NAVY   = colors.HexColor('#1a237e')
        BLACK  = colors.black
        GRAY   = colors.HexColor('#555555')
        WHITE  = colors.white
        BORDER = colors.HexColor('#cccccc')
        LGRAY  = colors.HexColor('#f9f9f9')

        buf = BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=A4,
            leftMargin=15*mm, rightMargin=15*mm,
            topMargin=10*mm, bottomMargin=12*mm)
        W = A4[0] - 30*mm

        # -- Styles ------------------------------------------------
        inst_name_s = ParagraphStyle('IN', fontName=BOLD_FONT, fontSize=18, textColor=NAVY, alignment=TA_CENTER, spaceAfter=6)
        inst_addr_s = ParagraphStyle('IA', fontName=MAIN_FONT, fontSize=9, textColor=GRAY, alignment=TA_CENTER, leading=16, spaceAfter=4)
        label_s     = ParagraphStyle('L',  fontName=BOLD_FONT, fontSize=9.5, textColor=BLACK)
        value_s     = ParagraphStyle('V',  fontName=MAIN_FONT, fontSize=9.5, textColor=BLACK)
        footer_s    = ParagraphStyle('F',  fontName=MAIN_FONT, fontSize=8, textColor=GRAY, alignment=TA_RIGHT)
        note_s      = ParagraphStyle('N',  fontName=MAIN_FONT, fontSize=8, textColor=GRAY, alignment=TA_LEFT)

        story = []

        # -- Receipt & Date info -----------------------------------
        receipt_no      = f"IIE{fee.student.branch.upper()[:3]}{fee.id:04d}"
        date_str        = fee.updated_at.strftime('%Y-%m-%d')
        received_amount = float(transactions.last().amount) if transactions.exists() else float(fee.amount_paid)
        already_paid    = float(fee.amount_paid) - received_amount

        # -- Logo --------------------------------------------------
        # -- Logo --------------------------------------------------
        from django.conf import settings

        logo_path = None
        possible_paths = [
            # Django app static folder
            os.path.join(os.path.dirname(__file__), 'static', 'IIE.png'),
            # Project root static folder  
            os.path.join(settings.BASE_DIR, 'static', 'IIE.png'),
            # Frontend assets
            os.path.join(settings.BASE_DIR, '..', 'frontend', 'src', 'assets', 'IIE.png'),
            os.path.join(settings.BASE_DIR, 'frontend', 'src', 'assets', 'IIE.png'),
            # Absolute fallback
            r'C:\Users\vivek\OneDrive\Desktop\iie_connect_fullstack\frontend\src\assets\IIE.png',
        ]

        for p in possible_paths:
            p = os.path.normpath(p)
            print(f"Checking: {p} ? {os.path.exists(p)}")
            if os.path.exists(p):
                logo_path = p
                break

        logo_cell = RLImage(logo_path, width=28*mm, height=28*mm) if logo_path else Paragraph('<b>IIE</b>', label_s)
        # ----------------------------------------------------------
        # HEADER — Logo left, Institution center
        # ----------------------------------------------------------
        inst_table = Table([
        [Paragraph('INDRA INSTITUTE OF EDUCATION', inst_name_s)],
        [Paragraph('65/1, Tatabad, 7th Street, Dr Rajendra Prasad Rd, near BEA,\nGandhipuram, Coimbatore - 641012', inst_addr_s)],
        [Paragraph('IT Training and Testing Services', inst_addr_s)],
        [Paragraph('Ph : +91-9159779111', inst_addr_s)],
        ], colWidths=[W * 0.80])
        inst_table.setStyle(TableStyle([
            ('TOPPADDING',    (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ]))

        header_t = Table([[logo_cell, inst_table]], colWidths=[W*0.20, W*0.80])
        header_t.setStyle(TableStyle([
            ('VALIGN',        (0,0), (-1,-1), 'MIDDLE'),
            ('BOX',           (0,0), (-1,-1), 1, BORDER),
            ('TOPPADDING',    (0,0), (-1,-1), 10),
            ('BOTTOMPADDING', (0,0), (-1,-1), 10),
            ('LEFTPADDING',   (0,0), (-1,-1), 8),
            ('RIGHTPADDING',  (0,0), (-1,-1), 8),
        ]))
        story += [header_t, Spacer(1, 16)]

        # ----------------------------------------------------------
        # MAIN DETAILS TABLE
        # ----------------------------------------------------------
        LW = W * 0.20   # label col
        VW = W * 0.30   # value col

        def lbl(text): return Paragraph(f'<b>{text}</b>', label_s)
        def val(text): return Paragraph(str(text) if text else '—', value_s)

        # Calculate installments
        # -- Calculate installment rows ----------------------------
        total = float(fee.total_fee)
        num_transactions = transactions.count()

        # Build dynamic installment info for right side
        right_extra_rows = []
        for idx, t in enumerate(transactions, 1):
            right_extra_rows.append((f'{idx}{"st" if idx==1 else "nd" if idx==2 else "rd" if idx==3 else "th"} Fee:', f"Rs. {float(t.amount):,.0f}"))

        details_data = [
            [lbl('Date :'), val(date_str), lbl('Receipt No :'), val(receipt_no)],
            [lbl('Student Name :'), val(f"{fee.student.first_name} {fee.student.last_name or ''}"),
            lbl('Batch Date :'), val(str(fee.batch.start_date) if fee.batch.start_date else '—')],
            [lbl('Mobile No :'), val(fee.student.mobile_no or '—'),
            lbl('Course Fees :'), val(f"Rs. {total:,.0f}")],
            [lbl('Email ID :'), val(fee.student.email or '—'),
            lbl('Received Fees :'), val(f"Rs. {float(fee.amount_paid):,.0f}")],
            [lbl('Address:'), val(f"{fee.student.city or '—'},"),
            lbl(right_extra_rows[0][0] if len(right_extra_rows) > 0 else ''),
            val(right_extra_rows[0][1] if len(right_extra_rows) > 0 else '')],
            [Paragraph('', value_s), val(f"{fee.student.state or ''}"),
            lbl(right_extra_rows[1][0] if len(right_extra_rows) > 1 else ''),
            val(right_extra_rows[1][1] if len(right_extra_rows) > 1 else '')],
        ]

        # Add extra fee rows if more than 2 transactions
        for idx in range(2, len(right_extra_rows)):
            details_data.append([
                Paragraph('', value_s), Paragraph('', value_s),
                lbl(right_extra_rows[idx][0]),
                val(right_extra_rows[idx][1]),
            ])

        # Add course + balance row
        details_data.append([
            lbl('Course:'),
            val(fee.batch.course_name.course_name if fee.batch.course_name else '—'),
            lbl('Balance Due:'),
            val('Completely Paid' if fee.is_fully_paid else f"Rs. {float(fee.balance):,.0f}"),
        ])

        # Add status row
        details_data.append([
            Paragraph('', value_s), Paragraph('', value_s),
            lbl('Status:'),
            val('Completely Paid' if fee.is_fully_paid else 'Pending'),
        ])

        details_t = Table(details_data, colWidths=[LW, VW, LW, VW])
        details_t.setStyle(TableStyle([
        ('BOX',           (0,0), (-1,-1), 1, BORDER),
        ('LINEBELOW',     (0,0), (-1,-1), 0.3, BORDER),
        ('TOPPADDING',    (0,0), (-1,-1), 12),      # ? increased from 7
        ('BOTTOMPADDING', (0,0), (-1,-1), 12),      # ? increased from 7
        ('LEFTPADDING',   (0,0), (-1,-1), 12),      # ? increased from 8
        ('RIGHTPADDING',  (0,0), (-1,-1), 12),      # ? increased from 8
        ('VALIGN',        (0,0), (-1,-1), 'TOP'),
        ('ROWBACKGROUNDS',(0,0), (-1,-1), [WHITE, LGRAY]),
        ]))
        story += [details_t, Spacer(1, 16)]

        # ----------------------------------------------------------
        # PAYMENT HISTORY (if multiple transactions)
        # ----------------------------------------------------------
        if transactions.count() > 1:
            tx_data = [[
                Paragraph('<b>DATE</b>', label_s),
                Paragraph('<b>AMOUNT</b>', label_s),
                Paragraph('<b>MODE</b>', label_s),
                Paragraph('<b>NOTES</b>', label_s),
            ]]
            for t in transactions:
                tx_data.append([
                    val(t.paid_at.strftime('%Y-%m-%d')),
                    val(f"Rs. {float(t.amount):,.0f}"),
                    val(t.get_payment_mode_display()),
                    val(t.notes or '—'),
                ])
            tt = Table(tx_data, colWidths=[W*0.20, W*0.20, W*0.20, W*0.40])
            tt.setStyle(TableStyle([
                ('BOX',           (0,0), (-1,-1), 1, BORDER),
                ('LINEBELOW',     (0,0), (-1,-1), 0.3, BORDER),
                ('BACKGROUND',    (0,0), (-1,0), LGRAY),
                ('ROWBACKGROUNDS',(0,1), (-1,-1), [WHITE, LGRAY]),
                ('TOPPADDING',    (0,0), (-1,-1), 6),
                ('BOTTOMPADDING', (0,0), (-1,-1), 6),
                ('LEFTPADDING',   (0,0), (-1,-1), 8),
            ]))
            story += [tt, Spacer(1, 10)]

        # ----------------------------------------------------------
        # FOOTER
        # ----------------------------------------------------------
        story.append(HRFlowable(width=W, thickness=0.5, color=BORDER, spaceAfter=4))
        footer_data = [[
            Paragraph('* Fees once paid cannot be refunded.', note_s),
            Paragraph('This is computer generated bill no signature required.', footer_s),
        ]]
        footer_t = Table(footer_data, colWidths=[W*0.5, W*0.5])
        footer_t.setStyle(TableStyle([
            ('TOPPADDING',    (0,0), (-1,-1), 3),
            ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ]))
        story.append(footer_t)

        # -- Build -------------------------------------------------
        doc.build(story)
        buf.seek(0)
        transactions.update(bill_generated=True)

        student_name = f"{fee.student.student_id}_{fee.student.first_name}_{fee.student.last_name or ''}".strip('_')
        course_name  = fee.batch.course_name.course_name.replace(' ', '_') if fee.batch.course_name else 'Course'
        response = HttpResponse(buf.read(), content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="{student_name}_{course_name}.pdf"'
        return response

    except FeePayment.DoesNotExist:
        return Response({'error': 'Fee record not found'}, status=404)
    except Exception as e:
        import traceback
        traceback.print_exc()
        return Response({'error': str(e)}, status=500)




@api_view(['POST'])
@permission_classes([IsAuthenticated])
@parser_classes([MultiPartParser, FormParser, JSONParser])
def create_fee_payment_request(request, fee_id):
    """Counselor submits payment request to admin for approval"""
    try:
        user = request.user
        try:
            emp = Employee.objects.get(user=user)
            if emp.designation.lower() != 'counselor':
                return Response({'error': 'Access denied'}, status=403)
        except Employee.DoesNotExist:
            return Response({'error': 'Access denied'}, status=403)

        fee = FeePayment.objects.get(id=fee_id, student__branch=emp.branch)

        amount = request.data.get('amount')
        payment_mode = request.data.get('payment_mode', 'cash')
        notes = request.data.get('notes', '')
        screenshot = request.FILES.get('screenshot', None)

        if not amount or float(amount) <= 0:
            return Response({'error': 'Enter valid amount'}, status=400)
        if float(amount) > float(fee.balance):
            return Response({'error': f'Amount exceeds balance of ?{fee.balance}'}, status=400)

        # Check if there's already a pending request
        existing = FeePaymentRequest.objects.filter(
            fee_payment=fee, status='pending'
        ).first()
        if existing:
            return Response({'error': 'A pending request already exists for this student'}, status=400)

        req = FeePaymentRequest.objects.create(
            student=fee.student,
            fee_payment=fee,
            amount=amount,
            payment_mode=payment_mode,
            notes=notes,
            screenshot=screenshot,
            status='pending',
        )

        # Notify admin via email (async)
        try:
            admins = User.objects.filter(is_superuser=True)
            admin_emails = list(admins.values_list('email', flat=True))
            if admin_emails:
                send_email_async(
                    subject=f'Fee Payment Request — {fee.student.first_name} {fee.student.last_name or ""}',
                    message=f"""Dear Admin,

Counselor {emp.first_name} {emp.last_name or ''} has submitted a payment request for:

Student    : {fee.student.first_name} {fee.student.last_name or ''} ({fee.student.student_id})
Branch     : {fee.student.branch}
Amount     : Rs. {amount}
Mode       : {payment_mode}
Notes      : {notes or 'N/A'}
Batch      : {fee.batch.batch_number}
Balance Due: Rs. {fee.balance}

Please review and approve/reject in the Fee Management ? Payment Requests section.

Best regards,
IIE Pulse System
""",
                    recipient_list=admin_emails,
                )
        except Exception as e:
            print(f"Email queueing error: {e}")

        return Response({
            'success': True,
            'message': 'Payment request submitted to admin for approval!',
            'request_id': req.id,
        })

    except FeePayment.DoesNotExist:
        return Response({'error': 'Fee record not found'}, status=404)
    except Exception as e:
        import traceback; traceback.print_exc()
        return Response({'error': str(e)}, status=500)




@api_view(['GET'])
@permission_classes([IsAuthenticated])
def admin_fee_payment_requests(request):
    """Admin views all payment requests from students"""
    requests_qs = FeePaymentRequest.objects.select_related(
        'student', 'fee_payment', 'fee_payment__batch', 'fee_payment__batch__course_name'
    ).order_by('-requested_at')

    status_filter = request.query_params.get('status', 'pending')
    if status_filter:
        requests_qs = requests_qs.filter(status=status_filter)

    data = []
    for r in requests_qs:
        # -- Always fetch fresh balance from fee_payment -----------
        fee_payment = r.fee_payment
        fee_payment.refresh_from_db()
        current_balance = float(fee_payment.balance)
        balance_after = max(current_balance - float(r.amount), 0)
        # ---------------------------------------------------------

        # -- Screenshot URL ----------------------------------------
        screenshot_url = None
        if r.screenshot:
            try:
                screenshot_url = request.build_absolute_uri(r.screenshot.url)
            except Exception:
                screenshot_url = None
        # ---------------------------------------------------------

        data.append({
            'id': r.id,
            'student_name': f"{r.student.first_name} {r.student.last_name or ''}",
            'student_id': r.student.student_id,
            'branch': r.student.branch,
            'fee_id': fee_payment.id,
            'amount': float(r.amount),
            'payment_mode': r.payment_mode,
            'notes': r.notes,
            'status': r.status,
            'requested_at': r.requested_at,
            'batch_number': fee_payment.batch.batch_number,
            'course_name': fee_payment.batch.course_name.course_name if fee_payment.batch.course_name else '—',
            'current_balance': current_balance,
            'balance_after': balance_after,
            'screenshot': screenshot_url,   # ? ADD THIS
        })
    return Response({'results': data})



@api_view(['POST'])
@permission_classes([IsAuthenticated])
def process_fee_payment_request(request, request_id):
    """Admin approves or rejects a student payment request"""
    try:
        pay_req = FeePaymentRequest.objects.get(id=request_id)
        action = request.data.get('action')  # 'approve' or 'reject'

        if action == 'approve':
            fee = pay_req.fee_payment
            from decimal import Decimal

            FeeTransaction.objects.create(
                fee_payment=fee,
                amount=pay_req.amount,
                payment_mode=pay_req.payment_mode,
                notes=pay_req.notes or 'Student payment request approved',
                collected_by=request.user,
                bill_generated=False,
            )

            # -- Force refresh and recalculate ---------------------
            fee.refresh_from_db()
            fee.amount_paid = fee.amount_paid + Decimal(str(pay_req.amount))
            fee.balance = fee.total_fee - fee.amount_paid
            fee.is_fully_paid = fee.balance <= Decimal('0')
            fee.save()
            # -----------------------------------------------------

            # -- Mark request as approved --------------------------
            pay_req.status = 'approved'
            pay_req.reviewed_at = timezone.now()
            pay_req.reviewed_by = request.user
            pay_req.save()
            # -----------------------------------------------------

            # -- Notify student ------------------------------------
            # -- Notify student (async) --------------------------------
            try:
                send_email_async(
                    subject='Fee Payment Verified ?',
                    message=f"""Dear {pay_req.student.first_name},

Your payment of Rs. {pay_req.amount} has been verified and recorded.

Remaining Balance: Rs. {fee.balance}
Status: {'Fully Paid' if fee.is_fully_paid else 'Partial Payment'}

Best regards,
IIE Pulse Team
""",
                    recipient_list=[pay_req.student.email],
                )
            except Exception as e:
                print(f"Email queueing error: {e}")
            # -----------------------------------------------------

            return Response({
                'success': True,
                'message': f'Payment of Rs. {pay_req.amount} approved and recorded!',
                'new_balance': float(fee.balance),
                'is_fully_paid': fee.is_fully_paid,
            })

        elif action == 'reject':
            pay_req.status = 'rejected'
            pay_req.reviewed_at = timezone.now()
            pay_req.reviewed_by = request.user
            pay_req.save()
            return Response({'success': True, 'message': 'Payment request rejected.'})

        return Response({'error': 'Invalid action'}, status=400)

    except FeePaymentRequest.DoesNotExist:
        return Response({'error': 'Request not found'}, status=404)
    except Exception as e:
        import traceback; traceback.print_exc()
        return Response({'error': str(e)}, status=500)







# -- COUNSELOR ANNOUNCEMENTS ----------------------------------------------------

from connect.serializers import CounselorAnnouncementSerializer


class CounselorAnnouncementListView(generics.ListAPIView):
    serializer_class = CounselorAnnouncementSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        user = self.request.user
        counselor = Employee.objects.filter(user=user, designation__iexact='counselor').first()
        if not counselor:
            return CounselorAnnouncement.objects.none()
        return CounselorAnnouncement.objects.filter(created_by=user).order_by('-created_at')

    def list(self, request, *args, **kwargs):
        try:
            return super().list(request, *args, **kwargs)
        except Exception as e:
            logger.exception("Unhandled exception in CounselorAnnouncementListView.list")
            return Response({
                'error': 'Counselor announcements retrieval failed',
                'details': str(e),
                'view': 'CounselorAnnouncementListView.list'
            }, status=500)


class CounselorAnnouncementCreateView(generics.CreateAPIView):
    serializer_class = CounselorAnnouncementSerializer
    permission_classes = [IsAuthenticated]

    def perform_create(self, serializer):
        user = self.request.user
        try:
            emp = Employee.objects.get(user=user, designation='counselor')
        except Employee.DoesNotExist:
            from rest_framework.exceptions import PermissionDenied
            raise PermissionDenied("Only counselors can use this endpoint.")
        
        # Save the announcement first
        announcement = serializer.save(created_by=user, branch=emp.branch)
        
        # Handle specific students
        specific_student_ids = self.request.data.get('specific_student_ids', [])
        if specific_student_ids:
            from connect.models import Students
            students = Students.objects.filter(id__in=specific_student_ids)
            announcement.specific_students.set(students)

        branch_values = _branch_match_values(emp.branch)
        mentors = Employee.objects.filter(Q(designation__iexact='mentor') | Q(designation__iexact='trainer'))
        if branch_values:
            mentors = mentors.filter(branch__in=branch_values)
        for mentor in mentors.select_related('user').distinct():
            _queue_user_notification(
                mentor.user,
                user,
                'announcement',
                'Counselor Announcement',
                f"Counselor Announcement: {announcement.title}",
                False,
            )
        
        return announcement

@api_view(['PATCH'])
@permission_classes([IsAuthenticated])
def counselor_update_announcement(request, pk):
    try:
        Employee.objects.get(user=request.user, designation='counselor')
        ann = CounselorAnnouncement.objects.get(id=pk, created_by=request.user)
        serializer = CounselorAnnouncementSerializer(ann, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data)
        return Response(serializer.errors, status=400)
    except Employee.DoesNotExist:
        return Response({'error': 'Permission denied'}, status=403)
    except CounselorAnnouncement.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)


@api_view(['PATCH'])
@permission_classes([IsAuthenticated])
def counselor_toggle_announcement(request, pk):
    try:
        Employee.objects.get(user=request.user, designation='counselor')
        ann = CounselorAnnouncement.objects.get(id=pk, created_by=request.user)
        ann.is_published = not ann.is_published
        ann.save()
        return Response({'is_published': ann.is_published})
    except Employee.DoesNotExist:
        return Response({'error': 'Permission denied'}, status=403)
    except CounselorAnnouncement.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)


@api_view(['DELETE'])
@permission_classes([IsAuthenticated])
def counselor_delete_announcement(request, pk):
    try:
        Employee.objects.get(user=request.user, designation='counselor')
        CounselorAnnouncement.objects.get(id=pk, created_by=request.user).delete()
        return Response(status=204)
    except Employee.DoesNotExist:
        return Response({'error': 'Permission denied'}, status=403)
    except CounselorAnnouncement.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)


def _get_trainer_for_request(request):
    emp = Employee.objects.filter(user=request.user).first()
    if emp and (emp.designation or '').strip().lower() in ['trainer', 'mentor']:
        return emp
    return None


class TrainerAnnouncementListView(generics.ListAPIView):
    serializer_class = CounselorAnnouncementSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        trainer = _get_trainer_for_request(self.request)
        if not trainer:
            return CounselorAnnouncement.objects.none()
        return CounselorAnnouncement.objects.filter(created_by=self.request.user).order_by('-created_at')


class TrainerAnnouncementCreateView(generics.CreateAPIView):
    serializer_class = CounselorAnnouncementSerializer
    permission_classes = [IsAuthenticated]

    def get_serializer(self, *args, **kwargs):
        if 'data' in kwargs:
            data = kwargs['data'].copy()
            data.pop('specific_student_ids', None)
            kwargs['data'] = data
        return super().get_serializer(*args, **kwargs)

    def perform_create(self, serializer):
        trainer = _get_trainer_for_request(self.request)
        if not trainer:
            raise PermissionDenied("Only trainers can use this endpoint.")

        recipient_type = self.request.data.get('recipient_type')
        batch_id = self.request.data.get('specific_batch')
        if recipient_type not in ['specific_batch', 'specific_student']:
            raise PermissionDenied("Trainer announcements must target a specific batch or student.")
        if not batch_id:
            raise PermissionDenied("Select a batch for trainer announcement.")

        batch = Batches.objects.filter(id=batch_id, faculty=trainer).first()
        if not batch:
            raise PermissionDenied("You can announce only to your own batches.")

        announcement = serializer.save(
            created_by=self.request.user,
            branch=trainer.branch,
            specific_batch=batch,
            is_important=recipient_type == 'important' or self.request.data.get('announcement_type') == 'important',
        )
        specific_student_ids = self.request.data.get('specific_student_ids', [])
        if isinstance(specific_student_ids, str):
            specific_student_ids = [sid for sid in specific_student_ids.split(',') if sid]

        if recipient_type == 'specific_student':
            students = Students.objects.filter(
                id__in=specific_student_ids,
                assigned_batch=batch,
            ).filter(Q(assigned_staff=trainer) | Q(assigned_batch__faculty=trainer)).distinct()
            if not students.exists():
                raise PermissionDenied("Select at least one student from this batch.")
            announcement.specific_students.set(students)
        else:
            announcement.specific_students.clear()

        return announcement


@api_view(['PATCH'])
@permission_classes([IsAuthenticated])
def trainer_update_announcement(request, pk):
    trainer = _get_trainer_for_request(request)
    if not trainer:
        return Response({'error': 'Permission denied'}, status=403)

    try:
        ann = CounselorAnnouncement.objects.get(id=pk, created_by=request.user)
    except CounselorAnnouncement.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)

    recipient_type = request.data.get('recipient_type', ann.recipient_type)
    batch_id = request.data.get('specific_batch') or ann.specific_batch_id
    batch = Batches.objects.filter(id=batch_id, faculty=trainer).first()
    if recipient_type not in ['specific_batch', 'specific_student'] or not batch:
        return Response({'error': 'Select one of your batches.'}, status=400)

    serializer_data = request.data.copy()
    serializer_data.pop('specific_student_ids', None)
    serializer = CounselorAnnouncementSerializer(ann, data=serializer_data, partial=True)
    if not serializer.is_valid():
        return Response(serializer.errors, status=400)
    ann = serializer.save(
        branch=trainer.branch,
        specific_batch=batch,
        is_important=request.data.get('announcement_type', ann.announcement_type) == 'important',
    )

    specific_student_ids = request.data.get('specific_student_ids', [])
    if isinstance(specific_student_ids, str):
        specific_student_ids = [sid for sid in specific_student_ids.split(',') if sid]

    if recipient_type == 'specific_student':
        students = Students.objects.filter(
            id__in=specific_student_ids,
            assigned_batch=batch,
        ).filter(Q(assigned_staff=trainer) | Q(assigned_batch__faculty=trainer)).distinct()
        if not students.exists():
            return Response({'error': 'Select at least one student from this batch.'}, status=400)
        ann.specific_students.set(students)
    else:
        ann.specific_students.clear()

    return Response(CounselorAnnouncementSerializer(ann).data)


@api_view(['PATCH'])
@permission_classes([IsAuthenticated])
def trainer_toggle_announcement(request, pk):
    trainer = _get_trainer_for_request(request)
    if not trainer:
        return Response({'error': 'Permission denied'}, status=403)
    try:
        ann = CounselorAnnouncement.objects.get(id=pk, created_by=request.user)
        ann.is_published = not ann.is_published
        ann.save()
        return Response({'is_published': ann.is_published})
    except CounselorAnnouncement.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)


@api_view(['DELETE'])
@permission_classes([IsAuthenticated])
def trainer_delete_announcement(request, pk):
    trainer = _get_trainer_for_request(request)
    if not trainer:
        return Response({'error': 'Permission denied'}, status=403)
    try:
        CounselorAnnouncement.objects.get(id=pk, created_by=request.user).delete()
        return Response(status=204)
    except CounselorAnnouncement.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def trainer_announcement_batches(request):
    trainer = _get_trainer_for_request(request)
    if not trainer:
        return Response({'error': 'Permission denied'}, status=403)

    batches = Batches.objects.filter(faculty=trainer).select_related('course_name').order_by('batch_number')
    return Response([
        {
            'id': batch.id,
            'batch_number': batch.batch_number,
            'batch_timing': batch.batch_timing,
            'course_name': batch.course_name.course_name if batch.course_name else '',
            'display_text': f"{batch.batch_number} - {batch.course_name.course_name if batch.course_name else 'Course'} - {batch.batch_timing}",
        }
        for batch in batches
    ])


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def trainer_announcement_students(request):
    trainer = _get_trainer_for_request(request)
    if not trainer:
        return Response({'error': 'Permission denied'}, status=403)

    batch_id = request.query_params.get('batch')
    students = Students.objects.filter(Q(assigned_staff=trainer) | Q(assigned_batch__faculty=trainer)).distinct()
    if batch_id:
        students = students.filter(assigned_batch_id=batch_id)

    return Response([
        {
            'id': student.id,
            'student_id': student.student_id,
            'name': f"{student.first_name} {student.last_name or ''}".strip(),
            'batch_id': student.assigned_batch_id,
            'batch_number': student.assigned_batch.batch_number if student.assigned_batch else '',
        }
        for student in students.select_related('assigned_batch').order_by('first_name', 'student_id')
    ])



# -- BRANCH ANNOUNCEMENTS (for mentor & student) -------------------------------
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def branch_announcements(request):
    try:
        user = request.user
        branch = None
        allowed_types = ['all']

        student = Students.objects.filter(user=user).first()
        if student:
            branch = student.branch
            allowed_types = ['all', 'students']
        else:
            emp = Employee.objects.filter(user=user).first()
            if emp:
                branch = emp.branch
                if emp.designation.lower() in ['mentor', 'trainer']:
                    allowed_types = ['all', 'mentors', 'staff']
                elif emp.designation.lower() == 'counselor':
                    allowed_types = ['all', 'counselors', 'staff']

        if not branch:
            return Response({'results': []})

        announcements = Announcement.objects.filter(
            is_published=True,
            recipient_type__in=allowed_types
        ).order_by('-created_at')

        visible = []
        for ann in announcements:
            creator_emp = Employee.objects.filter(user=ann.created_by).first()
            if creator_emp and (creator_emp.branch or '').strip().lower() == branch.strip().lower():
                visible.append(ann)

        serializer = AnnouncementSerializer(visible, many=True)
        return Response({'results': serializer.data, 'count': len(visible)})

    except Exception as e:
        logger.exception("Unhandled exception in branch_announcements")
        return Response({
            'error': 'Branch announcements retrieval failed',
            'details': str(e),
            'view': 'branch_announcements'
        }, status=500)
# -- COUNSELOR FORM HELPERS ----------------------------------------------------
@api_view(['GET'])
@permission_classes([IsAuthenticated])
def counselor_branch_batches(request):
    """Return batches in counselor's branch with trainer name for the announcement form."""
    try:
        emp = Employee.objects.get(user=request.user, designation='counselor')
        from connect.models import Batches
        
        # Fetch batches with related faculty (trainer) and course_name
        batches = Batches.objects.filter(branch=emp.branch).select_related('faculty', 'course_name')
        
        data = []
        for batch in batches:
            # Get trainer/faculty name
            trainer_name = "No trainer assigned"
            if batch.faculty:
                trainer_name = f"{batch.faculty.first_name} {batch.faculty.last_name or ''}".strip()
            
            # Get course name
            course_name = batch.course_name.course_name if batch.course_name else "—"
            
            # Get timing display
            timing_display = batch.get_batch_timing_display() if hasattr(batch, 'get_batch_timing_display') else batch.batch_timing
            
            data.append({
                'id': batch.id,
                'batch_number': batch.batch_number,
                'batch_timing': timing_display or batch.batch_timing,
                'trainer_name': trainer_name,
                'course_name': course_name,
                'display_text': f"{batch.batch_number} — {course_name} (Trainer: {trainer_name}) — {timing_display or batch.batch_timing}"
            })
        
        print(f"? Returning {len(data)} batches with trainer names")  # Debug log
        return Response(data)
        
    except Employee.DoesNotExist:
        return Response({'error': 'Permission denied'}, status=403)

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def counselor_branch_students(request):
    """Return only active students (not fully completed) in counselor's branch.
    Partially completed (reassigned) students ARE included — they are still active.
    """
    try:
        emp = Employee.objects.get(user=request.user, designation='counselor')
        from connect.models import Students, CompletedStudent

        # Only FULLY completed students should be excluded
        try:
            fully_completed_emails = set(
                CompletedStudent.objects.filter(completion_type='full').values_list('email', flat=True)
            )
            fully_completed_sids = set(
                CompletedStudent.objects.filter(completion_type='full').values_list('original_student_id', flat=True)
            )
        except DatabaseError:
            logger.exception("Database error while querying CompletedStudent.completion_type in counselor_branch_students")
            fully_completed_emails = set()
            fully_completed_sids = set()

        all_students = Students.objects.filter(branch__iexact=emp.branch)

        batch_id = request.query_params.get('batch')
        if batch_id:
            all_students = all_students.filter(assigned_batch_id=batch_id)

        print(f"[DEBUG] branch={emp.branch}, total={all_students.count()}, fully_completed={len(fully_completed_emails)}")

        active = []
        for s in all_students:
            if s.email in fully_completed_emails or s.student_id in fully_completed_sids:
                continue
            active.append({
                'id':         s.id,
                'name':       f'{s.first_name} {s.last_name}'.strip(),
                'student_id': s.student_id,
            })

        print(f"[DEBUG] active count={len(active)}")
        return Response(active)
    except Employee.DoesNotExist:
        return Response({'error': 'Permission denied'}, status=403)



# -- COURSE TYPES --------------------------------------------------------------

@api_view(['GET'])
@permission_classes([IsAuthenticated])
def course_type_list(request):
    """List all active course types — used in dropdowns."""
    from connect.models import CourseType
    qs = CourseType.objects.filter(is_active=True).order_by('name')
    data = [{'id': ct.id, 'name': ct.name, 'value': ct.value} for ct in qs]
    return Response(data)


@api_view(['GET', 'POST'])
@permission_classes([IsAuthenticated])
def admin_course_types(request):
    from connect.models import CourseType

    if request.method == 'GET':
        qs = CourseType.objects.all().order_by('name')
        data = [
            {
                'id': ct.id,
                'name': ct.name,
                'value': ct.value,
                'is_active': ct.is_active,
                'created_at': ct.created_at,
            }
            for ct in qs
        ]
        return Response(data)

    name = request.data.get('name', '').strip()
    value = request.data.get('value', '').strip().lower().replace(' ', '_')

    if not name:
        return Response({'error': 'Name is required'}, status=400)

    if not value:
        value = name.lower().replace(' ', '_').replace('&', 'and').replace('/', '_')

    value = value.replace('/', '_')

    if CourseType.objects.filter(value=value).exists():
        return Response({'error': 'A course type with this value already exists'}, status=400)

    if CourseType.objects.filter(name__iexact=name).exists():
        return Response({'error': 'A course type with this name already exists'}, status=400)

    ct = CourseType.objects.create(
        name=name,
        value=value,
        is_active=True
    )

    return Response({
        'id': ct.id,
        'name': ct.name,
        'value': ct.value,
        'is_active': ct.is_active,
        'created_at': ct.created_at,
    }, status=201)
@api_view(['PATCH', 'DELETE'])
@permission_classes([IsAuthenticated])
def admin_course_type_detail(request, pk):
    """Admin: toggle active / delete a course type."""
    from connect.models import CourseType
    try:
        ct = CourseType.objects.get(pk=pk)
    except CourseType.DoesNotExist:
        return Response({'error': 'Not found'}, status=404)

    if request.method == 'DELETE':
        ct.delete()
        return Response(status=204)

    # PATCH — toggle is_active or rename
    if 'is_active' in request.data:
        ct.is_active = request.data['is_active']
    if 'name' in request.data:
        ct.name = request.data['name'].strip()
    ct.save()
    return Response({'id': ct.id, 'name': ct.name, 'value': ct.value, 'is_active': ct.is_active})
